#!/usr/bin/env python3
"""Embedded Word report, exact-pixel output panels and target-skeleton overlays."""
from __future__ import annotations
import argparse
import hashlib
import json
from pathlib import Path
import sys
import xml.etree.ElementTree as ET
from zipfile import ZipFile

PROJECT=Path(__file__).resolve().parents[2]
sys.path.insert(0,str(PROJECT))
import numpy as np
from PIL import Image,ImageDraw,ImageFont
from docx import Document
from docx.shared import Cm,Pt
from docx.enum.section import WD_SECTION_START
from scripts.eval.eval_hi4d_depth_combinations import DEFAULT,PREVIOUS,ARMS,LABELS,FONT,sha,write,read_rows,verify_lock,key
from scripts.eval.score_hi4d_depth_combinations import records,CONTRASTS
from scripts.reports.build_hi4d_native_binding_report import crop,configure,table
from scripts.reports.build_hi4d_separate_depth_report import centered

REPORT=Path('/home/shangguanrz/nihaochen/docs/Hi4D双人_深度输入组合对比_20261013.docx')
CHANGE={'improved':'改善','equal':'相当','worse':'变差','uncertain':'无法判断'}


def input_tile(refs,language):
    canvas=Image.new('RGB',(1024,700),'white');d=ImageDraw.Draw(canvas)
    labels={'A':'A RGB','S1':'人物1骨架','S2':'人物2骨架','depth_joint':'合并深度','D1':'人物1深度','D2':'人物2深度'}
    for i,r in enumerate([*refs,{'role':'Prompt'}]):
        x=12+(i%3)*336;y=12+(i//3)*336
        if r['role']=='Prompt':
            d.rounded_rectangle((x+10,y+48,x+314,y+290),radius=18,fill=(212,231,253),outline=(105,153,210),width=3)
            centered(d,'Prompt\n'+('英文' if language=='en' else '中文'),(x+10,y+48,x+314,y+290),ImageFont.truetype(FONT,48),(30,65,110))
        else:
            canvas.paste(Image.open(r['path']).convert('RGB').resize((300,300),Image.Resampling.LANCZOS),(x+12,y+32))
            centered(d,labels[r['role']],(x,y,x+324,y+30),ImageFont.truetype(FONT,34))
    return canvas


def make_panel(root,case,rows,manual):
    pid=case['pair_id'];jobs={(r['method'],r['language']):r for r in rows if r['pair_id']==pid}
    paths=[case['source']['rgb'],case['target']['rgb'],*[jobs[a]['image'] for a in ARMS]]
    folder=root/'inputs/1024'/pid
    side=1024;gap=10;left=280;header=145;inputs=710;notes=145
    starts=[header,header+inputs+gap]
    starts.extend(starts[1]+i*(side+gap) for i in range(1,4))
    canvas=Image.new('RGB',(left+7*(side+gap),starts[-1]+side+notes),'white');draw=ImageDraw.Draw(canvas)
    font=ImageFont.truetype(FONT,54);small=ImageFont.truetype(FONT,40)
    headings=['输入 RGB：A','目标 RGB：B\n仅核对和展示',*[LABELS[a] for a in ARMS]]
    checks=[]
    for col,(path,heading) in enumerate(zip(paths,headings)):
        x=left+col*(side+gap);centered(draw,heading,(x,0,x+side,header-4),font)
        if col<2:
            tile=Image.open(folder/('A.png' if col==0 else 'B.png')).convert('RGB').resize((650,650),Image.Resampling.LANCZOS)
            canvas.paste(tile,(x+187,header+25))
        else:canvas.paste(input_tile(jobs[ARMS[col-2]]['references'],ARMS[col-2][1]),(x,header))
        full=Image.open(folder/('A.png' if col==0 else 'B.png')).convert('RGB') if col<2 else Image.open(path).convert('RGB')
        canvas.paste(full,(x,starts[1]))
        assert np.array_equal(np.asarray(canvas.crop((x,starts[1],x+side,starts[1]+side))),np.asarray(full))
        if col==0:
            overlay=Image.new('RGB',(side,side),(241,241,241));centered(ImageDraw.Draw(overlay),'不适用',(0,0,side,side),font,(105,105,105))
        else:
            overlay_path=folder/'B_target_overlay.png' if col==1 else root/'overlays'/Path(path).name
            overlay=Image.open(overlay_path).convert('RGB')
            checks.append({'column':col,'overlay':str(overlay_path),'overlay_sha256':sha(overlay_path)})
        canvas.paste(overlay,(x,starts[2]))
        assert np.array_equal(np.asarray(canvas.crop((x,starts[2],x+side,starts[2]+side))),np.asarray(overlay))
        raw=Image.open(path).convert('RGB');rec=records(root,path);review=manual['images'][path]
        for identity in (0,1):
            pp=next(p for p in review['people'] if p['source_id']==identity);idx=pp.get('face_detection_index')
            f=rec['faces']['detections'][idx] if idx is not None and pp.get('face_real') else None
            if f:
                face=crop(raw,f['box_native'],side)
            else:
                face=Image.new('RGB',(side,side),(238,238,238))
                centered(ImageDraw.Draw(face),'无法可靠定位\n该人物的实际脸部',(0,0,side,side),font,(95,95,95))
            y=starts[3+identity];canvas.paste(face,(x,y))
            label=None
            if f and col>=2:
                if pp.get('original_person_residue'):label='未移走的原人物：目标动作由新增人物执行'
                elif pp.get('duplicate_source_appearance'):label='该外观出现重复：显示已确认实例'
                elif pp.get('identity_status') in ('mixed','uncertain'):label='身份或外观不确定，仅供观察'
                elif not pp.get('face_identity_usable'):label='侧后脸或遮挡，仅供观察'
            if f and f['short_side_at_512']<32:label=(label+'；小脸' if label else '原图脸较小，放大仅供观察')
            if label:
                draw.rectangle((x,y,x+side,y+88),fill='white');centered(draw,label,(x,y,x+side,y+85),small,(115,75,25))
        for y in starts:draw.line((x,y-3,x+side,y-3),fill=(206,212,219),width=3)
    labels=('实际输入','输出整图\nA/B为原图','B目标骨架\n叠加核对','A人物1脸\nB列为目标1','A人物2脸\nB列为目标2')
    for i,label in enumerate(labels):
        end=starts[i+1] if i<4 else starts[i]+side
        centered(draw,label,(0,starts[i],left-10,end),small)
    centered(draw,f'{pid} · 1024 · 原生Qwen · 直接对应，无swap\nB目标骨架叠加仅供对位核对；深度组未输入骨架；不是生成图检测结果',
        (left,starts[-1]+side+10,canvas.width,canvas.height-3),small)
    out=root/'panels'/f'{pid}_1024.png';out.parent.mkdir(parents=True,exist_ok=True);canvas.save(out)
    prev=root/'previews'/f'{pid}_1024.jpg';prev.parent.mkdir(parents=True,exist_ok=True)
    canvas.resize((3200,round(canvas.height*3200/canvas.width)),Image.Resampling.LANCZOS).save(prev,quality=96,subsampling=0)
    return out,prev,checks


def build(root,report):
    verify_lock(root)
    rows=read_rows(root/'generations.jsonl');cases=json.loads((root/'data/cases.json').read_text())
    manual=json.loads((root/'review/manual.json').read_text());conclusions=json.loads((root/'review/conclusions.json').read_text())
    summary=json.loads((root/'evaluation/summary.json').read_text())
    assert len(rows)==30 and len({key(r) for r in rows})==30 and len(cases)==6
    panels=[make_panel(root,c,rows,manual) for c in cases]
    doc=Document();configure(doc);doc.sections[0].footer.paragraphs[0].text='Hi4D 双人编辑 · 骨架与三种深度输入'
    doc.add_heading('双人动作编辑：深度能替代或补充骨架吗？',0)
    doc.add_paragraph('6对固定案例｜1024×1024｜五组原生Qwen｜英文提示词',style='Subtitle')
    doc.add_heading('先看结论',1)
    for text in conclusions['main']:doc.add_paragraph(text,style='List Bullet')
    doc.add_heading('这次比较什么',1)
    doc.add_paragraph('五组分别为：分开骨架、分开骨架＋合并深度、合并深度、分开深度，以及合并＋分开深度。'
        'A照片始终提供两人的身份、服装和场景；B仅供核对与展示，不输入模型。五组全部重新生成，未复用旧输出。')
    doc.add_paragraph('独立深度保留各自完整人体表面，合并深度反映两人共同遮挡后的可见表面；三张深度使用相同相机、灰度尺度和场景坐标。'
        '没有修复原有姿态估计。英文提示词共用外观与背景保留要求，仅按图片角色调整指认说明，不人工描述动作和手臂绕行。')
    doc.add_paragraph('本轮没有训练，没有启用Adapter、LoRA或提示词重写。统一30步生成采样、种子42、true CFG 1.0、关闭KV缓存。'
        '六例复用三张A，共六位源人物；一次固定种子实验不证明跨种子的稳定性。')
    doc.add_heading('五组结果',1)
    ss=[next(s for s in summary if (s['method'],s['language'])==a) for a in ARMS]
    fields=[('恰好两人 ↑','exactly_two_people',0,'/6',True),('多出人物 ↓','extra_people_images',0,'/6',False),
        ('人物对应正确 ↑','correspondence_correct',0,'/6',True),
        ('明显异常图片 ↓','artifact_images',0,'/6',False),('动作 PCK ↑','pose_pck_002_common',3,'',True),
        ('动作 NME ↓','pose_nme_common',3,'',False),('共同可评动作 n','pose_common_n',0,'',None),
        ('背景误差 ↓','background_mae',2,'',False),('共同可靠脸 n','face_common_n',0,'',None),
        ('源脸相似度 ↑','face_cosine_common',3,'',None),('耗时 / 秒 ↓','seconds',1,'',False),
        ('峰值显存 / GiB ↓','peak_cuda_gib',2,'',False)]
    values=[];best=set()
    for i,(label,k,digits,suffix,up) in enumerate(fields):
        values.append([label,*[('无法判断' if s[k] is None else f'{s[k]:.{digits}f}{suffix}') for s in ss]])
        nums=[s[k] for s in ss if s[k] is not None]
        if up is not None and nums:
            target=(max if up else min)(nums);best|={(i,j+1) for j,s in enumerate(ss) if s[k]==target}
    t=table(doc,['指标','分开骨架','骨架＋深度','合并深度','分开深度','合并＋分开'],values,best)
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for run in p.runs:run.font.size=Pt(9)
    doc.add_paragraph('PCK越高、NME越低越贴合B的二维关键点；均值使用五组共同可评的“人物×案例”，n不是独立人数。'
        '人物按A外观指认，不按最佳动作分数重新配人。B不是理想生成图，不计算与B的PSNR或SSIM。'
        '姿态与背景统一在512网格评价；脸部从实际输出位置检测，可靠脸不足3例不排名。'
        '人物数量、对应及异常为AI辅助逐图核对结果，可按后附图片复核，未做独立人工盲评。',style='Caption')
    doc.add_paragraph(conclusions['metric_caution'])
    doc.add_heading('逐例观察与边界',1)
    table(doc,['案例','主要观察'],[[c['pair_id'],conclusions['cases'][c['pair_id']]['brief']] for c in cases])
    for text in conclusions['limits_next']:doc.add_paragraph(text)
    section=doc.add_section(WD_SECTION_START.NEW_PAGE)
    section.page_width=Cm(42);section.page_height=Cm(29.7)
    section.left_margin=section.right_margin=Cm(1.4);section.top_margin=section.bottom_margin=Cm(1.4)
    for index,(c,(png,preview,checks)) in enumerate(zip(cases,panels)):
        if index:doc.add_page_break()
        pid=c['pair_id'];doc.add_heading(f'{pid}｜输入、结果与目标骨架对位',1)
        doc.add_paragraph(conclusions['cases'][pid]['detail'])
        doc.add_picture(str(png),width=Cm(34))
        doc.paragraphs[-1].alignment=1
        doc.add_paragraph('第三行统一用B目标骨架按原场景坐标叠加，未根据输出拟合；仅供对位核对，不表示深度组输入了骨架。'
            '脸图按A身份从实际位置裁剪，无法确认时留空。放大不增加原图细节。',style='Caption')
    report.parent.mkdir(parents=True,exist_ok=True);doc.save(report)
    with ZipFile(report) as z:
        assert z.testzip() is None
        media=[p for p in z.namelist() if p.startswith('word/media/')]
        hashes={hashlib.sha256(z.read(p)).hexdigest() for p in media}
        assert len(media)==6 and all(sha(png) in hashes for png,_,_ in panels)
        assert z.read('word/document.xml').count(b'<w:drawing>')==6
        assert b'TargetMode="External"' not in z.read('word/_rels/document.xml.rels')
    reopened=Document(report);assert len(reopened.inline_shapes)==6 and len(reopened.tables)==2
    for i,(_,k,digits,suffix,_) in enumerate(fields):
        for j,s in enumerate(ss):assert reopened.tables[0].rows[i+1].cells[j+1].text==values[i][j+1]
    write(root/'report_checks.json',{'report':str(report),'report_sha256':sha(report),'zip_valid':True,
        'embedded_lossless_png_count':6,'external_images':False,'formal_results':30,'target_skeleton_not_output_estimation':True,
        'depth_routes_have_no_skeleton_reference':True,'full_output_cells_exact_pixels':True,'overlay_cells_exact_pixels':True,
        'table_values_match_summary':True,
        'panels':[{'png':str(p),'preview':str(v),'png_sha256':sha(p),'overlay_sources':q} for p,v,q in panels]})
    print(report)


def build_preview(root,report,preview_report):
    """A separately labelled reading copy; the full lossless report is retained."""
    checks=json.loads((root/'report_checks.json').read_text())
    if sha(report)!=checks['report_sha256']:raise ValueError('Full report changed')
    by_hash={p['png_sha256']:Path(p['preview']) for p in checks['panels']}
    temp=root/'report_preview_work.docx';replacements={};original_hash=sha(report)
    with ZipFile(report) as source:
        for name in source.namelist():
            if name.startswith('word/media/'):
                h=hashlib.sha256(source.read(name)).hexdigest()
                if h not in by_hash:raise ValueError('Unexpected report media')
                replacements[name]=(str(Path(name).with_suffix('.jpg')),by_hash[h].read_bytes())
        assert len(replacements)==6
        with ZipFile(temp,'w',compression=8) as dest:
            for name in source.namelist():
                value=source.read(name)
                if name in replacements:
                    new,value=replacements[name];dest.writestr(new,value);continue
                if name=='word/_rels/document.xml.rels':
                    for before,(after,_) in replacements.items():
                        value=value.replace(before.removeprefix('word/').encode(),after.removeprefix('word/').encode())
                if name=='[Content_Types].xml':
                    tree=ET.fromstring(value);ns='{http://schemas.openxmlformats.org/package/2006/content-types}'
                    if not any(e.get('Extension')=='jpg' for e in tree):
                        ET.SubElement(tree,ns+'Default',Extension='jpg',ContentType='image/jpeg')
                    value=ET.tostring(tree,encoding='utf-8',xml_declaration=True)
                dest.writestr(name,value)
    doc=Document(temp);doc.paragraphs[0].add_run('（轻量阅读版）')
    paragraph=doc.add_paragraph('本文件使用聊天预览尺寸的JPEG，方便阅读；数据和结论与高清版相同。高清PNG、原尺寸输出和无损Word另行保留。',style='Caption')
    doc.element.body.insert(1,paragraph._p)
    doc.save(preview_report);temp.unlink()
    with ZipFile(preview_report) as z:
        assert z.testzip() is None
        media=[n for n in z.namelist() if n.startswith('word/media/')]
        assert len(media)==6 and all(n.endswith('.jpg') for n in media)
        assert {hashlib.sha256(z.read(n)).hexdigest() for n in media}=={hashlib.sha256(v).hexdigest() for _,v in replacements.values()}
    small=Document(preview_report);full=Document(report)
    assert [[c.text for row in t.rows for c in row.cells] for t in small.tables]==[[c.text for row in t.rows for c in row.cells] for t in full.tables]
    assert sha(report)==original_hash
    write(root/'preview_report_checks.json',{'report':str(preview_report),'sha256':sha(preview_report),
        'size_mib':preview_report.stat().st_size/2**20,'embedded_jpeg_count':6,'tables_identical':True,
        'full_report_unchanged':True,'full_report_sha256':original_hash,'preview_only_no_additional_generation':True})
    print(preview_report)


def main():
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('--output',type=Path,default=DEFAULT)
    p.add_argument('--report',type=Path,default=REPORT);p.add_argument('--preview-only',action='store_true')
    args=p.parse_args()
    if args.preview_only:build_preview(args.output,args.report,args.report.with_stem(args.report.stem+'_轻量版'))
    else:build(args.output,args.report)

if __name__=='__main__':main()
