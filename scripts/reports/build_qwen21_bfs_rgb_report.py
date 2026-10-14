#!/usr/bin/env python3
"""Concise Word report for the fixed 1024 A-RGB + B-RGB BFS comparison."""
from __future__ import annotations
import argparse
import hashlib
import json
from pathlib import Path
import zipfile

from docx import Document
from docx.shared import Cm, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from scripts.eval.eval_qwen21_bfs_rgb import DEFAULT, PAIR_IDS, sha256
from scripts.reports.build_qwen21_native_value_report import font, para, number
from scripts.reports import build_qwen21_bfs_skeleton_report as shared_report

OUTPUT = Path('/home/shangguanrz/nihaochen/docs/BFS_LoRA两张RGB输入_对比结果_20261010.docx')


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--root', type=Path, default=DEFAULT)
    parser.add_argument('--output', type=Path, default=OUTPUT)
    args = parser.parse_args()
    root = args.root
    stats = json.loads((root / 'evaluation/metric_summary.json').read_text())
    rows = json.loads((root / 'evaluation/evaluated_samples.json').read_text())
    observations = json.loads((root / 'evaluation/observations.json').read_text())
    panels = json.loads((root / 'report_figures/manifest.json').read_text())
    if len(rows) != 12 or len(panels) != 4 or not stats['formal_identity_review_complete']:
        raise ValueError('Report needs all 12 outputs and completed source/B/generated face review')
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(1.4)
    sec.left_margin = sec.right_margin = Cm(1.4)
    sec.header_distance = sec.footer_distance = Cm(.6)
    normal = doc.styles['Normal']
    normal.font.name = 'Noto Sans CJK SC'
    normal.font.size = Pt(10)
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Noto Sans CJK SC')
    settings = OxmlElement('w:doNotAutoCompressPictures')
    doc.settings.element.append(settings)
    footer = sec.footer.paragraphs[0]
    font(footer.add_run('BFS 两张照片输入对比 · '), 8)
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    para(doc, '给两张真实照片，BFS LoRA 有帮助吗？', 17, True)
    para(doc, 'A-RGB＋B-RGB · 四对固定案例 · 1024×1024 · 2026年10月10日', 9)
    para(doc, '这次实际做了什么', 11, True)
    para(doc, '第一张照片 A 提供要保留的人物、服装、体型和场景；第二张照片 B 展示另一个人的目标动作和头向。三组都给 Qwen 完整 A、完整 B，使用完全相同的逐样本详细提示词。比较原生、BFS Body V1.0、BFS Head V1.1，两个 LoRA 分别加载，强度均为1。', 10)
    para(doc, '本轮没有训练，没有二次编辑，没有启用项目自训的 Body／Face Adapter，也没有加入骨架或 A 脸图。BFS 是外部已训练权重，其 Body／Head 名称不代表项目控制模块。', 10)
    para(doc, '三组统一：1024×1024、30步采样、种子42、true CFG 1.0，等比例缩放并补边，关闭KV缓存。共12张结果：4张原生图经输入、提示词及配置核验后复用，8张LoRA图本轮生成。', 9)
    para(doc, observations['headline'], 11, True)
    for text in observations['conclusions']:
        para(doc, text, 10)
    shared_report.LABELS['native'] = '原生 A+B'
    para(doc, '关键数据｜两种数据集分别看', 11, True)
    for dataset in ('iper', 'tiktok'):
        shared_report.table(doc, stats['groups'][f'{dataset}_1024_steps30'], f'{"iPER" if dataset == "iper" else "TikTok"} · 每组2对')
    para(doc, '动作PCK越高、位置误差和背景误差越低越好。面部相似度是特征余弦，不是身份保持百分比；n是三组共同可靠的人脸配对数。每个数据集最多2对，身份分数只作案例参考，不加粗排名；检出脸不等于脸保持得好。其他可比指标按方向加粗最优，可能并列。', 8.5)
    para(doc, 'TikTok的严格位置PCK很低，但三组都明显改变了动作。它要求关节落在目标的绝对位置附近，跨人物的体型、取景和人物大小会影响得分，不能把0分理解为“没有听从动作”。', 8.5)
    para(doc, '* 原生耗时和显存为核验后复用的历史测量；LoRA为本轮测量。姿态与背景统一缩至512评价，人脸使用原尺寸。B不是理想结果图，没有计算相对B的PSNR、SSIM或人体重建误差。', 8.5)
    doc.add_page_break()
    para(doc, '人物是否向 B 靠拢？', 14, True)
    for text in observations['identity_interpretation']:
        para(doc, text, 10)
    table = doc.add_table(rows=1, cols=5)
    table.style = 'Table Grid'
    for cell, text in zip(table.rows[0].cells, ['案例', '方案', '与A原始相似度', '与B原始相似度', 'A身份共同子集？']):
        font(cell.paragraphs[0].add_run(text), 8.5, True)
    for pid in PAIR_IDS:
        for method in shared_report.METHODS:
            row = next(r for r in rows if r['pair_id'] == pid and r['method'] == method)
            b_score = number(row['metrics']['identity_b_cosine_raw'], 3)
            if not row['face_selection']['target']['formal_reliable']:
                b_score += '*'
            values = [pid, shared_report.LABELS[method], number(row['metrics']['identity_a_cosine_raw'], 3),
                      b_score,
                      '是' if row['formal_common_reliable_identity_sample'] else '否，仅原始分']
            for cell, text in zip(table.add_row().cells, values):
                font(cell.paragraphs[0].add_run(text), 8.5)
    para(doc, '这张表保留检测原始分数，便于追溯；不可可靠评价的脸也没有隐藏。* 表示B脸被遮挡，分数仅供追溯；“A身份共同子集”只要求A和三组生成脸可靠，不表示B也可评价。与B更接近只是串脸线索，仍需结合A、B本身相似程度、头向、模糊程度和实际外观检查，不能单凭两列大小认定串脸。', 8.5)
    para(doc, '与既有 A＋骨架结果的关系', 11, True)
    for text in observations['skeleton_context']:
        para(doc, text, 10)
    para(doc, '范围与正确性', 11, True)
    para(doc, '仅4对案例，A来自2个人物；这是决定是否扩大验证的小规模检查，不足以证明跨人物泛化。BFS作者设计的换人／换头图片分工与本项目不同，本轮也不能代表官方工作流或项目专用LoRA重新训练后的能力。', 9)
    checks = []
    for method in shared_report.METHODS[1:]:
        check = json.loads((root / 'development' / method / 'preflight.json').read_text())
        checks.append(f'{shared_report.LABELS[method]}零强度误差{check["zero_strength_max_abs"]:.6g}，强度1预测变化{check["strength_one_max_abs_change"]:.3f}')
    para(doc, '加载检查：' + '；'.join(checks) + '。全部完整采样通过有限值检查，参数冻结，自训Adapter调用为0；B-RGB按设计作为第二张参考输入，所有低分结果保留。', 9)
    para(doc, '后续建议', 11, True)
    for text in observations['recommendations']:
        para(doc, text, 10)
    for pid in PAIR_IDS:
        doc.add_page_break()
        panel = next(p for p in panels if p['pair_id'] == pid)
        if sha256(panel['path']) != panel['sha256']:
            raise ValueError('Report panel checksum differs')
        para(doc, pid + '｜A、B 与三种生成结果', 14, True)
        para(doc, observations['captions'][pid], 10)
        para(doc, '第一行整图；第二行从每张图实际检测位置裁出人脸并等宽放大。检测不足或脸太小会单独标明；放大不会恢复原图缺失的细节。A、B两张完整照片均是本轮实际输入。', 9)
        doc.add_picture(panel['path'], width=Cm(18.2))
        para(doc, '观察重点：动作和目标头向、A的脸与衣服、背景，以及是否出现B的人物外观。', 9)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.output)
    with zipfile.ZipFile(args.output) as archive:
        assert archive.testzip() is None
        media = [n for n in archive.namelist() if n.startswith('word/media/')]
        assert len(media) == 4
        assert {hashlib.sha256(archive.read(n)).hexdigest() for n in media} == {p['sha256'] for p in panels}
    result = {'document': str(args.output), 'sha256': sha256(args.output), 'bytes': args.output.stat().st_size,
              'embedded_images': 4, 'lossless_embedded_media_verified': True,
              'tables': len(doc.tables), 'metric_summary_sha256': sha256(root / 'evaluation/metric_summary.json'),
              'observations_sha256': sha256(root / 'evaluation/observations.json')}
    (root / 'word_report_manifest.json').write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n')
    print(json.dumps(result, ensure_ascii=False), flush=True)


if __name__ == '__main__':
    main()
