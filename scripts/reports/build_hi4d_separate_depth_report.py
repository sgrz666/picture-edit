#!/usr/bin/env python3
"""Concise embedded-image Word report and six full-resolution comparison panels."""
from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path
import sys
from zipfile import ZipFile

PROJECT=Path(__file__).resolve().parents[2]
sys.path.insert(0,str(PROJECT))
from PIL import Image,ImageDraw,ImageFont
from docx import Document
from docx.shared import Cm
from scripts.eval.eval_hi4d_separate_depth import DEFAULT,METHODS,sha,write,read_rows,verify_lock
from scripts.eval.score_hi4d_separate_depth import records
from scripts.reports.build_hi4d_native_binding_report import crop,configure,table,number

FONT='/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc'
REPORT=Path('/home/shangguanrz/nihaochen/docs/Hi4D清晰分开骨架与共同深度_对比结果_20261012.docx')
LABELS={'separate':'分开骨架基线','separate_depth':'分开骨架＋共同深度'}
CHANGE={'improved':'改善','equal':'相当','worse':'变差','uncertain':'无法判断'}


def centered(draw,text,rect,font,fill='black'):
    b=draw.multiline_textbbox((0,0),text,font=font,spacing=8,align='center')
    x0,y0,x1,y1=rect
    draw.multiline_text(((x0+x1-b[2]+b[0])/2,(y0+y1-b[3]+b[1])/2-b[1]),text,font=font,fill=fill,spacing=8,align='center')


def input_tile(refs):
    canvas=Image.new('RGB',(1024,700),'white');d=ImageDraw.Draw(canvas)
    font=ImageFont.truetype(FONT,34);labels={'A':'A RGB','S1':'人物1骨架','S2':'人物2骨架','depth_joint':'共同深度'}
    items=[*refs,{'role':'Prompt'}]
    for i,r in enumerate(items):
        x=12+(i%3)*336;y=12+(i//3)*336
        if r['role']=='Prompt':
            d.rounded_rectangle((x+10,y+48,x+314,y+290),radius=18,fill=(212,231,253),outline=(105,153,210),width=3)
            centered(d,'Prompt',(x+10,y+48,x+314,y+290),ImageFont.truetype(FONT,49),(30,65,110))
        else:
            canvas.paste(Image.open(r['path']).convert('RGB').resize((300,300),Image.Resampling.LANCZOS),(x+12,y+32))
            centered(d,labels[r['role']],(x,y,x+324,y+30),font)
    return canvas


def make_panel(root,case,rows,manual):
    pid=case['pair_id'];jobs={r['method']:r for r in rows if r['pair_id']==pid}
    paths=[case['source']['rgb'],case['target']['rgb'],*[jobs[m]['image'] for m in METHODS]]
    folder=root/'inputs/1024'/pid
    side=1024;gap=10;left=255;header=138;inputs=710;notes=120
    starts=[header,header+inputs+gap,header+inputs+gap+side+gap,header+inputs+gap+2*(side+gap)]
    canvas=Image.new('RGB',(left+4*(side+gap),starts[-1]+side+notes),'white');draw=ImageDraw.Draw(canvas)
    font=ImageFont.truetype(FONT,59);small=ImageFont.truetype(FONT,44)
    headings=['输入 RGB：A','目标 RGB：B\n仅核对和展示',*[LABELS[m] for m in METHODS]]
    for col,(path,heading) in enumerate(zip(paths,headings)):
        x=left+col*(side+gap)
        centered(draw,heading,(x,0,x+side,header-4),font)
        if col<2:
            tile=Image.open(folder/('A.png' if col==0 else 'B.png')).convert('RGB').resize((650,650),Image.Resampling.LANCZOS)
            canvas.paste(tile,(x+187,header+25))
        else:canvas.paste(input_tile(jobs[METHODS[col-2]]['references']),(x,header))
        full=Image.open(folder/('A.png' if col==0 else 'B.png')) if col<2 else Image.open(path)
        canvas.paste(full.convert('RGB'),(x,starts[1]))
        raw=Image.open(path).convert('RGB');rec=records(root,path);review=manual['images'][path]
        for identity in (0,1):
            pp=next(p for p in review['people'] if p['source_id']==identity)
            idx=pp.get('face_detection_index')
            f=rec['faces']['detections'][idx] if idx is not None and pp.get('face_real') else None
            # Actual native-resolution detector box from this image, never B's box.
            if f:
                im=crop(raw,f['box_native'],side)
            else:
                im=Image.new('RGB',(side,side),(238,238,238));missing_draw=ImageDraw.Draw(im)
                missing_label='未确认对应人脸' if rec['faces']['detections'] else '未检出人脸'
                centered(missing_draw,missing_label,(0,0,side,side),font,(90,90,90))
            canvas.paste(im,(x,starts[2+identity]))
            if f is None:
                draw.rectangle((x,starts[2+identity],x+side,starts[2+identity]+90),fill='white')
                centered(draw,'脸部无法可靠定位',(x,starts[2+identity],x+side,starts[2+identity]+85),small,(95,95,95))
            elif col>=2 and (not pp.get('face_identity_usable') or pp.get('identity_status')=='duplicated'
                             or pp.get('duplicate_source_appearance')):
                label=('人物重复：仅显示原位置可见脸' if pp.get('identity_status')=='duplicated' or pp.get('duplicate_source_appearance')
                       else '遮挡或侧后脸：仅供定位参考')
                draw.rectangle((x,starts[2+identity],x+side,starts[2+identity]+90),fill='white')
                centered(draw,label,(x,starts[2+identity],x+side,starts[2+identity]+85),small,(115,75,25))
        for y in starts:draw.line((x,y-3,x+side,y-3),fill=(206,212,219),width=3)
    for i,label in enumerate(('实际输入','输出整图\nA/B 为原图','人物1脸\nB列为目标1','人物2脸\nB列为目标2')):
        end=starts[i+1] if i<3 else starts[i]+side
        centered(draw,label,(0,starts[i],left-10,end),small)
    decision=manual['pairwise'][pid]
    footer=f'{pid} · 1024 · 直接对应 · 加深度后遮挡：{CHANGE[decision["occlusion_change"]]}'
    if decision.get('assignment_text'):footer+='\n'+decision['assignment_text']
    centered(draw,footer,
             (left,starts[-1]+side+10,canvas.width,canvas.height-3),small)
    out=root/'panels'/f'{pid}_1024.png';out.parent.mkdir(parents=True,exist_ok=True);canvas.save(out)
    prev=root/'previews'/f'{pid}_1024.jpg';prev.parent.mkdir(parents=True,exist_ok=True)
    canvas.resize((2400,round(canvas.height*2400/canvas.width)),Image.Resampling.LANCZOS).save(prev,quality=95,subsampling=0)
    return out,prev


def build(root,report):
    verify_lock(root)
    rows=read_rows(root/'generations.jsonl');cases=json.loads((root/'data/cases.json').read_text())
    manual=json.loads((root/'review/manual.json').read_text())
    summary=json.loads((root/'evaluation/summary.json').read_text())
    visual=json.loads((root/'evaluation/paired_visual_summary.json').read_text())
    conclusions=json.loads((root/'review/conclusions.json').read_text())
    assert len(rows)==12 and len(cases)==6 and set(METHODS)=={r['method'] for r in summary}
    panels=[make_panel(root,c,rows,manual) for c in cases]
    doc=Document();configure(doc);doc.sections[0].footer.paragraphs[0].text='Hi4D 双人交互 · 分开骨架与共同深度'
    doc.add_heading('共同深度能否改善双人交互？',0)
    doc.add_paragraph('Hi4D · 6对固定案例 · 1024×1024 · 原生Qwen对照',style='Subtitle')
    doc.add_heading('先看结果',1)
    for text in conclusions['main']:doc.add_paragraph(text,style='List Bullet')
    counts=visual['occlusion_change_counts']
    icounts=visual['interaction_change_counts']
    doc.add_paragraph('加入深度后的整体交互：'+ '、'.join(f'{CHANGE[k]}{icounts[k]}对' for k in CHANGE)+'；'
                      '单看前后遮挡：'+ '、'.join(f'{CHANGE[k]}{counts[k]}对' for k in CHANGE)+'。'
                      '逐图观察由AI辅助完成，并非独立人工盲评；看不清的部位不强行判断。')
    doc.add_heading('这次比较了什么',1)
    doc.add_paragraph('A提供要保留的两个人及场景。基线给Qwen“完整A图＋人物1骨架＋人物2骨架”；另一组只多给一张两人共同深度图。'
                      '骨架说明每人的关节位置，深度说明交叠表面谁在前、谁在后。B原图只用于检查和展示，没有交给模型。')
    doc.add_paragraph('骨架按原坐标重绘：提亮线条、加粗手脚线、去掉图内标签，原本缺失的关节仍保留缺失。'
                      '提示词只用服装、发型说明“谁对应哪张骨架”，加通用的动作、交互和外观保留要求；'
                      '没有人工逐条描述动作或手臂遮挡，也没有自动重写提示词。'
                      '同一张A搭配两种目标骨架时，提示词逐字相同，动作变化只来自条件图。')
    doc.add_paragraph('两组均为30步采样、种子42、true CFG 1.0，关闭KV缓存。'
                      '本轮没有训练，没有启用项目自训的Body/Face Adapter，也没有LoRA；只做直接对应，不做swap。'
                      '六对案例来自三张A图，共六位源人物；同一个A搭配两个目标，不能当成六组独立人物。')
    doc.add_heading('数据对比',1)
    rr=[next(s for s in summary if s['method']==m) for m in METHODS]
    metrics=[('人物对应正确 ↑','correspondence_correct',0,'/6',True),
             ('存在明显异常的图片 ↓','artifact_images',0,'/6',False),
             ('动作 PCK ↑','pose_pck_002_common',3,'',True),('动作 NME ↓','pose_nme_common',3,'',False),
             ('共同动作人数 n','pose_common_n',0,'',None),('有效关节检出比例 ↑','pose_joint_recall_common',3,'',True),
             ('背景误差 ↓','background_mae',2,'',False),('源脸相似度 ↑','face_cosine_common',3,'',None),
             ('共同可靠脸 n','face_common_n',0,'',None),('单图耗时 / 秒 ↓','seconds',1,'',False),
             ('峰值显存 / GiB ↓','peak_cuda_gib',2,'',False)]
    values=[];best=set()
    for i,(label,k,digits,suffix,up) in enumerate(metrics):
        values.append([label,*[('无法判断' if r[k] is None else f'{r[k]:.{digits}f}{suffix}') for r in rr]])
        vv=[r[k] for r in rr if r[k] is not None]
        if up is not None and vv:
            target=(max if up else min)(vv)
            best|={(i,j+1) for j,r in enumerate(rr) if r[k]==target}
    numeric=table(doc,['指标','分开骨架','分开骨架＋共同深度'],values,best)
    numeric.autofit=False
    for i,width in enumerate((7.,5.5,5.5)):
        numeric.columns[i].width=Cm(width)
        for row in numeric.rows:row.cells[i].width=Cm(width)
    doc.add_paragraph('“对应正确”只核对谁采用哪副骨架，不代表动作细节或脸已完全保留。异常主要指额外人物、重复头、明显服装丢失。'
                      'PCK越高、NME越低，关节越接近目标骨架；仅在两组都能按A外观辨认并检出的人物共同子集比较。'
                      '对应错误仍按指令目标计分，不按生成位置重新配人。n是“人物×案例”次数，不是独立人数。背景误差按512网格计算，排除补边及A/B人物区域。'
                      '人脸要求原图与输出都真实可见，且映射到512后脸框短边至少32像素；共同可靠脸不足3时不作身份排名。'
                      'B中的人不是理想输出，不计算与B的PSNR或SSIM。',style='Caption')
    doc.add_heading('六对案例分别发生了什么',1)
    vv=[]
    for c in cases:
        r=manual['pairwise'][c['pair_id']]
        vv.append([c['pair_id'],CHANGE[r['interaction_change']],CHANGE[r['occlusion_change']],r['brief']])
    observations=table(doc,['案例','整体交互','前后遮挡','主要观察'],vv)
    observations.autofit=False
    for i,width in enumerate((1.7,2.,2.,12.3)):
        observations.columns[i].width=Cm(width)
        for row in observations.rows:row.cells[i].width=Cm(width)
    doc.add_heading('如何使用这次结论',1)
    for text in conclusions['limits_next']:doc.add_paragraph(text)
    doc.add_paragraph('输入、完整提示词、逐人物检测及排除原因、资源记录与原尺寸图片均保存在本轮独立实验目录。'
                      '历史图片继续保留；本次同时改了绘图和提示词，新旧差异不能全部归因于线条清晰化。',style='Caption')
    for c,(png,preview) in zip(cases,panels):
        doc.add_page_break();pid=c['pair_id'];doc.add_heading(f'{pid}｜整图与两位人物的脸部',1)
        decision=manual['pairwise'][pid]
        if decision.get('assignment_text'):doc.add_paragraph('本组对应：'+decision['assignment_text'])
        doc.add_paragraph(decision['detail'])
        doc.add_picture(str(png),width=Cm(18.0))
        doc.add_paragraph('A/B列为原图；B与两张骨架保持官方人物编号。输出脸部来自实际检测框，未检出时保留空位。'
                          '裁剪放大仅便于观察，不代表原图增加了细节。',style='Caption')
    report.parent.mkdir(parents=True,exist_ok=True);doc.save(report)
    # Images must be genuinely embedded, lossless PNGs with original bytes.
    with ZipFile(report) as z:
        bad=z.testzip()
        if bad:raise ValueError(f'Bad DOCX ZIP member: {bad}')
        xml=z.read('word/document.xml').decode()
        media=[p for p in z.namelist() if p.startswith('word/media/')]
        hashes={hashlib.sha256(z.read(p)).hexdigest() for p in media}
        assert len(media)==6 and all(sha(png) in hashes for png,_ in panels)
        assert xml.count('<w:drawing>')==6
        rels=z.read('word/_rels/document.xml.rels').decode()
        assert 'TargetMode="External"' not in rels
    reopened=Document(report);assert len(reopened.inline_shapes)==6 and len(reopened.tables)==2
    write(root/'report_checks.json',{'report':str(report),'report_sha256':sha(report),'zip_valid':True,
        'embedded_png_count':6,'all_original_png_bytes_preserved':True,'external_images':False,
        'formal_results':12,'panels':[{'png':str(p),'preview':str(v),'png_sha256':sha(p)} for p,v in panels],
        'source_media':{'resolution':1024,'methods':list(METHODS),'mapping':'direct'}})
    print(report)


def main():
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('--output',type=Path,default=DEFAULT)
    p.add_argument('--report',type=Path,default=REPORT);a=p.parse_args();build(a.output,a.report)


if __name__=='__main__':main()
