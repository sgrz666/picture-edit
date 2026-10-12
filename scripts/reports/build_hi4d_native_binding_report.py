#!/usr/bin/env python3
"""Embedded-image Word report for the fixed native Hi4D binding experiment."""
from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path
import sys
from zipfile import ZipFile

PROJECT=Path(__file__).resolve().parents[2]
sys.path.insert(0,str(PROJECT))
import numpy as np
from PIL import Image,ImageDraw,ImageFont
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from scripts.eval.eval_hi4d_native_binding import DEFAULT,METHODS,MAPPINGS,sha,read_rows,write
from scripts.eval.score_hi4d_native_binding import records

FONT='/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc'
REPORT=Path('/home/shangguanrz/nihaochen/docs/Hi4D双人人物对应与骨架输入_对比结果_20261010.docx')
LABELS={'rgb':'A＋B照片','combined':'A＋合并骨架','separate':'A＋分开骨架'}
FIXED=((1024,'hi4d_00','direct'),(1024,'hi4d_00','swap'),
       (1024,'hi4d_03','direct'),(1024,'hi4d_03','swap'),
       (512,'hi4d_04','direct'),(1024,'hi4d_04','direct'))


def crop(image,box,size):
    if box is None:
        out=Image.new('RGB',(size,size),(238,238,238));d=ImageDraw.Draw(out)
        d.text((size*.27,size*.45),'未检出',font=ImageFont.truetype(FONT,max(18,round(size*.08))),fill=(90,90,90))
        return out
    x0,y0,x1,y1=box;cx,cy=(x0+x1)/2,(y0+y1)/2
    width=max(x1-x0,y1-y0)*1.7
    rect=(round(cx-width/2),round(cy-width/2),round(cx+width/2),round(cy+width/2))
    return image.crop(rect).resize((size,size),Image.Resampling.LANCZOS)


def panel(root,res,pid,mapping,rows,cases,manual):
    case=cases[pid];jobs={r['method']:r for r in rows if r['resolution']==res and r['pair_id']==pid and r['mapping']==mapping}
    paths=[case['source']['rgb'],case['target']['rgb'],*[jobs[m]['image'] for m in METHODS]]
    side=res;header=round(side*.11);gap=8;left=round(side*.22)
    canvas=Image.new('RGB',(left+5*(side+gap),header+3*(side+gap)),(255,255,255))
    draw=ImageDraw.Draw(canvas);font=ImageFont.truetype(FONT,round(side*.08));small=ImageFont.truetype(FONT,round(side*.065))
    for col,(path,label) in enumerate(zip(paths,('A 原图','B 动作参考',*[LABELS[m] for m in METHODS]))):
        im=Image.open(path).convert('RGB');x=left+col*(side+gap)
        draw.text((x+8,8),label,font=font,fill=(20,20,20))
        # Full-scene images retain the exact same letterbox as native inference.
        if col<2:
            from src.data.letterbox import Letterbox
            full=Letterbox.from_size(im.size,side).image(im,fill=(128,128,128))
        else:full=im
        canvas.paste(full,(x,header))
        rec=records(root,path);review=manual['images'][path]
        for source_id in (0,1):
            identity=MAPPINGS[mapping][source_id] if col==1 else source_id
            pp=next(p for p in review['people'] if p['source_id']==identity)
            face_index=pp.get('face_detection_index')
            f=rec['faces']['detections'][face_index] if face_index is not None else None
            box=f['box_native'] if f and pp.get('face_real') else None
            face=crop(im,box,side);canvas.paste(face,(x,header+(source_id+1)*(side+gap)))
            status=pp.get('identity_status')
            message=('源人物缺失' if col>=2 and status=='replaced' else
                     '身份混合' if col>=2 and status in ('mixed','uncertain') else
                     '小脸参考' if f is not None and f['short_side_at_512']<32 else None)
            if message:
                draw.rectangle((x,header+(source_id+1)*(side+gap),x+side,header+(source_id+1)*(side+gap)+round(side*.09)),fill='white')
                draw.text((x+6,header+(source_id+1)*(side+gap)+2),message,font=small,fill=(100,70,30))
    for row,label in enumerate(('整图','A1脸','A2脸')):
        draw.text((3,header+row*(side+gap)+12),label,font=small,fill='black')
    out=root/'panels'/f'{pid}_{res}_{mapping}.png';out.parent.mkdir(parents=True,exist_ok=True);canvas.save(out)
    return out


def configure(doc):
    section=doc.sections[0];section.page_width=Cm(21);section.page_height=Cm(29.7)
    section.top_margin=section.bottom_margin=Cm(1.4);section.left_margin=section.right_margin=Cm(1.4)
    for name in ('Normal','Body Text'):
        st=doc.styles[name];st.font.name='Arial';st.font.size=Pt(10.5)
        st.element.rPr.rFonts.set(qn('w:eastAsia'),'等线')
        st.paragraph_format.space_after=Pt(5)
    for name in ('Title','Heading 1','Heading 2'):
        doc.styles[name].element.rPr.rFonts.set(qn('w:eastAsia'),'等线')
    node=OxmlElement('w:doNotAutoCompressPictures');doc.settings.element.append(node)
    footer=section.footer.paragraphs[0];footer.alignment=2
    footer.add_run('Hi4D 双人编辑 · 原生 Qwen 条件核对')


def table(doc,headers,rows,best=None):
    t=doc.add_table(rows=1,cols=len(headers));t.style='Table Grid'
    for c,h in zip(t.rows[0].cells,headers):c.text=h
    for i,values in enumerate(rows):
        cells=t.add_row().cells
        for j,(cell,value) in enumerate(zip(cells,values)):
            run=cell.paragraphs[0].add_run(str(value))
            if best and (i,j) in best:run.bold=True
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(3)
                for run in p.runs:run.font.size=Pt(9)
    return t


def number(x,digits=3):return '无法判断' if x is None else f'{x:.{digits}f}'


def build(root,report):
    rows=read_rows(root/'generations.jsonl');cases={c['pair_id']:c for c in json.loads((root/'data/cases.json').read_text())}
    manual=json.loads((root/'review/manual.json').read_text());summaries=json.loads((root/'evaluation/summary.json').read_text())
    conclusions=json.loads((root/'review/conclusions.json').read_text())
    panels={}
    for res,pid,mapping in sorted({(r['resolution'],r['pair_id'],r['mapping']) for r in rows}):
        panels[(res,pid,mapping)]=panel(root,res,pid,mapping,rows,cases,manual)
    doc=Document();configure(doc)
    doc.add_heading('Hi4D 双人人物对应与骨架输入',0)
    doc.add_paragraph('原生 Qwen 对比结果｜6组跨人物案例｜48张生成结果',style='Subtitle')
    doc.add_heading('先看结论',1)
    for text in conclusions['main']:doc.add_paragraph(text,style='List Bullet')
    doc.add_heading('这次实际做了什么',1)
    doc.add_paragraph('A 图提供要保留的两个人、服装和背景，B 图提供两人的目标动作及交互。'
      '比较“直接给 B 照片”“给两人合在一张图中的骨架”“分别给两张骨架”。骨架经过重绘，保留身体、手和面部点。'
      '两张独立骨架沿用同一画布，合起来的位置与合并骨架相同。')
    doc.add_paragraph('每组同时要求两种人物分配：A1→S1、A2→S2，以及A1→S2、A2→S1。'
      '人物编号由原始数据的人物标注确定，不等于左右位置。通过反转指令，检查模型是否真的能按文字指定“谁做哪个动作”。'
      '三条路线共享详细动作、人物保留和防串脸要求，只调整输入图片的角色说明。')
    doc.add_paragraph('六组案例由三张A原图分别搭配两种目标动作组成，共六位源人物。重复使用A是为了检查同一人物是否能响应不同动作；这六组并非六组独立人物。')
    doc.add_paragraph('本轮没有训练；没有使用项目自训的 Body／Face Adapter，也没有加载 LoRA。'
      '所有结果固定30步生成、种子42。1024是完整六组比较，512只复核预定的两组；没有使用最终测试集。')
    doc.add_heading('1024：六组案例 × 两种分配',1)
    rr=[r for r in summaries if r['resolution']==1024]
    values=[[LABELS[r['method']],f'{r["correspondence_correct"]}/{r["images"]}',
        f'{r["cases_both_assignments_correct"]}/6',f'{r["artifact_images"]}/{r["images"]}',
        f'{r["seconds"]:.1f}秒'] for r in rr]
    best={(i,1) for i,r in enumerate(rr) if r['correspondence_correct']==max(x['correspondence_correct'] for x in rr)}
    best|={(i,3) for i,r in enumerate(rr) if r['artifact_images']==min(x['artifact_images'] for x in rr)}
    best|={(i,4) for i,r in enumerate(rr) if r['seconds']==min(x['seconds'] for x in rr)}
    best|={(i,2) for i,r in enumerate(rr) if r['cases_both_assignments_correct']==max(x['cases_both_assignments_correct'] for x in rr)}
    table(doc,['输入路线','对应正确↑','双向均正确↑','异常图片↓','单图耗时↓'],values,best)
    doc.add_paragraph('“对应正确”按服装、发型等可辨认的源人物外观检查两位A是否处于指定角色；它不代表面部身份已严格保留，也不代表每个关节都已精确匹配。'
      '“异常”统计线条残留、多余人物、畸形及局部服装污染；整个人沿用B的情况已计入“对应错误”，不重复算作这一栏的可见异常。完整原因保存在逐图记录中。',style='Caption')
    doc.add_paragraph('“双向均正确”要求同一组案例的两种相反分配都执行成功，比只做一种分配更能检查文字指定是否有效。分开骨架的12张中，有4张同时对应正确且没有上述已记录异常，仍不等于所有细节都完美。',style='Caption')
    doc.add_page_break()
    doc.add_heading('如何读动作与身份指标',1)
    doc.add_paragraph('先确认生成图里谁是A1、谁是A2，再按指令检查目标骨架。'
      '动作分数只在三条路线均能辨认、均有有效检测的同一批人物上比较；没有为了提高分数而按位置重新配人。'
      '近距离拥抱可能让检测器把两人合成一个框，这时缺失一方会明确记录。')
    for res in (1024,512):
        doc.add_heading(f'{res} 指标对照',2);rr=[r for r in summaries if r['resolution']==res]
        vv=[[LABELS[r['method']],number(r['pose_pck_002_common']),number(r['pose_nme_common']),
             str(r['pose_common_n']),number(r['background_mae'],2)] for r in rr]
        best=set()
        for j,k,up in ((1,'pose_pck_002_common',True),(2,'pose_nme_common',False),(4,'background_mae',False)):
            available=[r[k] for r in rr if r[k] is not None]
            if available:
                target=(max if up else min)(available)
                best|={(i,j) for i,r in enumerate(rr) if r[k]==target}
        table(doc,['输入路线','动作 PCK↑','动作 NME↓','共同人物数 n','背景误差↓'],vv,best)
    doc.add_paragraph('PCK越高，指定关节越接近目标位置；NME越低，关节位置误差越小。n为可共同比较的“人物×案例×分配”次数，并非独立人物人数。'
      '本次共同动作子集很小（1024为4次、512为1次），只作局部诊断，不能代替全案例的人物对应检查。'
      '背景误差只计算两图人物区域之外的有效画面，排除补边；人物体型差异和移动后的背景显露仍会影响该指标。'
      'B中的人并非理想输出，因此不计算与B的PSNR或SSIM。',style='Caption')
    doc.add_heading('面部与分辨率',2)
    for text in conclusions['face_resolution']:doc.add_paragraph(text)
    doc.add_heading('下一步建议',2)
    for text in conclusions['next']:doc.add_paragraph(text,style='List Bullet')
    doc.add_paragraph('范围：这次验证的是原生模型对人物分配和重绘骨架的理解，尚未执行体型适配、接触优化等上游重定向。'
      '结果限于六组固定案例及一个随机种子，不能据此宣称已实现可靠的任意双人身份对齐。')
    doc.add_page_break();doc.add_heading('固定案例对照',1)
    doc.add_paragraph('每图第一行是整图，第二、三行分别观察A1、A2的脸。B列按当前分配展示对应目标角色的脸；'
        '结果列只裁实际检测且确认的人脸，不按B的位置强行裁脸。缺失者的替代脸、混合身份另行标注。'
        '文中对照图按版面缩放，脸部始终从原图裁取；原尺寸单图和全部16组对照另行保存。')
    embedded=[]
    for i,key in enumerate(FIXED):
        if i and i%2==0:doc.add_page_break()
        res,pid,mapping=key
        assignment='A1→S1，A2→S2' if mapping=='direct' else 'A1→S2，A2→S1'
        doc.add_heading(f'{pid} · {res} · {assignment}',2)
        note=conclusions.get('figure_notes',{}).get(f'{pid}_{res}_{mapping}','')
        if note:doc.add_paragraph(note)
        path=panels[key]
        # 768 pixels per column remains >500 dpi at this page width. Keep the
        # unscaled panel and original outputs separately; embed this PNG intact.
        media=root/'report_media'/path.name;media.parent.mkdir(exist_ok=True)
        with Image.open(path) as im:
            width=round(im.width*min(1,768/res));height=round(im.height*min(1,768/res))
            im.resize((width,height),Image.Resampling.LANCZOS).save(media)
        doc.add_picture(str(media),width=Cm(18.1));embedded.append(media)
    report.parent.mkdir(parents=True,exist_ok=True);doc.save(report)
    with ZipFile(report) as z:
        if z.testzip() is not None:raise AssertionError('Broken docx ZIP')
        names=[n for n in z.namelist() if n.startswith('word/media/')]
        hashes={hashlib.sha256(z.read(n)).hexdigest() for n in names}
        if not all(sha(p) in hashes for p in embedded):raise AssertionError('Images not embedded losslessly')
        if b'Target="file:' in z.read('word/_rels/document.xml.rels'):raise AssertionError('External image link')
    write(root/'report_validation.json',{'docx':str(report),'sha256':sha(report),'formal_images':len(rows),
          'panels':len(panels),'embedded_panels':len(embedded),'zip_valid':True,'images_losslessly_embedded':True})
    print(str(report))


if __name__=='__main__':
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('--output',type=Path,default=DEFAULT);p.add_argument('--report',type=Path,default=REPORT)
    args=p.parse_args();build(args.output,args.report)
