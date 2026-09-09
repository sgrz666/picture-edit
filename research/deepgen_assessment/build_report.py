from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parent
OUT = ROOT.parent / 'DeepGen人物姿态图像编辑适配性评估.docx'
d = Document()
s = d.sections[0]
s.page_width=Inches(8.5); s.page_height=Inches(11)
s.top_margin=Inches(.72); s.bottom_margin=Inches(.68)
s.left_margin=Inches(.78); s.right_margin=Inches(.78)
s.header_distance=Inches(.28); s.footer_distance=Inches(.3)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
    st=d.styles[name]; st.font.name='Calibri'; st.font.color.rgb=RGBColor(0,0,0)
    st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'宋体' if name=='Normal' else '微软雅黑')
n=d.styles['Normal']; n.font.size=Pt(11)
n.paragraph_format.line_spacing=1.15; n.paragraph_format.space_after=Pt(6)
for st in d.styles:
    for el in list(st.element.iter(qn('w:pBdr'))): el.getparent().remove(el)
for name in ['Title','Subtitle','Caption']:
    d.styles[name].font.italic=False
    d.styles[name].font.underline=False
for name,size in [('Title',23),('Subtitle',11),('Heading 1',16),('Heading 2',12)]:
    st=d.styles[name];st.font.size=Pt(size)
    st.paragraph_format.space_before=Pt(9);st.paragraph_format.space_after=Pt(9)
    if name.startswith('Heading'): st.font.bold=True
head=s.header.paragraphs[0];head.text='DeepGen 人物姿态图像编辑适配性评估';head.style=d.styles['Caption']
head.runs[0].font.size=Pt(9)
f=s.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=f.add_run('技术评估  ·  2026年9月8日    ');r.font.size=Pt(9)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');f._p.append(field)

def p(t,bold=False):
    x=d.add_paragraph();r=x.add_run(t);r.bold=bold;return x
def h(t): d.add_heading(t,level=2)
def page(t):
    x=d.add_heading(t,level=1);x.paragraph_format.page_break_before=True
def table(headers,rows,widths):
    t=d.add_table(rows=1, cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    pr=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
    pr.append(borders)
    for c,w in zip(t.columns,widths):c.width=Inches(w)
    for vals,row in [(headers,t.rows[0])]+[(v,t.add_row()) for v in rows]:
        ishead=row is t.rows[0] # proxies are not identity stable
        for i,(cell,text) in enumerate(zip(row.cells,vals)):
            cell.width=Inches(widths[i]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text=str(text)
            cp=cell._tc.get_or_add_tcPr(); margins=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),'90');e.set(qn('w:type'),'dxa');margins.append(e)
            cp.append(margins)
            for para in cell.paragraphs:
                para.paragraph_format.line_spacing=1.12;para.paragraph_format.space_after=Pt(1)
                for r in para.runs:r.font.size=Pt(10)
        cant=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(cant)
    for i,row in enumerate(t.rows):
        for cell in row.cells:
            sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E9EDF1' if i==0 else ('F8F9FA' if i%2==0 else 'FFFFFF'));cell._tc.get_or_add_tcPr().append(sh)
            if i==0:
                for para in cell.paragraphs:
                    for r in para.runs:r.bold=True
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    gap=d.add_paragraph();gap.paragraph_format.space_after=Pt(0);gap.paragraph_format.space_before=Pt(0);gap.paragraph_format.line_spacing=Pt(3);gap.add_run().font.size=Pt(3)
    return t

REFS=[
('DeepGen 官方项目与基准结果','https://github.com/deepgenteam/deepgen'),
('DeepGen 技术报告  arXiv 2602.12205','https://arxiv.org/abs/2602.12205'),
('DeepGen Diffusers 权重与配置  固定版本','https://huggingface.co/deepgenteam/DeepGen-1.0-diffusers/tree/85c2ed257b9406d8531965c68e140dc7fdb81382'),
('DeepGen Diffusers 推理与动态 Transformer 实现','https://huggingface.co/deepgenteam/DeepGen-1.0-diffusers/blob/85c2ed257b9406d8531965c68e140dc7fdb81382/deepgen_pipeline.py'),
('DeepGen 动态 Transformer 源码','https://github.com/deepgenteam/deepgen/blob/fbaaaf4ae1c79450d6a34b30beb4f59840568c69/src/models/sd3_kontext/transformer_sd3_dynamic.py'),
('DeepGen 原生 SFT 配置','https://github.com/deepgenteam/deepgen/blob/fbaaaf4ae1c79450d6a34b30beb4f59840568c69/configs/finetune/deepgen_joint_sft_scb.py'),
('DeepGen 训练说明','https://github.com/deepgenteam/deepgen/blob/main/TRAIN.md'),
('CHI3D 数据集与接触标注说明','https://ci3d.imar.ro/chi3d'),
('EgoBody 数据集与第三人称视图说明','https://github.com/sanweiliti/EgoBody'),
('Hi4D 数据格式说明','https://github.com/yifeiyin04/Hi4D/blob/main/dataset.md'),
('BEDLAM 官方数据集','https://bedlam.is.tue.mpg.de/'),
('CFLD 姿态引导人物图像合成','https://github.com/YanzuoLu/CFLD'),
('T2I Adapter 官方实现','https://github.com/TencentARC/T2I-Adapter'),
('PyTorch 自动微分机制','https://docs.pytorch.org/docs/main/notes/autograd.html'),
]
def cite(*ids):
    para=d.add_paragraph();para.paragraph_format.space_after=Pt(4)
    for i in ids:
        title,url=REFS[i-1];link=OxmlElement('w:hyperlink');link.set(qn('r:id'),para.part.relate_to(url,RT.HYPERLINK,is_external=True))
        r=OxmlElement('w:r');pr=OxmlElement('w:rPr');sz=OxmlElement('w:sz');sz.set(qn('w:val'),'18');pr.append(sz)
        col=OxmlElement('w:color');col.set(qn('w:val'),'34536E');pr.append(col);r.append(pr)
        tx=OxmlElement('w:t');tx.set(qn('xml:space'),'preserve');tx.text=f'[{i}] {title}  ';r.append(tx);link.append(r);para._p.append(link)

d.add_paragraph('DeepGen用于人物姿态图像编辑的适配性评估',style='Title')
d.add_paragraph('面向三维人体控制与双人交互编辑的技术评估报告',style='Subtitle')
p('评估对象  DeepGen 1.0    ｜    项目范围  研究目标三    ｜    资料截至  2026年9月8日')
h('评估结论')
p('建议将 DeepGen 纳入第三阶段的优先原型基座，先完成小规模适配验证，再决定是否承担正式实验。它已有参考图像编辑路径，生成主干规模低于6B，并提供可利用的控制残差接口；这些条件能够减少从头构建编辑模型的工作。单卡84GB适合启动冻结主干与轻量外挂的训练验证，但现阶段没有本项目的训练峰值显存或效果实测。',True)
p('DeepGen 的主要风险是双人接触和外观保持能力尚未得到专项验证。公开通用编辑成绩也不足以支持其优于 Qwen-Edit 的结论。外挂能够增强目标姿态约束，但不能预先保证完全弥补基座在手指结构、遮挡重建和服饰细节上的能力差距。')
table(['评估维度','当前判断','对后续实验的含义'],[
('生成主干规模','符合当前口径','DiT 权重张量约2.47B；全组件存储口径约6.74B。'),
('原生编辑与外挂接口','具备工程起点','保留参考图像路径，新增结构编码与残差注入分支。'),
('双人交互与手部效果','需要专项验证','先检查身份交换、接触偏移和手部结构错误。'),
('单卡84GB训练','值得直接验证','从512分辨率和小批次启动，再评估1024。'),
('Qwen Edit 类适配要求','存在验收边界','使用 Qwen 编码器不等于属于 Qwen Image Edit 模型家族。')],[1.35,1.4,4.19])
h('评估依据与适用范围')
p('本报告依据项目任务书、已提供的两份讨论材料、官方论文和公开代码编写。公开事实、权重元数据计算与拟议实验分别陈述。本文未运行 DeepGen 完整权重推理或训练，不将工程推断表述为性能实测。对于蒸馏路线，质量应按具体检查点评价，不能仅凭“蒸馏”判断优劣。')
p('任务书要求完成与 Qwen-Edit 类模型的适配，并以 Qwen-Edit 为主观评测基线。若验收严格限定 Qwen-Image-Edit 家族，DeepGen 不能替代该项交付；若允许同类图像编辑模型，则可作为技术候选。前一项约束应在正式选型时落实。')

page('1 模型结构与项目接口')
p('DeepGen 以 Qwen2.5-VL 为视觉语言理解模块，通过 Stacked Channel Bridging（SCB，分层特征连接模块）向扩散 Transformer（DiT）提供语义条件；生成模块来源于 UniPic2-SD3.5M-Kontext。它具备图像编辑训练基础，因此本项目可以围绕已有参考图像条件增加几何控制。它并非 Qwen-Image-Edit 的小参数版本。')
cite(1,2)
table(['组件','权重张量元素数','在本项目中的处理'],[
('生成 DiT','2,469,663,936','冻结，接收外挂输出的控制残差。'),
('Qwen 视觉语言模块','3,754,622,976','冻结，保留发布版本的已合并权重。'),
('SCB 连接模块','432,449,532','冻结，保留原有语义到生成条件的映射。'),
('VAE 图像编解码器','83,819,683','冻结，编码源图和训练目标图。'),
('合计','6,740,556,127','全组件统计，不作为主干6B限制的口径。')],[1.7,1.55,3.69])
p('上表为已固定版本 safetensors 文件头中张量形状的乘积汇总，含存储层面的张量统计；共享参数与缓冲区仍应在模型加载后复核。其用途是核查体量和权重占用，不能替代运行时峰值测量。官方“3B＋2B”的简写也不应直接用于精确显存预算。')
cite(3)
h('与第二阶段的衔接')
p('第三阶段的输入应明确包含源图、两人的稳定身份编号、重定向后的 SMPL-X、目标相机、接触部位对应关系及有效性标记。若希望保留目标角色外观，源图中的角色编号必须与第二阶段输出一致。仅有两组人体参数而缺少相机或身份对应，无法稳定构造图像控制信号。')
p('建议将网格投影成深度、法线、身体部位与人物实例图，并生成关节和接触区域提示。每个人的可见性、左右手和深度顺序需要保留。SMPL-X 描述人体表面，不描述服装外轮廓，训练时应允许衣物偏离裸身轮廓，避免把宽松服装压成身体形状。')
h('保留原有外观路径')
p('源图同时承担视觉语义条件与参考潜变量条件。后者参与生成 Transformer 的图像序列计算，是保留人物和服饰的现有基础。优先复用这条路径；只有身份和纹理的消融结果显示不足时，再增加按人物区分的外观特征分支。')
cite(4)

page('2 外挂控制的适配方案')
h('建议的首版结构')
p('首版采用冻结 DeepGen 与可训练结构适配器。适配器融合目标深度、法线、人物编号、部位和接触提示，结合扩散时间步生成若干层的控制残差，注入 DiT。建议把100M至300M作为初始可训练参数预算，并通过实际参数统计约束规模；这个范围是工程方案，尚无充分证据说明它一定满足双人精细控制。')
p('不同控制图优先在通道或紧凑特征层融合。接触关系还应编码为“人物A的左手—人物B的右肩”等部位对应，辅以局部接触热图。单张二值接触图无法完整表达参与者、左右侧和遮挡顺序。控制分支输出层采用零初始化或零门控，使训练初始行为接近原始编辑模型。')
h('已存在的接口与必须修改的部分')
p('已核查的动态 Transformer 接受 block_controlnet_hidden_states，并在部分 Transformer 块之后加到隐藏状态；同时提供梯度检查点路径。由此可以避免重新设计主干内部的全部控制机制，但仍需实现适配器、数据管线和训练入口。')
cite(4,5)
table(['接口问题','适配要求'],[
('目标与参考图像序列拼接','残差应匹配目标 token、参考 token 与批内填充后的完整形状。首版可仅在目标区域输出非零残差。'),
('Transformer 的特征组织','按实际层宽和序列长度投影；不能直接沿用 U-Net 多尺度特征的尺寸假设。'),
('控制强度与条件丢弃','为不同几何信号设置消融开关；检查推理时 CFG 分支与训练时条件丢弃的一致性。'),
('多视图和多控制图','同一目标相机渲染并同步裁剪。不能默认把每张控制图当作额外参考图直接拼入。')],[1.85,5.09])
h('可以复用的研究工作')
p('T2I-Adapter 可提供轻量条件编码和冻结生成器训练的设计参考；CFLD 可提供人物外观与姿态分离的实验思路及数据准备经验。它们的预训练外挂与 DeepGen 的层结构、序列组织和条件维度并不一致，现成权重不能据此认定可直接加载。复用重点是数据处理、控制设计和评价协议。')
cite(12,13)
p('研究增量宜集中在双人身份绑定、接触部位对应和遮挡感知控制，以及在主干完全冻结时的有效注入方式。是否构成方法创新，需要结合后续消融与相关工作比较，单纯连接 SMPL-X 渲染器和现有控制接口不足以证明贡献。')

page('3 效果证据与能力边界')
p('官方发布的编辑基准可以证明 DeepGen 具备通用编辑能力，但不能直接回答其对拥抱、握手等任务的适配效果。下表为作者报告的分数，并非本项目复测；各列评价体系不同，不能横向相加。')
table(['模型','GEdit EN','ImgEdit','RISE','UniREdit'],[
('DeepGen SFT','7.12','4.09','13.3','77.5'),
('DeepGen RL','7.17','4.14','10.8','75.7'),
('Qwen Image Edit 2509','7.54','4.35','8.9','56.5')],[2.7,1.06,1.06,1.06,1.06])
p('在上述结果中，DeepGen 的通用编辑分数低于 Qwen-Image-Edit 2509，推理编辑分数较高。这个差异支持其具备一定指令理解能力，却不能证明几何精度、身份保持或手部质量更好。RL 版本也并未在每项编辑指标上超过 SFT，后续应核对检查点来源并进行任务内比较。')
cite(1)
table(['项目要求','DeepGen 可以利用的基础','最需要验证的失败情形'],[
('单人姿态变化','原生图像编辑与几何残差','原姿态残留，或姿态变对但服装纹理改变。'),
('双人身份保持','源图参考潜变量','两人换脸、衣服串色、肢体归属错误。'),
('拥抱与搂肩','目标深度与实例控制','前后关系错误，手臂消失或多出肢体。'),
('握手与牵手','手部关节及接触区域提示','手指粘连，接触位置偏移，左右手混淆。'),
('不同体型的重定向','目标人体形状与投影','回退到源人体比例，或破坏衣物形状。')],[1.25,2.15,3.54])
h('对效果的合理预期')
p('小幅、可见关节充分的单人动作适合作为接通控制的初始任务；双人整体构图与躯干方向是下一步验证对象。大角度转身、严重遮挡、交叉手指等情况需要重建源图中不存在的外观信息，应单独统计，不能以平均图像分数掩盖其失败。以上顺序是实验设计判断，不是模型已有能力分级。')
p('主干冻结保留了基座参数，但不自动保证输出身份不变。控制残差过强仍可能覆盖外观条件。若提高控制强度持续导致身份或纹理损失，应优先检查条件冲突、数据对应和人物绑定，再考虑增加外挂容量。')

page('4 数据复用与训练样本构建')
p('现有数据能够减少采集和标注成本，但没有一个已核查的数据源能够直接覆盖“任意目标体型重定向后的双人真实编辑图像”。建议先利用同一人物或同一对人物的视频帧构造真实源图与目标图，再补充三维几何监督。')
table(['数据源','适合的用途','使用边界'],[
('CHI3D [8]','双人接触控制与验证','有双方 SMPL-X；第二人的网格属于伪真值。接触签名不是逐帧完整标注。'),
('EgoBody [9]','真实双人外观与姿态配对','选择第三人称 Kinect 视图中双方可见的帧；不能把第一人称图像当作双人完整源图。'),
('Hi4D [10]','近距离双人遮挡与交互','发布的人体参数为 SMPL，不能直接当作带完整手部标注的 SMPL-X。'),
('BEDLAM [11]','几何控制预训练与体型覆盖','合成图像存在域差异；筛选稳定相机和合适人物实例。'),
('DeepFashion 配对 [12]','单人姿态与服饰保持预实验','不能直接承担双人接触监督；三维参数需要另行估计。')],[1.25,2.05,3.64])
p('CHI3D 的631段交互记录各包含一个建立接触的标注时间点，四个视角对应2524组图像与接触签名。可将这些标注点用于接触验证；对其余帧，从网格距离生成的接触标签应标为派生标签，并记录阈值和可信度，不能等同于人工标注。')
cite(8,9,10,11,12)
h('推荐的数据单元')
p('每个训练单元保存源图、真实目标图、目标 SMPL-X 与相机、两人身份对应、控制渲染图、接触标签来源及有效区域。源图与目标图优先来自同一人物组合、同一场景和同一相机。相机移动或人物出画的样本需要过滤或显式建模，否则背景变化会混入姿态学习目标。')
p('在有真实目标图的阶段，使用该图对应的目标几何监督编辑。第二阶段任意重定向产生的体型和动作组合通常没有真实目标照片，应作为泛化测试，或通过可控合成数据补充，不能将不对应的原视频帧当作目标图。')
h('数据划分和准入')
p('按人物组合与视频序列划分训练和测试，保留不同身高差、接触动作和遮挡程度。连续帧不应随机分散到两侧。下载前记录各数据集、人体模型和权重的许可及项目用途范围；可公开获取不等于可以用于任意企业项目训练。')

page('5 冻结主干的训练实现')
h('先建立可验证的冻结边界')
p('冻结 DiT、视觉语言模块、SCB 和 VAE，只把新增适配器参数交给优化器。发布模型中已经合并的 VLM LoRA 属于所选基座权重，应随基座冻结；不要替换成未合并的原始 Qwen 编码器而仍假定模型等价。实际加载代码优先读取权重目录内的合并 VLM。')
cite(4)
p('原生 SFT 配置包含 freeze_transformer=False 和 LoRA 设置，其目的并非本项目的外挂训练。因此不能直接启动官方 SFT 配置后声称主干保持不变。需要独立配置冻结边界、优化器参数组、适配器保存格式以及推理接入方式。')
cite(6,7)
h('冻结权重仍需要反向传播')
p('当适配器残差进入 DiT 后，损失到适配器之间的计算图必须保留。将 DiT 全部包在 no_grad 中会切断这条路径。正确做法是冻结参数的 requires_grad，同时允许对输入残差求梯度；仅对无需回传的源图编码等路径关闭梯度。公开推理 pipeline 使用 no_grad，训练必须调用可微的底层模块。')
cite(4,14)
h('训练目标与推进顺序')
p('首轮沿用原模型的流匹配或去噪训练定义，核对噪声插值、时间步采样、预测目标和潜变量缩放。一个训练样本随机选择一个时间步即可建立基础监督，初期不需要把整条多步采样过程展开反向传播。')
p('先用单人配对图验证结构控制，再加入双人整体姿态，最后提高接触和手部样本比例。对手部或接触区域增加损失权重时，应区分真实标注与伪标签；潜空间加权误差只是局部训练代理，不能作为手部解剖结构正确的保证。')
p('身份、衣物和背景损失应在对应区域与有效可见区域内定义。姿态变化后直接逐像素约束整个人体会惩罚合理变化。需要图像空间辅助监督时，可对适当噪声区间的预测图进行低频解码评估，并计入其额外显存与计算成本。')
h('验明训练确实有效')
p('训练前后核对基座参数校验值；完成一次反向后确认适配器有有效梯度而基座没有参数梯度；关闭适配器应恢复原始编辑行为。每个检查点同时保存基座版本、控制图规范和适配器配置，以便在另一台机器上重现。')

page('6 训练显存预算与84GB设备适配')
p('建议按单卡84GB资源推进首轮实验，而非先追加硬件。该判断基于约2.47B的冻结 DiT、可缓存的编码路径和100M至300M适配器设定。它说明资源条件值得开展实测，不代表已证明1024分辨率的任意配置都能运行。')
h('可以计算的静态占用')
table(['项目','计算假设','估算占用'],[
('全套冻结权重','所有组件统一 BF16，每元素2字节','12.56 GiB'),
('仅 DiT 驻留','编码结果已缓存且相应模块移出 GPU','4.60 GiB'),
('100M适配器训练状态','保守按每参数16字节预算','1.49 GiB'),
('300M适配器训练状态','同上，含梯度与优化器等状态','4.47 GiB'),
('全权重加300M适配器','以上静态项求和','17.03 GiB'),
('DiT加300M适配器','缓存条件下的静态项求和','9.07 GiB')],[2.05,3.55,1.34])
p('GiB按2³⁰字节计算。16字节是用于预算的保守混合精度 Adam 状态口径，实际取决于优化器实现和主权重副本。发布连接模块以FP32存储；若加载后仍保留FP32，全权重静态项比上表增加约0.81 GiB。文件体积、权重占用和训练峰值必须分别报告。')
p('训练峰值还包含 DiT 与适配器激活、注意力临时张量、图像编码、图像空间辅助损失、内存分配器预留和运行时开销。现有核查无法给这些动态项一个可靠的单值，因此不把17.03 GiB表述为训练只需17GB。')
h('分辨率为什么影响明显')
p('按 VAE 八倍下采样、DiT 两倍分块计算，512方图约1024个图像 token，1024方图约4096个。若同时使用同分辨率源图与目标图，图像序列分别约2048和8192个 token，另有语义条件。高效注意力减少显式注意力矩阵的存储，但序列加长仍显著增加计算与激活成本。')
cite(3,4)
h('建议的设备使用顺序')
p('先测512×512、微批次1、BF16、梯度检查点和高效注意力；通过后测试768和1024。梯度累积用于增加有效批次。将冻结且确定性的 VLM、SCB 和 VAE 编码离线缓存，可进一步腾出显存，但缓存必须与提示词、裁剪、分辨率和模型版本一致。')
p('以48GB作为512原型的可尝试档位，以80至84GB作为1024原型的优先资源档位；二者均为规划建议，不是最低显存实测。最终以设备实际可用容量为准，连续训练稳定后留出约10%至15%余量。')

page('7 分阶段验证与正式实验设计')
p('推荐用四步验证决定是否继续投入。下表样本量为启动建议，不代表达到统计显著性所需的最终规模。前两步通过后，再扩大双人数据和训练时长。')
table(['阶段','建议设置','应取得的证据'],[
('原生能力复现','20例单人、40例双人；固定模型、提示和随机种子','确定原始外观保持水平，记录接触与手部失败类型。'),
('外挂连通验证','16至32组配对图过拟合；512、微批次1','主干校验值不变；梯度有效；更换目标控制能产生对应变化。'),
('小规模有效性','500至2000组去重训练对；按人物与序列留出验证集','相同基座下，姿态与交互改善，同时身份和服饰没有不可接受退化。'),
('正式对照','确定数据与协议后扩大样本；多随机种子','与消融及 Qwen-Edit 基线比较，报告置信区间与失败分布。')],[1.2,2.65,3.09])
h('必须保留的对照')
p('同一 DeepGen 基座下比较原生编辑、仅二维姿态控制、三维几何控制、加入接触关系的完整方案。分别移除人物编号、深度或接触条件，可以判断增益来自哪类信号。保持训练数据和主要计算设置一致，单独报告可训练参数量。')
p('Qwen-Edit 对照需固定具体版本、输入信息、提示词、输出分辨率、采样预算及选择样本的规则。原生基线与额外几何条件方案的比较属于完整系统比较，不应表述为同等输入下的纯基座优劣。若基线可以接收目标渲染图，应另列信息尽量匹配的对照。')
h('评价应对应任务书')
p('姿态采用可见关节的误差或正确率；身份按人物A和B分别匹配评价；服装和背景在有效区域评价。严重遮挡下的姿态估计与接触重建会产生评估器误差，应将其作为代理指标，并保留人工核验。整体图像相似度不能代替接触正确性。')
p('主观盲评分别记录身份可辨识、动作准确、交互自然及手部结构。任务书的“相比 Qwen-Edit 基线高于5%”应明确分母、平局处理及相对提升或百分点口径，不能自行改写为固定55%胜率。统计应按案例或人物组合聚类，避免把同一案例多个随机种子视作完全独立样本。')
h('显存测量与停止条件')
p('预热并实际执行优化器更新后，连续运行至少100步，记录 max_memory_allocated、max_memory_reserved、设备侧占用、每步时间和精度设置；另测验证采样峰值。若控制失效，先排查坐标、残差形状和梯度；若几何已对齐却持续出现换脸或手部错误，再决定扩展外观分支或更换基座。')

page('8 选型建议与依据索引')
p('技术选型上，建议以 DeepGen 完成首轮可行性实验。它使主要工作集中在几何条件编码、双人接触控制和评价，而无需重新训练一个完整图像编辑模型。84GB设备足以支持有价值的验证起点；训练规模应由第7节的峰值测量与效果结果逐步确定。')
p('正式实验主基座的确定应同时满足三个条件：冻结边界与控制接口已验证；在留出人物上的结构改善没有明显损害外观；项目方确认其满足模型适配范围。若严格要求 Qwen-Image-Edit 家族，DeepGen仍可用于方法预研，但不能宣称已经完成该家族的适配交付。')
h('版本与可复核材料')
p('权重核查版本为 deepgenteam/DeepGen-1.0-diffusers，revision 85c2ed257b9406d8531965c68e140dc7fdb81382。代码核查使用 DeepGen GitHub revision fbaaaf4ae1c79450d6a34b30beb4f59840568c69。Diffusers 发行版与论文 SFT 或 RL 检查点的具体映射应在下载时再次核对，不能仅按仓库名称推定。')
p('项目依据为《人物图像姿态编辑任务书_v3.docx》及两份已提供的讨论材料。参数统计方法与结果已保存在本报告同目录的 audit_checkpoint_headers.py 和 checkpoint_audit.json。当前证据包括公开资料与元数据核查，不包括目标设备训练实测。')
h('公开依据')
for i in range(1,len(REFS)+1):cite(i)

props=d.core_properties
props.title='DeepGen用于人物姿态图像编辑的适配性评估'
props.subject='华为项目研究目标三的基座适配与训练资源评估'
props.author='项目技术评估'
props.keywords='DeepGen, SMPL-X, 图像编辑, 双人交互, 外挂训练, 显存'
d.save(OUT)
print(str(OUT))
