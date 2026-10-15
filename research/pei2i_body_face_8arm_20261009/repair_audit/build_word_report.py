"""Build the stage Word report from locked metadata and native images."""
import collections,datetime,hashlib,json,pathlib,zipfile
from PIL import Image
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=pathlib.Path(__file__).resolve().parents[1]
AUDIT=ROOT/'repair_audit'
REPORT=ROOT/'report'
OUT=REPORT/'PEI2I_BodyFace_实验报告_阶段版_20261011.docx'
conditions=json.loads((REPORT/'conditions.json').read_text(encoding='utf-8'))
summary=json.loads((REPORT/'summary.json').read_text(encoding='utf-8'))
history=json.loads((REPORT/'history/references.json').read_text(encoding='utf-8'))
inputs=json.loads((AUDIT/'word_inputs/manifest.json').read_text(encoding='utf-8'))
qa=json.loads((REPORT/'development_qa8_v1/records.json').read_text(encoding='utf-8'))
smoke=json.loads((AUDIT/'measurement_cpu_smoke.json').read_text(encoding='utf-8'))
look={r['id']:{i['pair_id']:i for i in r['images']} for r in history['records']}
order=[r['id'] for r in history['records']]
zone=datetime.timezone(datetime.timedelta(hours=8))
snapshot=datetime.datetime.fromtimestamp(inputs['captured_at'],zone).strftime('%Y年%m月%d日 %H:%M')
embedded=[]

def check(path,expected):
    if hashlib.sha256(path.read_bytes()).hexdigest()!=expected:raise ValueError(f'Hash mismatch: {path}')
for a in history['records']:
    for i in a['images']:check(REPORT/i['path'],i['sha256'])
for i in inputs['files']:check(AUDIT/i['path'],i['sha256'])
assert sum(len(r['images']) for r in history['records'])==81
assert summary['train_updates']==0 and summary['completed_eval_jobs']==0

doc=Document()
sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=Inches(.65);sec.bottom_margin=Inches(.65)
sec.left_margin=Inches(.65);sec.right_margin=Inches(.65)
sec.footer_distance=Inches(.28)
WIDTH=7.2
for name,size in [('Normal',11),('Title',24),('Subtitle',12),('Heading 1',17),('Heading 2',13),('Heading 3',11),('Caption',9)]:
    s=doc.styles[name];s.font.name='Calibri';s.font.size=Pt(size);s.font.color.rgb=RGBColor(0,0,0)
    s.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Microsoft YaHei')
    s.paragraph_format.space_after=Pt(7)
    s.paragraph_format.line_spacing=1.15
    for e in list(s.element.get_or_add_pPr()):
        if e.tag==qn('w:pBdr'):s.element.get_or_add_pPr().remove(e)
doc.styles['Normal'].paragraph_format.widow_control=True
for name in ['Heading 1','Heading 2','Heading 3']:
    doc.styles[name].paragraph_format.keep_with_next=True
    doc.styles[name].paragraph_format.space_before=Pt(9)
    doc.styles[name].paragraph_format.space_after=Pt(8)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=footer.add_run('第 ');r.font.size=Pt(9)
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');footer._p.append(f)
r=footer.add_run(' 页');r.font.size=Pt(9)
doc.core_properties.title='TikTok PEI2I Body Face 联合适配实验报告'
doc.core_properties.subject='实验条件 历史结果 PE提示词预检与八组正式实验设计'
doc.core_properties.author='PictureEdit 项目'
doc.core_properties.comments=''

def p(text='',bold=False,size=None,after=7):
    a=doc.add_paragraph();a.paragraph_format.space_after=Pt(after)
    r=a.add_run(text);r.bold=bold
    if size:r.font.size=Pt(size)
    return a
def h(text,level=1):doc.add_heading(text,level=level)
def page(title):doc.add_page_break();h(title)
def cap(text):
    a=doc.add_paragraph(text,style='Caption');a.paragraph_format.keep_with_next=False
    return a
def table(headers,rows,widths=None,size=9.5):
    t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
    widths=widths or [WIDTH/len(headers)]*len(headers)
    for col,w in zip(t.columns,widths):col.width=Inches(w)
    props=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for key in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+key);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    props.append(borders)
    margin=OxmlElement('w:tblCellMar')
    for key in ['top','bottom','left','right']:
        e=OxmlElement('w:'+key);e.set(qn('w:w'),'65');e.set(qn('w:type'),'dxa');margin.append(e)
    props.append(margin)
    rep=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(rep)
    for i,label in enumerate(headers):
        c=t.rows[0].cells[i];c.text=str(label)
        sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'F2F2F2');c._tc.get_or_add_tcPr().append(sh)
    for row in rows:
        cells=t.add_row().cells
        for c,text in zip(cells,row):c.text=str(text)
    for row in t.rows:
        no=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no)
        for c,w in zip(row.cells,widths):
            c.width=Inches(w)
            for a in c.paragraphs:
                a.paragraph_format.space_after=Pt(3);a.paragraph_format.space_before=Pt(2);a.paragraph_format.line_spacing=1.05
                for r in a.runs:r.font.size=Pt(size);r.bold=(row==t.rows[0])
    return t
def image_into(a,path,width,height_limit=None,alt=''):
    path=pathlib.Path(path)
    with Image.open(path) as im:w,h=im.size
    use=width if height_limit is None else min(width,height_limit*w/h)
    r=a.add_run();shape=r.add_picture(str(path),width=Inches(use))
    shape._inline.docPr.set('descr',alt or path.name)
    embedded.append({'path':str(path),'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'size':[w,h],'alt':alt})
def picture_row(labels,paths,width=7.2,height=2.15):
    t=table(labels,[['']*len(labels)],[width/len(labels)]*len(labels),size=9)
    for c,path,label in zip(t.rows[1].cells,paths,labels):
        a=c.paragraphs[0];a.alignment=WD_ALIGN_PARAGRAPH.CENTER
        a.paragraph_format.space_before=Pt(0);a.paragraph_format.space_after=Pt(0)
        if path:image_into(a,path,width/len(labels)-.10,height,alt=label)
        else:
            rr=a.add_run('该阶段未生成');rr.font.size=Pt(8)
    return t

doc.add_paragraph('TikTok PEI2I Body Face\n联合适配实验报告',style='Title')
doc.add_paragraph('实验条件  已有结果  图像对照  提示词预检',style='Subtitle')
p('阶段版  2026年10月11日',bold=True)
p(f'运行服务器为 SSH sg（node01）。本报告的运行状态快照时间为 {snapshot}（北京时间）。')
h('当前结论',2)
p('历史联合训练已保留 81 张原尺寸推理图。部分身体动作能够跟随目标，但头部朝向、表情和手部仍有明显偏差；仅挂载 BFS Head 或短程续训，没有形成所有样本上一致的改善。历史图片可以定位问题，不能直接证明新提示词方案或 LoRA 底座优于其他方案。')
p('新的正式研究比较两个冻结底座与四种提示词协议，共八组。快照时正式训练更新数为 0，正式完成评测任务数为 0；PE 提示词第三版正在预检。因而本阶段没有八组排名、最终测试分数或正式训练曲线。')
table(['证据类别','已存在的产物','本报告如何使用'],[
    ['历史联合实验','J256 24图；Qwen续训24图；Head热挂载9图；Head续训24图','81张图逐样本对齐；注明实际预算'],
    ['PE 第一版预检','8对样本 × PG/PB/PBG = 24条真实提示词','保留输入图并说明拒绝原因'],
    ['PE 第三版预检',f"PG {inputs['qa_v3_counts']['pg']}/8；PB {inputs['qa_v3_counts']['pb']}/8；PBG {inputs['qa_v3_counts']['pbg']}/8",'生成中；未签署通过'],
    ['正式八组实验','0次训练更新；0项完成评测','列出锁定条件与验收标准']
],[1.4,2.8,3.0])
h('阅读顺序',2)
p('先查看实验矩阵、几何输入和训练预算；随后阅读历史结果的评价及 24 对样本全量图集；最后查看 PE 输入图、提示词问题、验证记录和文件索引。图集中的空格均有明确原因，未计入成功结果。',size=10)

page('实验目标与八组对照')
p('任务是保留源图 A 的人物身份、发型、服装设计和场景，将身体姿态、可靠的头部朝向及可观察表情调整到目标 B。服装颜色、花纹和材质保持；动作引起的褶皱、遮挡和阴影允许改变。')
table(['提示词协议','PE 可观察的图像','冻结 Qwen 底座','冻结 Qwen 加 Head'],[
    ['P0 固定短提示词','不调用 PE','qwen_p0','head_p0'],
    ['PG 几何描述','A + B depth + B normal + B过滤DWPose','qwen_pg','head_pg'],
    ['PB 目标图描述','A + B','qwen_pb','head_pb'],
    ['PBG 目标与几何描述','A + B + B depth + B normal + B过滤DWPose','qwen_pbg','head_pbg'],
],[1.3,3.0,1.45,1.45])
h('两个底座和共同起点',2)
p('底座一为 Qwen Image 2.1；底座二为同一模型加冻结 BFS Head V1.1，LoRA 强度 1.0，从本研究第 0 次更新起挂载。Qwen 和 LoRA 均冻结，仅 Body 与 Face Adapter 联合训练。BFS Body 不属于本轮八组矩阵。')
p('八组都加载同一份历史 J256 联合 Adapter，重新创建 AdamW，各追加 1024 次更新。P0 也从同一 J256 起点执行相同预算。不存在某些组先多训练 256 步再进入比较的安排。')
h('固定短提示词',2)
p('Change the person pose to match the target pose. Preserve identity, clothing, and background.',size=11)
p('P0 复用历史实验的实际提示词；PG、PB、PBG 为每个样本分别重写。三种 PE 结果在两套底座间复用，不能为某个底座单独人工优化。',size=10)
h('比较范围',2)
p('共同 J256 曾在纯 Qwen 上训练，因此本轮回答“已有联合 Adapter 在不同文本和底座上如何继续适配”，不能外推到随机初始化 Adapter 的最优底座。PB/PBG 的 PE 观察 B RGB，属于带目标图信息的提示词协议；PG 更接近仅有目标几何的部署方式。')

page('几何条件与图像输入边界')
table(['通道','本轮实际用途','训练状态'],[
    ['源图 A RGB','唯一的 Qwen 参考 RGB；提供身份、服装和场景','参考条件'],
    ['B normal 及 Body25','Body 的局部空间条件；Body25 为关节点热图','Body 可训练'],
    ['part 及人体 mask','Body 的分区与人体空间约束','Body 可训练'],
    ['26维 SMPL X global','全局几何及体型信息；并非全部 Face 输入','Body 可训练'],
    ['目标72点 jaw expression','Face 单独接收；并结合 A 身份特征','Face 可训练'],
    ['B depth 及过滤DWPose','给 PE 观察并保留可视化','Depth 与 Hand Adapter 关闭'],
    ['Hand Interaction 共享Detail','本轮不参加条件注入或更新','关闭'],
],[1.9,3.8,1.5])
p('旧网页的 pose 小点图是关节点热图可视化，不能把它理解为完整 DWPose 骨架，更不能据此断言面部和手部 SMPL X 信息完全未传入。Body 的 26 维 global 与 Face 的 72 点、jaw、expression 是不同输入通道。Hand Adapter 本轮关闭，手指精细控制因此仍是限制。')
h('原尺寸规则',2)
p('RGB 保持原始像素宽高，不统一 resize 为 512方图。VAE 仅补到 16 的倍数，例如 604×1080 补为 608×1088，输出再去补边回 604×1080。视觉编码器内部预处理单独记录。Word 中的缩放仅用于排版，嵌入的是原始 PNG。')
h('GT 的使用边界',2)
p('PB/PBG 允许 PE 看 B，Qwen 的参考 RGB 始终只有 A。训练时 B 仍需编码为 flow 监督目标和加噪训练状态；这属于监督训练，不是把 B 加为参考图。验证、诊断和测试仅推理，允许离线读取 B 算指标，禁止 backward、optimizer step 或训练损失调用。')
p('提示词无需告诉模型“由 Adapter 调整细节”。Adapter 在模型前向计算中注入几何残差；文本应直接描述可观察动作和保留项。PG 对不可可靠观察的表情、视线和遮挡手指应留白，不能要求冻结 A 的表情或凭裸模脸猜表情。',size=10)

page('数据划分与训练条件')
table(['条件','锁定设置'],[
    ['训练与验证','训练512对；验证24对；跨人物诊断6对；独立测试24对'],
    ['数据隔离','按原视频组隔离；测试仅按输入资产完整性筛选'],
    ['每组预算','J256起点 + 1024次额外更新；八组合计8192次正式更新'],
    ['数据顺序','两轮512样本；第二轮seed1709打乱；同组同序'],
    ['批量与精度','batch 1；无梯度累积；Adapter及优化器FP32；BF16计算'],
    ['优化器','AdamW；β=(0.9,0.999)；eps=1e−8；weight decay=0.01'],
    ['学习率与裁剪','20步warmup后恒定；梯度clip=1.0'],
    ['Body三类学习率','2.5e−5 / 1.25e−5 / 2.5e−6'],
    ['Face三类学习率','5e−5 / 2.5e−5 / 5e−6'],
    ['训练目标','flow + 0.5 face ROI + 0.1 identity + 0.05 perceptual'],
    ['辅助损失门控','每4步检查；仅真实Face有效且sigma不高于0.35时启用'],
    ['检查点','0 / 128 / 256 / 384 / 512 / 640 / 768 / 896 / 1024'],
    ['推理参数','30步；CFG=1；Body=1；Face=1；Hand=0；seed42、43'],
],[1.65,5.55],size=10)
p('上述是正式实验锁定条件。当前 actual_training_configuration_files 为空，尚没有正式训练 config 或 loss 曲线可用于确认执行完成。完成后应以每组实际 config、checkpoint、optimizer 参数和签名对应结果核对。',size=10)
h('步数的解释',2)
p('1024 次额外更新比原有 256 次续训更充分，但不能预先断言已收敛。需要比较中间检查点：若姿态改善伴随身份或背景退化，应选择验证表现较稳的检查点，同时保留最终 1024 步结果。多个推理 seed 不能替代独立训练重复。')

page('推理保存与评测规则')
table(['数据用途','样本与检查点','每组推理次数'],[
    ['训练可视化','固定8对 × 9检查点；seed42','72'],
    ['完整验证','24对 × 0/256/512/768/1024；seed42、43','240'],
    ['中间验证','固定9对 × 128/384/640/896；seed42','36'],
    ['跨人物诊断','6对 × 0/512/1024；seed42','18'],
    ['最终测试','24对 × 0/1024；seed42、43','96'],
    ['合计','462次每组；8组合计3696次','462'],
],[1.3,4.8,1.1],size=10)
p('最终测试在八组训练全部完成后统一开放。所有测试样本只进行前向推理与离线评分，不参加训练更新。PE 正式重写总预算为 566对 × 3协议 = 1698条；最终测试前1626条，测试72条。已通过 QA 的24条仅在请求签名完全相同时复用。')
table(['指标','具体定义及缺失规则'],[
    ['身份相似度','只比较 A 与生成图的唯一有效人脸；检测失败记缺失并报告原因'],
    ['姿态PCK与NME','对照 B 原尺寸Body25；检测失败记缺失，不能用热图自评分代替'],
    ['背景L1','原尺寸RGB；排除A/B人体mask并集扩张8像素及画布四边8像素'],
    ['GT像素L1','仅同人物配对可选；跨人物诊断禁止使用整图 B 像素误差'],
    ['汇总方法','seed先在pair内求均值，再按原视频组宏平均；公开有效样本数'],
],[1.4,5.8],size=9.5)
p('目前没有这些正式输出分数；检测器 CPU 冒烟测试中的 A 与自身相似度接近1，只能说明身份模型可工作，不能作为生成身份保持分数。',size=10)
h('每张图应保留的信息',2)
p('保留样本ID、A/B、depth/normal/DWPose、目标 Face 与 Hand ROI、完整提示词和原始 PE 回复、协议、底座、LoRA SHA、Adapter checkpoint SHA、追加更新数、原始宽高、seed、30步推理配置，以及失败记录。不得删除困难样本或用好看的结果替换固定样本。')

page('历史实验条件与结果')
table(['历史阶段','共同起点后追加更新','LoRA加载方式','已保存推理图'],[
    ['J common','0；名义联合步骤256','未挂载','24 / 24'],
    ['Qwen final','256；名义联合步骤512','未挂载','24 / 24'],
    ['BFS Head hotload','0；名义联合步骤256','J256后直接热挂载','9 / 9'],
    ['BFS Head final','256；名义联合步骤512','挂载后继续适配Adapter','24 / 24'],
],[1.6,2.0,2.2,1.4],size=10)
p('以上名义步骤只表示该历史联合阶段的256加256，不等于 Adapter 全部预训练历史。此前 Body/Face 权重的来源仍以旧初始化 provenance 为准。四阶段均使用实际固定短提示词，native尺寸，seed42，30步，CFG1，Body1及Face1；Head强度1。')
p('这24对在旧配置中标记为 pilot diagnostic pairs，并非最终独立测试。每个24图阶段有23对 Face condition 有效；这是条件可用性统计，不是人脸生成成功率。Head热挂载只有9图，其余15格缺失。')
h('逐图可观察的结果',2)
table(['样本','实际图像中的现象','结论边界'],[
    ['val_0000666','双臂举过头的宏观动作在几个阶段均能体现','说明有身体姿态控制能力'],
    ['val_0000868','B明显仰头，生成图仍大多朝前','头部方向未精确服从目标'],
    ['val_0000594','目标闭口及手势与生成图张口和手部配置有差异','表情、手部仍需独立核查'],
    ['val_0000941','前伸手臂周围可见重影或变形','原图与高分辨率局部共同检查'],
    ['val_0000175','本次关注的酒红上衣样本，手位与头部仍有偏差','不能凭单张判断Head获胜'],
],[1.4,3.6,2.2],size=9.5)
p('Qwen续训和Head续训的差异因样本而异，没有足够证据认定某个阶段稳定更优。热挂载0步与续训256步的对比含有不同训练预算，应分别解读。历史图集没有统一重新评分，因而不提供伪造的均值、排名或显著性结论。',size=10)

# Two paired comparisons per page, all existing 81 outputs are embedded natively.
labels=['源图 A','目标 B','J256','Qwen +256','Head热挂载','Head +256']
for start in range(0,24,2):
    page(f'历史图像对照 {start+1} 至 {start+2}')
    p('全部为历史固定短提示词；seed42，30步，CFG1，Body1，Face1。B仅供离线核对。',size=9,after=6)
    for pair in inputs['pairs'][start:start+2]:
        pid=pair['pair_id'];h(pid.replace('_',' '),2)
        p(f"视频 {pair['sequence_id']}  A帧 {pair['source_stem']}  B帧 {pair['target_stem']}  原视频组 {pair['source_video_group']}",size=9,after=4)
        paths=[AUDIT/pair['inputs'][k]['path'] for k in ['source_rgb','target_rgb']]
        paths += [REPORT/look[k][pid]['path'] if pid in look[k] else None for k in order]
        picture_row(labels,paths,height=1.98)
        with Image.open(paths[0]) as im:size=im.size
        cap(f"样本 {pid}  原尺寸 {size[0]}×{size[1]}。J256/热挂载追加0步；Qwen/Head续训追加256步。" + (' 热挂载未覆盖该样本。' if paths[4] is None else ''))
    p('此页为阶段诊断图，未计入新八组正式结果。图片在Word内嵌入，打开文档无需连接sg。',size=9,after=0)

page('PE 提示词预检结果')
p('第一版 QA 在固定8对上真实调用 PE，形成 PG/PB/PBG 共24条回复。预检发现部分编辑指令误读几何或改变了应保留的内容，未批准进入正式训练。第二版曾把 A 表情和头部方向锁死，与 Face 目标冲突，生成4条PG后停止。第三版统一修订模板并重新预检，快照时尚在生成。')
table(['版本','实际产物状态','发现与处理'],[
    ['第一版','PG8 PB8 PBG8，24条','拒绝；保留输入图和原始回复'],
    ['第二版','PG4，其余未生成','停止；避免冻结源头姿与表情'],
    ['第三版',f"PG{inputs['qa_v3_counts']['pg']} PB{inputs['qa_v3_counts']['pb']} PBG{inputs['qa_v3_counts']['pbg']}",'生成中；待24条逐一审阅后签名'],
],[1.1,2.1,4.0],size=10)
table(['第一版问题','对应样本或证据','修订方向'],[
    ['删除源道具','train_0007410 的PG要求去掉高尔夫球杆','保持A可见道具，不擅自删改'],
    ['猜不可见肢体','train_0002794及0000092的PG描述画外腿或错误手指','只写原画幅内有证据的姿态'],
    ['猜表情和视线','从裸模normal推断闭口、微笑或说话','PG不确定则省略，由Face条件控制'],
    ['依赖未给Qwen的参考图','compiled prompt仍多次指代几何guide','编译为可独立理解的编辑描述'],
],[1.4,3.0,2.8],size=9.5)
h('PE 与 Adapter 的协调要求',2)
p('PE 负责把目标动作转成明确、可执行的文本，同时保留 A 的身份、服装和场景。几何残差负责空间细节。第三版要求使用可靠的头部方向，未知表情不写；不得把“保持身份”扩展为“保持A表情和头部朝向”。')
p('下面四页为第一版 QA 的全部8对原始输入。其几何可视化用于解释错误来源，不能代表已经合格的正式提示词；完整24条原文和raw reply保留在本地QA网页。',size=10)

qa_index={proto:{a['pair_id']:a for a in rows} for proto,rows in qa.items()}
qa_ids=[a['pair_id'] for a in qa['pg']]
for start in range(0,8,2):
    page(f'PE 预检输入图 {start+1} 至 {start+2}')
    p('PG看A及三个几何图；PB看A和B；PBG看全部五图。B不作为Qwen参考RGB。',size=9,after=6)
    for pid in qa_ids[start:start+2]:
        h(pid.replace('_',' '),2)
        roles={a['role']:a for proto in ['pg','pb','pbg'] for a in qa_index[proto][pid]['inputs']}
        paths=[REPORT/'development_qa8_v1/assets'/f"{roles[k]['sha256']}.png" for k in ['source_A','target_B','target_depth','target_normal','target_dwpose']]
        # Actual archived roles are checked rather than inferred from table labels.
        for key,path in zip(['source_A','target_B','target_depth','target_normal','target_dwpose'],paths):check(path,roles[key]['sha256'])
        picture_row(['源图 A','目标 B','B depth','B normal','B DWPose'],paths,height=2.25)
        cap(f'样本 {pid}  第一版预检输入  原图604×1080  三协议各一条真实回复  未批准用于正式训练')

page('提示词实例与修订依据')
pid='train_0002794';old=qa_index['pg'][pid]['compiled_prompt']
p('样本 train_0002794 的 PG 第一版同时要求保持 calm closed-mouth expression，并描述画幅外腿部和具体交叠手指。A只是身份和衣服锚点，normal不能可靠证明这些表情或被遮挡手指。上述描述会与几何目标或画幅产生冲突。')
p('第一版实际指令片段',bold=True)
p('“keep the head upright and facing the camera with the same calm closed-mouth expression”\n“both hands fully visible at the center of the chest with fingers extended horizontally and overlapping”\n“adopt the leg stance ... with legs apart and the viewer-left knee bent outward”',size=10)
h('第三版同一样本的实际输出',2)
v3path=AUDIT/'qa8_v3_pg_current.jsonl'
v3=[json.loads(s) for s in v3path.read_text(encoding='utf-8').splitlines() if s.strip()][0]
assert v3['pair_id']==pid
p(v3['compiled_prompt'],size=9.5)
p('此条是第三版已生成实例，取消了冻结 A 表情的要求；仍需与全部输入图及其余23条联合审阅，不等于第三版整体通过，更不是生成效果的证据。',size=9)

page('验证记录与阶段结论')
table(['验证内容','实际结果','能够说明什么'],[
    ['服务器单元检查','326通过，5条warning','代码契约和防护检查通过'],
    ['本地网页与文件','237文件SHA核对；相对链接无缺失','当前网页及历史图可离线查看'],
    ['Word新增输入下载','24对 × A/B/normal/depth/DWPose = 120文件SHA通过','图像来源与本地副本一致'],
    ['真实CPU身份检测','8个A各有1张有效人脸','评分模型可以读取这些输入'],
    ['真实CPU目标姿态检测','8个B均检出；有效关节12至17个','这些输入可进入姿态评分'],
    ['实际生成数量','冒烟测试正式生成数0；八组正式评测0','不能把输入自检写成生成质量结果'],
],[1.7,2.6,2.9],size=9.5)
h('本阶段已经能够确认',2)
p('第一，Body与Face的信息通道需要分别核对，不能从26维global或网页pose小点图推断Face缺失。第二，PE看GT能获得更丰富的动作语义，但其优势来自更丰富的输入，不能与仅几何协议混为同一部署条件。第三，详细提示词并不天然更好，错误的表情、画外肢体或保留约束会造成条件冲突。')
h('正式实验完成后的判断标准',2)
p('在相同底座内部比较P0与三种PE，观察姿态提升是否以身份或背景退化为代价；在相同协议内部比较Qwen与Head，核对初始0步与1024步的差异。优先依据独立测试的按视频组汇总指标及缺失比例，再用全量图像检查脸、手、衣服纹理与失败案例。')
p('需要保留的限制包括：共同J256偏向纯Qwen；Head原始任务常用两张参考RGB而本研究只给A；24对测试规模有限；只有一条训练轨迹，42和43是推理随机种子；B的global包含体型信息，可能与A体型保持相互影响。',size=10)

index_start=len(doc.paragraphs)
page('文件索引与复核方法')
p('服务器独立工作目录',bold=True)
p('/data/shangguanrz/pic-edit-experiments/pei2i_body_face_8arm_20261009',size=10)
p('本地文档和图像根目录',bold=True)
p(str(ROOT).replace('\\','/'),size=10)
table(['本地相对路径','用途'],[
    ['report/index.html','中文阶段网页和历史逐图对照'],
    ['report/history/references.json','81图的实际配置、加载方式、预算及SHA'],
    ['report/history/*.png','历史推理原尺寸PNG'],
    ['report/development_qa8_v1/index.html','第一版24条PE原文和输入图'],
    ['report/conditions.json','八组锁定条件与实际训练config索引'],
    ['report/summary.json','正式进度及完成门禁'],
    ['repair_audit/word_inputs/manifest.json','新增120张图的来源、样本、视频组及SHA'],
    ['repair_audit/word_report_manifest.json','Word来源文件签名与内嵌图清单'],
],[3.6,3.6],size=9.5)
h('重要资产指纹',2)
for label,digest in [
    ('共同J256',conditions['common_adapter_initialization_sha256']),
    ('BFS Head V1.1',conditions['data_lock']['artifacts']['head_lora']['sha256']),
    ('训练512对','854133776d2067aa392442553f3003ef9e6d45107ccc9cffe5a4af8600fb929f'),
    ('验证24对','cece883b96dbbc9270fd77d4bb1e973470f52f941af85c2a0746947f5d91e10a')]:
    p(label+'\n'+digest,size=9,after=6)
h('使用和测试',2)
p('直接用Word打开本报告，放大比较历史图集，同一行按A、B、J256、Qwen续训、Head热挂载、Head续训顺序核对。所有图已嵌入文档，离线可读。需要完整PE原文或查看原尺寸单图时，打开上述本地网页或PNG目录。',size=10)
p('远端状态查看 repair_audit/qa8_v3_launch.json 与 full_pipeline_state.json；正式进度以summary.json和各组实际checkpoint为准。正在运行时不要重复启动同一阶段。运行及复核步骤见项目 docs/experiments/2026-10-10-pei2i-eight-arm-runbook.md。',size=10)

# Keep the final source index and usage instructions on one readable page.
for paragraph in doc.paragraphs[index_start:]:
    paragraph.paragraph_format.line_spacing=1.0
    paragraph.paragraph_format.space_after=Pt(4)

doc.save(OUT)
sources=[REPORT/'conditions.json',REPORT/'summary.json',REPORT/'history/references.json',REPORT/'development_qa8_v1/records.json',AUDIT/'word_inputs/manifest.json',AUDIT/'qa8_v3_pg_current.jsonl',AUDIT/'measurement_cpu_smoke.json']
record={
    'schema':'pictureedit_word_stage_report_v1','snapshot':snapshot,'docx':str(OUT),
    'docx_sha256':hashlib.sha256(OUT.read_bytes()).hexdigest(),
    'historical_generation_count':81,'formal_train_updates':0,'formal_completed_eval_jobs':0,
    'qa_v3_counts':inputs['qa_v3_counts'],
    'sources':[{'path':str(a),'sha256':hashlib.sha256(a.read_bytes()).hexdigest()} for a in sources],
    'embedded_images':embedded,
}
(AUDIT/'word_report_manifest.json').write_text(json.dumps(record,ensure_ascii=False,indent=2),encoding='utf-8')
with zipfile.ZipFile(OUT) as z:
    media=[a for a in z.namelist() if a.startswith('word/media/')]
print(json.dumps({'output':str(OUT),'bytes':OUT.stat().st_size,'image_instances':len(embedded),'unique_media':len(media),'historical_outputs':81},ensure_ascii=False))
