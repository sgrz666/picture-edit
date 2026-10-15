#!/usr/bin/env python3
"""Build a Word acceptance report and full-resolution two-row comparison PNGs.

Reads recorded results only; never generates images or runs model inference.
Missing evidence is explicitly marked incomplete. Use --require-complete for
the final delivery gate; an intermediate report can be built without it.
"""
from __future__ import annotations

import argparse
from collections import defaultdict
from datetime import datetime
import hashlib
import json
import math
from pathlib import Path
from zoneinfo import ZoneInfo

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from PIL import Image, ImageDraw, ImageFont
import torch


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_ROOT = Path("/home/shangguanrz/nihaochen/qwen21_full")
DEFAULT_OUT = Path("/home/shangguanrz/nihaochen/docs/qwen21_full_迁移与短训验收.docx")
MODES = ("initial", "body", "body_face")
LABELS = {"initial": "初始化", "body": "Body 100步", "body_face": "Body＋Face 100步"}
FONT = "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"
FIELDS = (("PSNR ↑", "psnr", True), ("SSIM ↑", "ssim", True),
          ("人体 MAE ↓", "human_mae", False), ("背景 MAE ↓", "background_mae", False),
          ("单张生成耗时／秒 ↓", "seconds_mean", False), ("峰值显存／GiB ↓", "peak_cuda_gib", False))


def arguments():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("--root", type=Path, default=DEFAULT_ROOT)
    p.add_argument("--out", type=Path, default=DEFAULT_OUT)
    p.add_argument("--require-complete", action="store_true")
    return p.parse_args()


def read_json(path, default=None):
    path = Path(path)
    return json.loads(path.read_text()) if path.is_file() else default


def read_jsonl(path):
    path = Path(path)
    return [json.loads(line) for line in path.read_text().splitlines() if line.strip()] if path.is_file() else []


def finite(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value)


def number(value, digits=3):
    return f"{value:.{digits}f}" if finite(value) else "未评测"


def status(value):
    return "通过" if value is True else "未通过" if value is False else "未完成"


def add_table(doc, headers, rows, *, font_size=10):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = str(text)
        for run in cell.paragraphs[0].runs:
            run.bold = True
    repeat = OxmlElement("w:tblHeader")
    table.rows[0]._tr.get_or_add_trPr().append(repeat)
    for values in rows:
        for cell, value in zip(table.add_row().cells, values):
            cell.text = str(value)
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(3)
                for run in paragraph.runs:
                    run.font.size = Pt(font_size)
    return table


def configure_doc(doc):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
    sec.left_margin = sec.right_margin = Inches(.65)
    sec.top_margin = sec.bottom_margin = Inches(.55)
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Caption"):
        style = doc.styles[name]
        style.font.name = "Arial"
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
    doc.styles["Normal"].font.size = Pt(10)
    doc.styles["Normal"].paragraph_format.space_after = Pt(6)
    settings = doc.settings.element
    compress = OxmlElement("w:doNotAutoCompressPictures")
    settings.append(compress)
    for section in doc.sections:
        footer = section.footer.paragraphs[0]
        footer.alignment = 2
        run = footer.add_run("qwen21_full｜迁移与短训验收  ")
        run.font.size = Pt(8)
        fld = OxmlElement("w:fldSimple")
        fld.set(qn("w:instr"), "PAGE")
        footer._p.append(fld)


def verify_copy(manifest):
    rows = []
    for item in manifest.get("files", []):
        source = Path(manifest["source"]) / item["relative_path"]
        target = Path(manifest["destination"]) / item["relative_path"]
        source_hash = hashlib.sha256(source.read_bytes()).hexdigest() if source.is_file() else None
        target_hash = hashlib.sha256(target.read_bytes()).hexdigest() if target.is_file() else None
        rows.append({"name": item["relative_path"], "source_snapshot_sha256": item["source_sha256"],
                     "current_source_sha256": source_hash, "current_target_sha256": target_hash,
                     "matches_snapshot": source_hash == target_hash == item["source_sha256"]})
    return rows


def evaluation(root, name, count):
    directory = root / name
    summary = read_json(directory / "summary.json", {})
    rows = read_jsonl(directory / "per_sample.jsonl")
    # A restarted evaluation can leave duplicate lines. Require the last record
    # per pair/mode, and require the actual saved output as well as the record.
    indexed = {(r["pair_index"], r["mode"]): r for r in rows}
    missing = []
    for index in range(count):
        for mode in MODES:
            row = indexed.get((index, mode))
            if row is None or not Path(row["image"]).is_file():
                missing.append(f"{name}/{index:03d}/{mode}")
            elif row.get("control_forwards") != 30:
                missing.append(f"{name}/{index:03d}/{mode}:非30步完整采样")
            else:
                calls = row.get("block_calls", {})
                if len(calls) != 32 or any(value != 30 for value in calls.values()):
                    missing.append(f"{name}/{index:03d}/{mode}:32层控制调用记录不完整")
                for field in ("source", "target", "image"):
                    path = Path(row.get(field, "/missing"))
                    if not path.is_file():
                        missing.append(f"{name}/{index:03d}/{mode}:缺少{field}图片")
                        continue
                    try:
                        with Image.open(path) as image:
                            if image.size != (row["resolution"], row["resolution"]):
                                missing.append(f"{name}/{index:03d}/{mode}:{field}图片尺寸不符")
                            image.verify()
                    except (OSError, ValueError) as exc:
                        missing.append(f"{name}/{index:03d}/{mode}:{field}图片不可读取({type(exc).__name__})")
    if not summary:
        missing.append(f"{name}/summary.json")
    else:
        for mode in MODES:
            if summary.get(mode, {}).get("n") != count:
                missing.append(f"{name}/summary.json:{mode}数量不是{count}")
    return {"name": name, "count": count, "summary": summary, "indexed": indexed,
            "missing": missing, "config": read_json(directory / "config.json", {})}


def add_metrics(doc, item, title):
    doc.add_heading(title, level=2)
    doc.add_paragraph(f"每组计划 {item['count']} 对；n 表示该组实际纳入汇总的图片数量。"
                      "加粗只表示本行已有数值中最好的一项，不代表统计显著。")
    summary = item["summary"]
    table = add_table(doc, ["指标", *(LABELS[m] for m in MODES)], [])
    counts = table.add_row().cells
    counts[0].text = "有效图片数 n"
    for i, mode in enumerate(MODES, 1):
        counts[i].text = str(summary.get(mode, {}).get("n", "未完成"))
    for label, key, higher in FIELDS:
        values = [summary.get(mode, {}).get(key) for mode in MODES]
        present = [v for v in values if finite(v)]
        best = (max(present) if higher else min(present)) if present else None
        cells = table.add_row().cells
        cells[0].text = label
        for i, value in enumerate(values, 1):
            run = cells[i].paragraphs[0].add_run(number(value))
            run.bold = best is not None and finite(value) and math.isclose(value, best, rel_tol=1e-10, abs_tol=1e-10)
    doc.add_paragraph("像素指标统一在512×512上计算。人体与背景MAE均相对同人物真实Target计算；"
                      "这与此前相对Source计算的背景指标口径不同，不能直接拼在旧表中。"
                      "本轮未计算ArcFace面部保持度和姿态PCK，因此不能据此判断面部身份或动作已经改善。")
    if item["missing"]:
        doc.add_paragraph("状态：未完成。仍缺少 " + "、".join(item["missing"]) + "。")


def crop_face(image, box, scale=1.35):
    if box is None or len(box) != 4:
        return None
    x0, y0, x1, y1 = [float(v) for v in box]
    if not all(math.isfinite(v) for v in (x0, y0, x1, y1)) or x1 <= x0 or y1 <= y0:
        return None
    w, h = image.size
    cx, cy = (x0 + x1) * w / 2, (y0 + y1) * h / 2
    side = max((x1 - x0) * w, (y1 - y0) * h) * scale
    return image.crop((math.floor(cx - side / 2), math.floor(cy - side / 2),
                       math.ceil(cx + side / 2), math.ceil(cy + side / 2)))


def image_rgb(path):
    if not path or not Path(path).is_file():
        return None
    with Image.open(path) as source:
        image = source.convert("RGBA")
        return Image.alpha_composite(Image.new("RGBA", image.size, "white"), image).convert("RGB")


def make_panel(root, item, index, panel_dir, caches):
    records = [item["indexed"].get((index, mode)) for mode in MODES]
    anchor = next((row for row in records if row), None)
    if anchor is None:
        return None
    meta = anchor["metadata"]
    source, target = image_rgb(anchor["source"]), image_rgb(anchor["target"])
    output_images = [image_rgb(row["image"]) if row else None for row in records]
    resolution = int(anchor["resolution"])
    tile = resolution  # Preserve original 512/1024 pixels; no thumbnail step.
    gap = max(10, tile // 64)
    label_h, middle_h = max(48, tile // 10), max(34, tile // 16)
    width = tile * 5 + gap * 4
    canvas = Image.new("RGB", (width, label_h + tile + middle_h + tile), "white")
    draw = ImageDraw.Draw(canvas)
    font = ImageFont.truetype(FONT, max(21, tile // 23))
    small = ImageFont.truetype(FONT, max(17, tile // 30))
    boxes = {}
    for role in ("source", "target"):
        app = meta[f"{role}_appearance"]
        if app not in caches:
            path = root / "face_cache" / f"{app}.pt"
            caches[app] = torch.load(path, map_location="cpu", weights_only=False) if path.is_file() else None
        cache = caches[app]
        boxes[role] = None
        if cache is not None and meta[f"{role}_stem"] in cache["stem_to_idx"]:
            fi = cache["stem_to_idx"][meta[f"{role}_stem"]]
            if bool(cache["geometry_valid"][fi]):
                boxes[role] = cache["face_boxes"][fi].tolist()
    titles = ["Source：A原图", "Target：B目标", *(LABELS[m] for m in MODES)]
    # Use a fixed union window for all generated outputs: early Body training
    # may leave the head near A instead of B, and B-only crops would hide it.
    available = [box for box in boxes.values() if box is not None]
    generated_box = ([min(box[0] for box in available), min(box[1] for box in available),
                      max(box[2] for box in available), max(box[3] for box in available)] if available else None)
    images = [source, target, *output_images]
    for col, (title, image) in enumerate(zip(titles, images)):
        left = col * (tile + gap)
        bbox = draw.textbbox((0, 0), title, font=font)
        draw.text((left + max(5, (tile - bbox[2]) // 2), 4), title, fill="#17273c", font=font)
        if image is not None:
            canvas.paste(image.resize((tile, tile), Image.Resampling.LANCZOS), (left, label_h))
        else:
            draw.rectangle((left, label_h, left + tile, label_h + tile), fill="#eeeeee")
            draw.text((left + 20, label_h + 30), "未完成／缺少图片", fill="#aa3333", font=font)
        crop_box = boxes['source'] if col == 0 else boxes['target'] if col == 1 else generated_box
        crop = crop_face(image, crop_box) if image is not None else None
        face_top = label_h + tile + middle_h
        if crop is not None:
            canvas.paste(crop.resize((tile, tile), Image.Resampling.LANCZOS), (left, face_top))
        else:
            draw.rectangle((left, face_top, left + tile, face_top + tile), fill="#eeeeee")
            draw.text((left + 20, face_top + 30), "未完成／缺少图片" if image is None else "缺少可靠脸框", fill="#aa3333", font=font)
        draw.text((left + 8, label_h + tile + 2), "源脸放大" if col == 0 else "目标脸放大" if col == 1 else "统一头部窗口放大", fill="#425466", font=small)
    panel_dir.mkdir(parents=True, exist_ok=True)
    path = panel_dir / f"{item['name']}_{index:03d}.png"
    canvas.save(path, optimize=True)
    return {"image": path, "metadata": meta, "resolution": resolution,
            "source_box": boxes["source"], "target_box": boxes["target"], 'generated_box': generated_box,
            "size": canvas.size, "pair_index": index}


def main():
    args = arguments()
    root = args.root
    manifest = read_json(root / "migration_manifest.json", {})
    copied = verify_copy(manifest)
    data = read_json(root / "data_manifest.json", {})
    preprocess = read_json(root / "preprocess_manifest.json", {})
    face_supervision = read_json(root / "face_supervision.json", {})
    acceptances = {p.stem: read_json(p) for p in sorted((root / "acceptance").glob("*.json"))}
    resume_summary = acceptances.get("resume_summary", {})
    face_resume = read_json(root / "face_512_100/resume_verification.json", {})
    branches = {name: read_json(root / f"{name}_512_100/summary.json") for name in ("body", "face")}
    loading = read_json(root / "body_512_100/weight_loading.json", {})
    evals = [evaluation(root, "validation_512", 4), evaluation(root, "validation_1024", 2),
             evaluation(root, "cross_person_512", 2)]
    missing = [item for result in evals for item in result["missing"]]
    missing += [f"{name}_512_100/summary.json" for name, value in branches.items() if value is None]
    for name, value in branches.items():
        if value and (value.get("optimizer_steps") != 100 or value.get("frozen_parameters_unchanged") is not True
                      or value.get("base_versions_unchanged") is not True):
            missing.append(f"{name}_512_100:100步完成及冻结验收尚未通过")
    for resolution in (512, 1024):
        runtime = acceptances.get(f"runtime_{resolution}")
        if runtime is None:
            missing.append(f"acceptance/runtime_{resolution}.json")
        elif runtime.get("passed") is not True or runtime.get("status") != "complete":
            missing.append(f"acceptance/runtime_{resolution}.json:尚未通过")
    if not resume_summary:
        missing.append("acceptance/resume_summary.json")
    elif resume_summary.get("state_loading_exact") is not True or resume_summary.get("deterministic_reference_next_step_exact") is not True:
        missing.append("acceptance/resume_summary.json:状态恢复或确定性参考验收尚未通过")
    if not all(face_resume.get(key) is True for key in ("parameters_equal", "optimizer_equal", "rng_restored")):
        missing.append("face_512_100/resume_verification.json:参数/优化器/RNG恢复尚未通过")
    if not face_supervision:
        missing.append("face_supervision.json")
    if not manifest:
        missing.append("migration_manifest.json")
    if args.require_complete and missing:
        raise SystemExit("结果尚未齐全，未生成最终报告：" + "; ".join(missing))
    doc = Document()
    configure_doc(doc)
    doc.add_heading("qwen21_full：迁移与短训验收", 0)
    now = datetime.now(ZoneInfo("Asia/Shanghai")).strftime("%Y年%m月%d日 %H:%M")
    doc.add_paragraph(now + "｜保留旧控制架构，适配Qwen-Image-2.1")
    doc.add_paragraph("这份报告主要回答：旧模块有没有完整保留、能否正确连接Qwen、是否能训练并完整生成。"
                      "100步短训用于检查这条流程是否运转，不代表已经训练到可交付的画质。"
                      + ("部分验证仍未完成，文中逐项标明。" if missing else "本次规定的短训与生成记录已齐全；质量仍需独立评估。"))
    review = read_json(root / 'visual_review.json', {})
    if review:
        doc.add_paragraph("本轮结论：结构与运行检查通过，但短训画质尚不合格。512验证图的Body动作跟随仍不充分，"
                          "加入Face后的4个例子都出现了脸部彩斑或局部破坏；1024和跨人案例也出现同类问题。"
                          "这些是实际失败结果，不能被整图平均指标的小幅改善掩盖。"
                          "因此新检查点目前只用于研究和排查，不作为可用的面部优化方案，也未替换现有成果。")
    doc.add_heading("一、先看懂本次做了什么", level=1)
    add_table(doc, ["名称", "容易理解的说明"], [
        ("完整旧控制模块", "将原DeepGen控制目录中的模块完整复制保留：身体、面部、手部、外观与身份绑定、Reasoner、注意力控制核心及强度调节。没有用简化网络替代它们。"),
        ("Qwen接口", "新增负责连接的代码，把Qwen的图文信息交给旧控制器，再把控制器的修正信号转换成Qwen可使用的形式。"),
        ("Body Adapter", "身体姿态控制模块，利用B的动作几何，帮助A中的人物改变姿态；也保留旧版外观与身份绑定。"),
        ("Face Adapter", "面部控制模块：从A取得ArcFace身份及DINO局部纹理，从B取得头部位置、面部关键点、下颌和表情等几何，再进行局部控制。"),
        ("初始化", "已继承旧控制模块权重，但新Qwen输出连接还没有训练、初始输出为零；因此它用于检查适配训练前后的变化，不能当作已训练好的Qwen控制方案。"),
        ("Body 100步", "让Body及相关连接完成100次优化后的结果。"),
        ("Body＋Face 100步", "以上述Body结果为起点，冻结Body，再让Face完成100次训练迭代后的结果；实际有梯度的次数另列。"),
    ])
    doc.add_paragraph("这次Qwen的图片参考通道只接收完整A图。Face虽然使用A的脸部特征，但没有把裁剪脸额外作为第二张原生参考图。"
                      "B的人脸照片与身份特征不作为生成条件；同人物Target只用于训练监督及离线对照。")
    doc.add_heading("二、模块是否保留，权重来自哪里", level=1)
    all_same = bool(copied) and all(row["matches_snapshot"] for row in copied)
    doc.add_paragraph(f"原目录记录了 {len(copied)} 个源文件。当前逐文件SHA-256核对："
                      f"{sum(row['matches_snapshot'] for row in copied)}/{len(copied)} 与源快照及原文件一致。"
                      "新增连接、数据与执行代码另列，不覆盖旧DeepGen工作流或之前的简化Qwen版本。")
    preservation = acceptances.get("preservation_results", {})
    cpu_ok = preservation.get("exit_code") == 0 if preservation else None
    add_table(doc, ["模块", "本轮实际状态"], [
        ("几何编码、Reasoner、外观与身份绑定", f"完整复制；原结构前向/梯度一致性检查：{status(cpu_ok)}。Body短训状态见下一节。"),
        ("循环注意力控制核心、六阶段适配、强度调节", "保留原设计；新增Qwen输入/输出转换。六组控制分配到Qwen各层，保留按组控制的方式。"),
        ("完整Face：DINO、ArcFace、融合、细致几何与局部注意力", "保留原模块；已准备真实A身份/纹理及B几何。Face短训状态见下一节。"),
        ("Hand Adapter与旧细节分支", "代码及结构保留；没有进行真实手部训练或手部效果验收。不能称为已训练好的Qwen手部模型。"),
        ("双人推理、交互及接触条件", "代码保留并参与结构检查；本轮没有双人真实训练。跨人A/B案例仍是单人编辑，不能替代双人交互验收。"),
    ])
    add_table(doc, ["部分", "初始化来源"], [
        ("Body及旧控制核心", f"DeepGen身体20k检查点；实际载入 {loading.get('body_tensors_loaded', '未记录')} 个张量。不是继承之前简化Qwen的控制权重。"),
        ("Face", f"DeepGen V6.5单样本1500步检查点，仅加载Face部分；实际载入 {loading.get('face_tensors_loaded', '未记录')} 个张量，不覆盖Body。"),
        ("新增接口及无旧训练权重的模块", "按清单新初始化。尤其Hand及部分细节参数没有可继承的已训权重，保留结构不等于已完成训练。"),
        ("Qwen基座", "使用本地官方Qwen-Image-2.1；本轮保持冻结，训练的是Adapter及连接部分。"),
    ])
    doc.add_paragraph("逐项加载清单区分了三类状态：继承Body的477个张量、继承Face的242个张量；"
                      "另有247个旧结构内的状态沿用初始化，以及11个新增Qwen接口状态。"
                      "Face检查点中的538个非Face张量明确不加载，以免覆盖Body或混入其它实验权重。"
                      "所以，完整保留结构不等于每个分支都拥有已经训练好的权重。")
    doc.add_heading("三、数据和预处理的实际范围", level=1)
    counts = data.get("counts", {})
    doc.add_paragraph(f"本轮选定 {counts.get('train', '未记录')} 对训练样本、{counts.get('val', '未记录')} 对验证样本、"
                      f"{counts.get('cross_person', '未记录')} 对A/B不同人的检查案例。"
                      "一对就是一张A源图和一张B目标动作图。训练图负责更新参数；验证及跨人图负责检查输出，本轮不更新参数。原50对最终测试集未参与此次验收。")
    train_rows = read_jsonl(root / "train_pairs.jsonl")
    identities = sorted({row.get("appearance", row.get("source_appearance", "")).split("_")[0] for row in train_rows})
    doc.add_paragraph(f"训练清单按既定顺序优先选择不同外观；实际包含{len(train_rows)}种外观、{len(identities)}个人物编号"
                      f"（{', '.join(identities) or '未记录'}）。这8对属于同一个人的不同外观，不能代表丰富的身份覆盖。"
                      "验证外观未参与本轮短训，不等于所有验证人物身份都未见过。旧检查点此前的训练经历也没有因此消失。")
    fallback_count = sum(v.get("arcface_preprocess") == "original_A_detection_five_point_alignment_112"
                         for v in preprocess.get("records", {}).values())
    doc.add_paragraph("发现原预处理中的一个实际问题：整张A图能检测到脸，但把脸裁得很紧后再次检测会失败。"
                      "因此在这种情况下，改为使用整张A图已检出的五个关键点进行标准112像素对齐，再送入同一个ArcFace模型。"
                      f"本次有{fallback_count}张A使用这一处理；DINO仍用旧版1.35倍扩展裁剪及完整局部特征。"
                      "这项预处理修正已逐图记录，没有把失败样本偷偷当成功。")
    if face_supervision:
        summaries = face_supervision.get("summary", {})
        eligibility = "；".join(f"{res}像素：{summaries.get(res, {}).get('valid', '未记录')}/{summaries.get(res, {}).get('total', '未记录')}个Target通过筛选"
                               for res in ("512", "1024"))
        short_sides = []
        for key, record in face_supervision.get("records", {}).items():
            choice = record.get("by_resolution", {}).get("512", {})
            if choice.get("valid") and choice.get("selected_detection") is not None:
                detection = record["detections"][choice["selected_detection"]]
                value = detection.get("short_side_output_px", {}).get("512")
                if finite(value):
                    short_sides.append((value, key))
        smallest = min(short_sides) if short_sides else None
        doc.add_paragraph("训练时，Target的脸还需要通过额外的可见性筛选，才计算面部身份与感知损失："
                          "检测分数至少0.5、检测位置与几何脸框相符，并且按实际输出尺寸计算的脸框短边至少32像素。"
                          + eligibility + "。这12个Target来自8对训练与4对验证清单；验证记录只用于检查。"
                          "Target的RGB及相关身份损失只用于监督，绝不进入生成参考或Adapter的身份/纹理条件。")
        if smallest:
            doc.add_paragraph(f"512情况下，最小通过样本的脸框短边只有{smallest[0]:.1f}像素（{smallest[1]}），接近32像素门槛。"
                              "检测通过只是可见性的自动筛选依据，不等于已经人工确认正脸、无遮挡或细节清晰。")
    else:
        doc.add_paragraph("Target面部监督的可见性筛选尚未完成，不能把所有Target都默认当作可靠面部监督。")
    doc.add_heading("四、工程检查与100＋100步短训", level=1)
    doc.add_paragraph("模块结构和这次训练安排需要分开理解：旧Face实验固定一对图片、固定噪声，"
                      "这次则轮流使用8对图片和不同噪声；每阶段前20步只训练新增输出连接，之后训练对应Adapter。"
                      "Face普通训练在旧面部分支已启用的噪声范围内采样，每4步在更低噪声下计算身份和感知损失。"
                      "四项Face损失的系数、辅助裁剪方式及脸外区域定义沿用旧版；参数与优化器状态使用FP32。"
                      "这些训练安排与接口变化均已写入迁移对照清单，不能理解为照搬了旧单样本实验的所有训练设置。")
    checks = [("47个旧文件完整保留", status(all_same if copied else None), f"{sum(r['matches_snapshot'] for r in copied)}/{len(copied)}文件一致")]
    if preservation:
        checks.append(("旧结构前向、梯度及原测试复用", status(cpu_ok), json.dumps(preservation.get("counts", {}), ensure_ascii=False)))
    else:
        checks.append(("旧结构检查", "未完成", "缺少验收记录"))
    data_check = acceptances.get("data_runtime", {})
    data_ok = all(data_check.get(k) is True for k in ("geometry_bundle_valid", "invalid_source_face_fallback", "all_cache_fingerprints_valid")) if data_check else None
    checks.append(("A身份/B几何分离、无效脸回退、尺寸读取", status(data_ok), "覆盖512/1024与14对固定样本；详见数据验收记录" if data_check else "缺少记录"))
    parity = read_json(root / "body_512_100/native_parity.json")
    parity_ok = bool(parity["disabled_equal"] and parity["zero_max_abs"] == 0) if parity else None
    checks.append(("关闭/零控制时与Qwen原生一致", status(parity_ok), f"最大绝对差：{parity.get('zero_max_abs')}" if parity else "缺少记录"))
    for branch in ("body", "face"):
        value = branches[branch]
        frozen = value.get("frozen_parameters_unchanged") and value.get("base_versions_unchanged") if value else None
        checks.append((f"{branch.title()}短训期间冻结参数保持不变", status(frozen), "基座及不在本阶段训练的参数未变化" if frozen else "按实际记录判定"))
        if branch == "body":
            resume_ok = resume_summary.get("state_loading_exact") if resume_summary else None
            checks.append(("Body保存与加载", status(resume_ok), "参数与优化器状态精确恢复；续训的CUDA数值差异见下文" if resume_summary else "缺少最终恢复结论"))
            reference_ok = resume_summary.get("deterministic_reference_next_step_exact") if resume_summary else None
            checks.append(("Body确定性参考续训", status(reference_ok), "独立参考实现中下一步损失、梯度、参数更新与优化器状态完全一致" if reference_ok else "未完成或未通过"))
        else:
            resume_ok = all(face_resume.get(k) is True for k in ("parameters_equal", "optimizer_equal", "rng_restored")) if face_resume else None
            checks.append(("Face保存与加载", status(resume_ok), "参数、优化器精确恢复，随机数状态恢复；不宣称默认CUDA下一步逐位相同" if face_resume else "缺少恢复检查记录"))
            face_reference = acceptances.get('face_resume_deterministic_reference', {}).get('trials', {}).get('deterministic_reference', {})
            face_exact = bool(face_reference) and all(face_reference.get(key, {}).get('exact_equal') for key in ('gradients','parameters','optimizer'))
            checks.append(('Face确定性参考续训', status(face_exact), '同一下一个样本的损失、梯度及更新一致；参考计算不写回正式检查点'))
    for resolution in (512, 1024):
        runtime = acceptances.get(f"runtime_{resolution}")
        ok = runtime.get("passed") if runtime else None
        checks.append((f"{resolution}实际Qwen运行与反传", status(ok),
                       "检查Face关闭/无效脸回退、目标RGB隔离、脸区残差与梯度归属；不是画质评测" if runtime else "缺少运行验收记录"))
    add_table(doc, ["检查项", "结果", "依据或限制"], checks, font_size=9)
    if resume_summary:
        maximum = resume_summary.get("default_cuda_update_max_abs")
        difference = f"{maximum:.8g}" if finite(maximum) else "未记录"
        doc.add_paragraph("恢复检查需要区分两件事：保存后重新加载的参数和优化器状态能够完全相同；"
                          "但默认CUDA反向计算含有非确定性操作，两次下一步计算的损失相同，参数更新仍存在微小差异。"
                          f"本次记录的更新最大绝对差为{difference}。独立的确定性参考实现中，下一步损失、梯度、参数更新及优化器状态完全一致。"
                          "因此这里确认的是状态恢复正确，不宣称默认CUDA训练能够逐位复现；正式训练没有被替换成该参考实现。")
    doc.add_paragraph("1024运行验收中的Face反传只验证主生成损失路径，未验证1024的VAE解码、身份及感知辅助损失反传，"
                      "不能据此宣称已经完成1024面部训练验收。")
    # Include any additional machine-readable acceptance result without turning
    # its existence alone into a 'passed' claim.
    extras = []
    for name, result in acceptances.items():
        if name in ("data_runtime", "preservation_results", "report_manifest", "resume_summary",
                    "resume_diagnosis", "resume_deterministic_reference", "runtime_512", "runtime_1024",
                    "weight_inventory", "face_resume_deterministic_reference"):
            continue
        detail = {k: v for k, v in result.items() if isinstance(v, (bool, int, float, str))}
        extras.append((name, json.dumps(detail, ensure_ascii=False)[:900] or "详见原始验收记录；未自动推断通过"))
    if extras:
        add_table(doc, ["补充验收记录", "记录中的结果"], extras, font_size=8)
    training_rows = []
    for branch, value in branches.items():
        if value is None:
            training_rows.append((branch.title(), "未完成", "未记录", "未记录", "未记录", "未记录"))
            continue
        before = value.get("probe_before", {}).get("loss")
        after = value.get("probe_after", {}).get("loss")
        training_rows.append((branch.title(), value.get("optimizer_steps", "未记录"),
                              f"{value.get('unique_training_pairs', '未记录')}对／{value.get('training_pair_uses', '未记录')}次",
                              f"{number(before)} → {number(after)}", number(value.get("seconds"), 1), number(value.get("peak_cuda_gib"), 2)))
    add_table(doc, ["阶段", "优化步数", "不同样本／使用次数", "固定探针损失：前→后", "阶段用时/秒", "峰值/GiB"], training_rows, font_size=8)
    doc.add_paragraph("这里的训练步数按优化器调用次数记录；门控为零时可能没有实际更新，实际有梯度的次数在下文另列。"
                      "它不是图片数量，也不是生成图片时的30步采样。"
                      "固定探针是在相同样本、噪声和时间点下检查损失，用于观察是否发生学习；数值下降不能直接证明脸更像或泛化更好。"
                      "本轮只有Body 100步及Face 100步，不是完整iPER训练，不是新的20k模型。")
    if branches.get('body') and branches.get('face'):
        body, face = branches['body'], branches['face']
        doc.add_paragraph(f"日志实际记录：Body的100次迭代中有{body.get('nonzero_gradient_steps')}次非零梯度，"
                          f"Face的100次迭代中有{face.get('nonzero_gradient_steps')}次非零梯度。"
                          "Face有一次位于旧强度调度的边界附近，BF16下门控为零，这一次没有实际参数更新。"
                          "Body固定诊断损失下降；Face固定诊断损失略升，不能据此声称Face训练已经改善质量。"
                          "Face的身份与感知辅助损失实际计算了25次，但固定诊断值本身不包含这两项。")
    doc.add_heading("五、完整采样的对照结果", level=1)
    doc.add_paragraph("统一30步采样、随机种子42和同一提示词。下表的初始化、Body和Body＋Face都采用同一Qwen基座；"
                      "本轮没有与旧DeepGen或此前简化Qwen做等训练量的正式排名。")
    add_metrics(doc, evals[0], "512×512：4对固定验证样本")
    add_metrics(doc, evals[1], "1024×1024：前2对验证样本的运行检查")
    doc.add_paragraph("1024结果使用上述512短训权重直接推理；它证明的是该尺寸能否运行，不能写成已经完成1024训练。"
                      "跨人A/B没有正确的像素级目标答案，因此不拿B的照片计算重建质量分数。")
    doc.add_heading("六、目前可以和不能得出的结论", level=1)
    if review:
        doc.add_paragraph("逐图观察已经发现Face局部画质问题；具体原因尚未确认，不能简单归因为训练步数不足。"
                          "继续正式训练前，应先检查面部残差数值、注入位置，以及训练噪声范围与实际采样调度的对应关系。"
                          "本轮保留100步终点和全部失败图片，没有调参挑图或延长训练来掩盖问题。")
        doc.add_paragraph("跨人案例中仍能看出A的服装，但受损的脸无法可靠判断是否保持身份，因此串脸检查结论为“无法判断”。"
                          "另有一条可追溯的排查线索：按原计划保留的推理步序调度，会在少数训练未覆盖的噪声水平启用Face。"
                          "这只是已观察到的训练与推理覆盖差异，尚不能认定它就是彩斑的原因。")
    doc.add_paragraph("目前能够根据哈希、结构测试、冻结检查和完整采样记录，判断迁移是否保留旧模块、训练链路是否正常。"
                      "能否更好地保持人脸，需要对可靠可见脸进行身份评分并逐图核查。当前没有ArcFace/PCK结果，且短训只有极少样本，无法判断完整方案是否优于此前版本。"
                      "图片中脸太小、转头、背面或模糊时，应标为无法判断，不能用放大显示替代可靠细节。")
    if missing:
        doc.add_paragraph("尚未完成的记录：" + "；".join(missing))
    doc.add_paragraph("完整文件哈希、加载参数清单、逐样本指标和检测处理记录均保留为JSON；报告只摘取可读结论。"
                      "本报告生成时再次核对了文件哈希，详细结果记录于report_manifest.json。")
    doc.add_heading("原模块文件清单", level=2)
    by_dir = defaultdict(list)
    for row in copied:
        path = Path(row["name"])
        by_dir[str(path.parent)].append(path.name)
    add_table(doc, ["原目录", "完整保留的文件"], sorted((name, "、".join(files)) for name, files in by_dir.items()), font_size=8)

    panels = []
    caches = {}
    for item in evals:
        for index in range(item["count"]):
            panel = make_panel(root, item, index, root / "report_panels", caches)
            if panel is not None:
                panels.append((item, panel))
    if panels:
        section = doc.add_section(WD_SECTION.NEW_PAGE)
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = Inches(11.69), Inches(8.27)
        section.left_margin = section.right_margin = Inches(.4)
        section.top_margin = section.bottom_margin = Inches(.4)
        titles = {"validation_512": "512验证", "validation_1024": "1024运行检查", "cross_person_512": "不同人物A/B检查"}
        for ordinal, (item, panel) in enumerate(panels):
            if ordinal:
                doc.add_page_break()
            meta = panel["metadata"]
            doc.add_heading(f"{titles[item['name']]} · 样本{panel['pair_index']:03d}", level=1)
            doc.add_paragraph(f"A：{meta['source_appearance']}／{meta['source_stem']}；B：{meta['target_appearance']}／{meta['target_stem']}。"
                              f"生成分辨率：{panel['resolution']}×{panel['resolution']}。")
            doc.add_picture(str(panel["image"]), width=Inches(10.85))
            doc.add_paragraph("第一排为整图，第二排为对应脸部区域，放大到与整图单栏相同宽度。"
                              "A、B各用自己的脸框；三张生成图统一使用包含A、B头部位置的窗口，避免动作未跟随后只截到半张脸。"
                              "没有按效果逐张挑选裁剪位置。"
                              "PNG按原始像素保存并嵌入，放大本身不会增加真实细节。看不清的面部无法判断。", style="Caption")
            if item["name"] == "cross_person_512":
                doc.add_paragraph("此处B是另一个人，只用于提供动作与几何。该例未计算可靠人脸相似度，不能仅凭这两张案例宣称已排除串脸。")
    args.out.parent.mkdir(parents=True, exist_ok=True)
    temporary = args.out.with_name(args.out.stem + ".tmp.docx")
    doc.save(temporary)
    temporary.replace(args.out)
    report = {"report": str(args.out), "created_shanghai": now, "complete_records": not missing,
              "missing": missing, "source_file_audit": copied,
              "panels": [{**panel, "image": str(panel["image"]), "group": item["name"]} for item, panel in panels],
              "identity_metric_computed": False, "pose_pck_computed": False,
              "resume_source": "acceptance/resume_summary.json", "default_cuda_bitwise_reproducibility_claimed": False,
              "required_runtime_acceptance": ["runtime_512", "runtime_1024"],
              "face_supervision_source": "face_supervision.json"}
    (root / "report_manifest.json").write_text(json.dumps(report, ensure_ascii=False, indent=2))
    print(json.dumps({"report": str(args.out), "panels": len(panels), "complete_records": not missing,
                      "missing_count": len(missing)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
