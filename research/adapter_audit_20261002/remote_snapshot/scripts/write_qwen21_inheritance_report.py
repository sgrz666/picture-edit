#!/usr/bin/env python3
"""Embed measured tables and original-pixel comparison panels in a Word report."""
from __future__ import annotations

from datetime import datetime
import hashlib
import json
import math
from pathlib import Path
import sys
from zoneinfo import ZoneInfo

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from scripts.run_qwen21_weight_inheritance import OUT, RUN_ROOT, CHECKPOINTS
from scripts.write_qwen21_full_report import configure_doc, add_table, crop_face, image_rgb, verify_copy
from scripts.audit_qwen21_weight_compatibility import file_sha

REPORT = RUN_ROOT.parent / 'docs/Qwen完整版_继承Body权重与直接出图比较.docx'
LABELS = {'baseline': '当前完整版', 'body512': '继承512来源', 'body1024': '继承1024来源'}
FONT = '/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc'
FIELDS = [('姿态 PCK@0.02 ↑','pose_pck_002',True),('姿态 NME ↓','pose_nme_diagonal',False),
          ('PSNR ↑','psnr',True),('SSIM ↑','ssim',True),('人体 MAE ↓','human_mae',False),
          ('背景 MAE（相对A）↓','background_mae_source',False),
          ('身份相似度（原始）↑','identity_cosine_raw',True),
          ('身份相似度（可靠共同子集）↑','identity_cosine_common',True),
          ('单图采样耗时／秒 ↓','seconds_mean',False),('峰值显存／GiB ↓','peak_cuda_gib',False)]


def make_panel(rows, set_name, resolution, index, mode):
    selected = [r for r in rows if (r['set'],r['resolution'],r['pair_index'],r['mode']) == (set_name,resolution,index,mode)]
    by_candidate = {r['candidate']:r for r in selected}
    if set(by_candidate) != set(CHECKPOINTS):
        raise ValueError('A panel is missing a comparison candidate')
    anchor = by_candidate['baseline']
    tile = resolution
    gap, top, middle = tile//64, tile//10, tile//14
    canvas = Image.new('RGB',(5*tile+4*gap,2*tile+top+middle),'white')
    draw = ImageDraw.Draw(canvas)
    font = ImageFont.truetype(FONT,max(20,tile//25))
    small = ImageFont.truetype(FONT,max(17,tile//34))
    pictures = [image_rgb(anchor['source']), image_rgb(anchor['target']),
                *[image_rgb(by_candidate[c]['image']) for c in CHECKPOINTS]]
    boxes = [anchor['face_boxes_normalized']['source'],anchor['face_boxes_normalized']['target']]
    crop_notes = ['A脸窗口','B脸窗口']
    for candidate in CHECKPOINTS:
        detections = by_candidate[candidate]['face_detection']['generated']
        selected_index = detections['reliable_index']
        if selected_index is None: selected_index = detections['raw_largest_index']
        if selected_index is not None:
            box = [v/512 for v in detections['detections'][selected_index]['box']]
            note = '检出脸放大' if detections['reliable'] else '检测框放大（待核查）'
        else:
            box, note = anchor['face_boxes_normalized']['target'], 'B头部窗口（未检出脸）'
        boxes.append(box)
        crop_notes.append(note)
    for col, (picture, label, box, note) in enumerate(zip(pictures,['A：身份与背景','B：目标动作',*[LABELS[c] for c in CHECKPOINTS]],boxes,crop_notes)):
        left = col*(tile+gap)
        draw.text((left+8,8),label,font=font,fill='#17273c')
        canvas.paste(picture.resize((tile,tile),Image.Resampling.LANCZOS),(left,top))
        crop = crop_face(picture,box)
        canvas.paste(crop.resize((tile,tile),Image.Resampling.LANCZOS),(left,top+tile+middle))
        draw.text((left+8,top+tile+3),note,font=small,fill='#425466')
    directory = OUT / 'panels'
    directory.mkdir(exist_ok=True)
    path = directory / f'{set_name}_{resolution}_{index:03d}_{mode}.png'
    canvas.save(path,optimize=True)
    return {'set':set_name,'resolution':resolution,'pair_index':index,'mode':mode,'path':str(path),
            'size':list(canvas.size),'boxes':boxes,'crop_notes':crop_notes,
            'image_tile_original_resolution':resolution,'no_super_resolution':True}


def comparison_table(doc, summary, resolution, mode):
    selected = [summary[f'validation_{resolution}'][f'{c}_{mode}'] for c in CHECKPOINTS]
    doc.add_heading(f'{resolution}生成｜'+('Body Adapter' if mode=='body' else 'Body＋Face Adapter'),level=2)
    table = add_table(doc,['指标',*[LABELS[c] for c in CHECKPOINTS]],[])
    for label, field, higher in FIELDS:
        values=[]
        for report in selected:
            entry=report[field]
            values.append((entry['mean'],entry['n']) if isinstance(entry,dict) else (entry,None))
        valid=[v for v,n in values if v is not None and math.isfinite(v)]
        best=(max(valid) if higher else min(valid)) if valid else None
        cells=table.add_row().cells
        cells[0].text=label
        for cell,(value,n) in zip(cells[1:],values):
            text=('无法判断' if value is None else f'{value:.4f}' if field not in ('seconds_mean','peak_cuda_gib','human_mae','background_mae_source','psnr') else f'{value:.2f}')
            if n is not None: text+=f'（n={n}）'
            run=cell.paragraphs[0].add_run(text)
            run.bold=best is not None and value is not None and math.isclose(value,best,rel_tol=1e-10,abs_tol=1e-10)
    for field,label in [('face_detected_n','生成脸检出数'),('reliable_face_n','可靠检测数')]:
        cells=table.add_row().cells
        cells[0].text=label
        for cell,report in zip(cells[1:],selected): cell.text=f'{report[field]}/4'
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for run in p.runs: run.font.size=Pt(9)


def main():
    rows=json.loads((OUT/'evaluated_samples.json').read_text())
    summary=json.loads((OUT/'metric_summary.json').read_text())
    runtime=json.loads((OUT/'runtime_checks.json').read_text())
    generation=json.loads((OUT/'generation_summary.json').read_text())
    copies=verify_copy(json.loads((RUN_ROOT/'migration_manifest.json').read_text()))
    if not generation['passed'] or len(rows)!=60 or len(runtime)!=6 or not all(x['passed'] for x in runtime.values()) or not all(x['matches_snapshot'] for x in copies):
        raise ValueError('Acceptance evidence is incomplete or architecture changed')
    (OUT/'architecture_preservation.json').write_text(json.dumps(copies,ensure_ascii=False,indent=2))
    panel_manifest=OUT/'panel_manifest.json'
    if panel_manifest.exists():
        panels=json.loads(panel_manifest.read_text())
        expected={(name,res,index,mode) for name,res,count in [('validation',512,4),('validation',1024,4),('cross_person',1024,2)] for index in range(count) for mode in ('body','body_face')}
        if len(panels)!=20 or {(p['set'],p['resolution'],p['pair_index'],p['mode']) for p in panels}!=expected:
            raise ValueError('Existing panel manifest differs from the fixed comparison')
        for row in rows:
            if file_sha(Path(row['image']))!=row['image_sha256']:
                raise ValueError('Generation changed since the panels were built')
        for panel in panels:
            with Image.open(panel['path']) as picture:
                if list(picture.size)!=panel['size']:
                    raise ValueError('Panel dimension changed')
    else:
        panels=[]
        for set_name,resolution,count in [('validation',512,4),('validation',1024,4),('cross_person',1024,2)]:
            for index in range(count):
                for mode in ('body','body_face'): panels.append(make_panel(rows,set_name,resolution,index,mode))
        panel_manifest.write_text(json.dumps(panels,ensure_ascii=False,indent=2))
    doc=Document()
    configure_doc(doc)
    for sec in doc.sections: sec.footer.paragraphs[0].text='Qwen完整版｜Body权重继承与直接出图比较'
    doc.add_heading('Qwen完整版：继承Body权重后的直接出图比较',0)
    doc.add_paragraph(datetime.now(ZoneInfo('Asia/Shanghai')).strftime('%Y年%m月%d日'))
    doc.add_heading('这次实际做了什么',1)
    doc.add_paragraph('Body Adapter负责把B的动作几何传给Qwen，让A中的人物改变姿态；Face Adapter负责结合A的身份和B的头部几何，保持人物的脸。两者都属于附加控制模块，Qwen本体使用现有预训练权重。')
    doc.add_paragraph('这次建立了两个完整架构候选：分别从过去的512、1024简化版Body检查点中，继承已经核实兼容且有训练证据的共有模块参数。这里的“512来源”和“1024来源”指权重来源；两个候选都分别生成了512和1024图片。')
    doc.add_paragraph('本轮没有训练，也没有使用最终50对测试样本。每个候选都以当前完整版为起点，只替换白名单允许的部分；完整控制核心、六组输出、Face Adapter和其他分支沿用当前完整版权重。Face彩斑是否变化单独记录。')
    doc.add_paragraph('本次观察：继承正确完成，但直接出图尚未获得稳定、全面的质量改善。两种来源的部分姿态指标略有提高，整体动作仍没有准确达到B；开启完整Face Adapter的30张结果均能看到彩斑或局部破损。1024的Body单独出图脸部更清楚，但这不等于继承来源更好。两个候选继续保留。')
    doc.add_paragraph('完整版控制核心及连接Qwen的接口仍沿用先前Body／Face各100步的短训结果，本次没有对继承后的组合继续训练。因此，这轮回答的是“把兼容参数搬过来后，能否直接获得更好的结果”；后续适应训练能否改善效果尚未验证。')
    add_table(doc,['显示名称','实际含义'],[
        ['当前完整版','Body、Face均来自既有完整版短训检查点'],
        ['继承512来源','用简化版512 Body 20k的兼容参数替换共有部分'],
        ['继承1024来源','用简化版1024 Body 3k的兼容参数替换共有部分'],
        ['Body','只开启Body Adapter'],['Body＋Face','同时开启完整Body和完整Face Adapter']])
    doc.add_heading('实验范围与读表方式',1)
    doc.add_paragraph('4对固定验证样本分别生成两种尺寸，并补充2对A/B不同人的1024案例。三个候选各有Body与Body＋Face两组，共60张完整结果。验证样本用于这次观察，不能据此断言完整数据集上的画质或泛化能力。')
    doc.add_paragraph('统一30步、种子42、true CFG 1.0。原生参考只输入完整A图；Face的身份和纹理来自A，B只给动作与头部几何。提示词保持为：Change the person pose to match the target pose. Preserve identity, clothing, and background.')
    doc.add_paragraph('所有姿态和像素指标统一在512网格评测；背景误差相对A计算。身份指标同时列出原始检测分数与六组都能可靠检测的共同子集。n是该数值实际采用的样本数；n=0时写“无法判断”。检出人脸只是检测成功，不代表脸已经好看或身份正确。每行加粗仅表示本表的最优数值，不代表显著优势。')
    reference_scores=[r['metrics']['source_target_identity_cosine_raw'] for r in rows if r['set']=='validation' and r['resolution']==512 and r['candidate']=='baseline' and r['mode']=='body']
    doc.add_paragraph(f'这4对同人A/B的原始身份相似度平均为{sum(reference_scores)/len(reference_scores):.4f}（n=4），用于理解真实动作和头向改变对指标的影响。可靠共同子集仅指检测、尺寸和位置条件通过；彩斑及低画质仍作为失败结果保留，不为提高平均分剔除。')
    for resolution in (512,1024):
        for mode in ('body','body_face'): comparison_table(doc,summary,resolution,mode)
    doc.add_heading('数值变化及其实际含义',1)
    def value(res,c,mode,field):
        e=summary[f'validation_{res}'][f'{c}_{mode}'][field]
        return e['mean'] if isinstance(e,dict) else e
    for res in (512,1024):
        baseline=value(res,'baseline','body','pose_pck_002')
        changes=[]
        for c in ('body512','body1024'):
            changes.append(f'{LABELS[c]}的Body姿态PCK为{value(res,c,"body","pose_pck_002"):.4f}，相对当前完整版变化{value(res,c,"body","pose_pck_002")-baseline:+.4f}；背景MAE变化{value(res,c,"body","background_mae_source")-value(res,"baseline","body","background_mae_source"):+.2f}')
        doc.add_paragraph(f'{res}生成：当前完整版Body的PCK为{baseline:.4f}。'+'。'.join(changes)+'。这些变化只描述4对案例。')
        gains=[]
        for c in CHECKPOINTS:
            delta=value(res,c,'body_face','pose_pck_002')-value(res,c,'body','pose_pck_002')
            bg=value(res,c,'body_face','background_mae_source')-value(res,c,'body','background_mae_source')
            before=value(res,c,'body','identity_cosine_common')
            after=value(res,c,'body_face','identity_cosine_common')
            identity_note='共同脸子集无法比较身份' if before is None or after is None else f'共同子集身份分数变化{after-before:+.4f}'
            gains.append(f'{LABELS[c]}开启Face后，PCK变化{delta:+.4f}、背景MAE变化{bg:+.2f}，{identity_note}')
        doc.add_paragraph('；'.join(gains)+'。Face带来的局部观感仍须结合下方放大图和人工记录。')
    for c in CHECKPOINTS:
        for mode in ('body','body_face'):
            pck_delta=value(1024,c,mode,'pose_pck_002')-value(512,c,mode,'pose_pck_002')
            t512,t1024=value(512,c,mode,'seconds_mean'),value(1024,c,mode,'seconds_mean')
            doc.add_paragraph(f'{LABELS[c]}／'+('Body' if mode=='body' else 'Body＋Face')+f'：512→1024的PCK变化{pck_delta:+.4f}；平均采样耗时从{t512:.2f}秒增至{t1024:.2f}秒（约{t1024/t512:.1f}倍）。原尺寸放大图用于判断脸部细节，指标仍统一缩到512计算。')
    manual_path=OUT/'visual_review.json'
    manual=json.loads(manual_path.read_text()) if manual_path.exists() else {}
    if manual:
        counts={}
        for candidate in CHECKPOINTS:
            for mode in ('body','body_face'):
                reviews=[v for case in manual.values() for v in case.get('variants',[]) if (v['candidate'],v['mode'])==(candidate,mode)]
                counts[f'{candidate}_{mode}']=sum(v.get('face_artifact',False) for v in reviews)
        add_table(doc,['方案','人工查看的彩斑／局部脸部破损案例数'],[
            [f'{LABELS[c]}｜'+('Body' if m=='body' else 'Body＋Face'),f'{counts[f"{c}_{m}"]}/10'] for c in CHECKPOINTS for m in ('body','body_face')])
        doc.add_paragraph('人工查看包括全部60张生成图、各案例A/B和放大脸部。彩斑、目标动作与头向，以及身份是否能可靠判断，均保留逐案例记录。彩斑数量只是这批案例的描述，不是通用检测指标。')
    else:
        doc.add_paragraph('人工面部审阅尚未完成；本版不能作为最终结果报告。')
    common=summary['validation_1024']['baseline_body']['identity_cosine_common']['n']
    doc.add_paragraph(f'1024比较的可靠共同身份子集只有{common}对；'+('样本不足，无法可靠比较身份保持度。' if common<2 else '人数很少，身份分数只能作为这批案例的辅助观察。'))
    doc.add_paragraph('两个候选继续保留，不自动替换默认权重。继承只改变了共有几何理解部分，不能认为简化版整套Body或Face能力已搬入完整架构，也不能把直接出图结果当作适应训练完成后的效果。')
    doc.add_heading('不同人A/B的检查',1)
    cross=[r for r in rows if r['set']=='cross_person']
    add_table(doc,['案例／方案','与A的身份分数','与B的身份分数','人工观察'],[
        [f'{r["pair_index"]+1}／{LABELS[r["candidate"]]}／'+('Body' if r['mode']=='body' else 'Body＋Face'),
         f'{r["metrics"]["identity_cosine_raw"]:.4f}',f'{r["metrics"]["identity_target_cosine_raw"]:.4f}',
         '脸部破损，无法可靠判断' if r['mode']=='body_face' else '更像A，未见明显变成B的人脸'] for r in cross])
    doc.add_paragraph('两对不同人案例的严格姿态PCK@0.02均为0，人工查看也没有达到B双手位于胸前的动作。原始身份分数更接近A，但Face开启后的破损使人工身份判断不可靠，不能宣称已通过完整串脸验收。与B的识别分数仅用于离线检查，B的身份图像和特征没有进入生成条件。')
    doc.add_paragraph('人物仍像A却没有完成B的动作，不能算姿态编辑成功。身份分数必须与目标动作完成情况一起看。')
    doc.add_heading('正确性与来源记录',1)
    add_table(doc,['检查项','结果'],[
        ['白名单加载','两份各208个参数张量；三个嵌入表仅复制第0行'],
        ['参数保护与重新加载','白名单外参数、未选中行及完整参数结构保持一致'],
        ['CPU共有模块对照','两份各8组固定验证输入，覆盖512/1024，输出一致'],
        ['GPU运行检查','三个候选×两种分辨率，共6组通过'],
        ['控制边界','关闭全部控制与原生预测一致；注入不改前缀；Face区外残差为零；无效脸回退'],
        ['架构保留','47个旧架构文件仍与迁移快照一致'],
        ['训练与默认设置','新增优化步数0；无优化器继承；默认权重保持原配置']])
    doc.add_paragraph('候选检查点是新的初始化文件：可以推理或作为后续训练的起点；没有旧优化器状态，不能把来源20k／3k当作新版续训步数。完整架构仍基于512短训权重；来自1024的共有参数不等于整套完整模型已经在1024训练。')
    doc.add_paragraph('完整参数映射、检查点校验值、逐图指标、共同子集原因和全部20张并排图保存在本次独立实验目录。下列展示固定采用前两对验证样本及两对不同人案例，涵盖两个Adapter开关；图片直接嵌入Word，未做超分辨率增强。')
    section=doc.add_section(WD_SECTION.NEW_PAGE)
    section.orientation=WD_ORIENT.LANDSCAPE
    section.page_width,section.page_height=Inches(11.69),Inches(8.27)
    section.left_margin=section.right_margin=Inches(.4)
    section.top_margin=section.bottom_margin=Inches(.4)
    shown=[p for p in panels if p['set']=='cross_person' or p['pair_index']<2]
    for index,panel in enumerate(shown):
        if index: doc.add_page_break()
        doc.add_heading(('不同人A/B' if panel['set']=='cross_person' else '固定验证')+f'｜案例{panel["pair_index"]+1}｜{panel["resolution"]}｜'+('Body Adapter' if panel['mode']=='body' else 'Body＋Face Adapter'),1)
        doc.add_paragraph('第一行是整图，第二行是对应脸部放大图。A提供身份，B展示目标动作；生成图未检出脸时，第二行展示B头部几何窗口并明确标注。')
        doc.add_picture(panel['path'],width=Inches(10.85))
    REPORT.parent.mkdir(parents=True,exist_ok=True)
    doc.save(REPORT)
    result={'report':str(REPORT),'sha256':file_sha(REPORT),'tables':len(doc.tables),'embedded_panels':len(shown),
            'all_panels':len(panels),'images_evaluated':60,'manual_review_present':bool(manual),
            'training_performed':False,'final_test_used':False,'architecture_preserved':True}
    (OUT/'report_manifest.json').write_text(json.dumps(result,ensure_ascii=False,indent=2))
    print(json.dumps(result,ensure_ascii=False))


if __name__=='__main__':
    main()
