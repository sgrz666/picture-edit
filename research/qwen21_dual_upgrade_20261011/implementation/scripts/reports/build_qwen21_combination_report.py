"""Revision of the Word value audit with the locked native combination grid."""
from __future__ import annotations

from collections import defaultdict
import hashlib
import json
from pathlib import Path
import zipfile

from docx import Document
from docx.shared import Cm
from docx.oxml import OxmlElement
from PIL import Image

from scripts.eval.qwen21_condition_combinations import (COMBINATION_METHODS, COMBINATION_LABELS,
    require_protocol_outputs, factor_edges)
from scripts.eval.score_qwen21_native_value import read_jsonl, save, sha, aggregate, paired_delta
from scripts.reports import build_qwen21_native_value_report as old


def quality_table(doc,nodes,methods,title,shared_n):
    old.para(doc,title,11,True)
    columns=(('pose_pck_002','动作PCK↑',3,True),('pose_nme_diagonal','动作NME↓',4,False),
             ('identity_a_cosine_formal_common','共同脸相似度↑',3,True),
             ('background_mae_source','背景误差↓',2,False),('seconds','秒/张↓',1,False),
             ('peak_cuda_gib','GiB↓',1,False))
    table=doc.add_table(rows=1,cols=7);table.style='Table Grid'
    for c,text in zip(table.rows[0].cells,['路线',*[x[1] for x in columns]]):old.font(c.paragraphs[0].add_run(text),8,True)
    best={}
    for metric,_,_,up in columns:
        vals=[old.summary_value(nodes[m],metric) for m in methods]
        vals=[v for v in vals if v is not None]
        best[metric]=(max(vals) if up else min(vals)) if vals else None
    for method in methods:
        cells=table.add_row().cells;old.font(cells[0].paragraphs[0].add_run(old.CN[method]),8)
        for cell,(metric,_,digits,_) in zip(cells[1:],columns):
            val=old.summary_value(nodes[method],metric)
            bold=val is not None and val==best[metric] and (metric!='identity_a_cosine_formal_common' or shared_n>=3)
            old.font(cell.paragraphs[0].add_run(old.number(val,digits)),8,bold)
    old.para(doc,f'每路线{nodes[methods[0]]["n_images"]}对；共同可靠脸 n={shared_n}。身份只比较这些路线都可评的同一批脸，少于3对不作稳定判断，也不加粗排名。其他加粗为同列最优，不代表整体最好。',8)


def shared_nodes(rows,methods):
    by_case=defaultdict(list)
    for r in rows:
        if r['method'] in methods:by_case[r['pair_id']].append({**r})
    for variants in by_case.values():
        valid={r['method'] for r in variants}==set(methods) and all(r['formal_reliable_identity_sample'] for r in variants)
        for r in variants:r['formal_common_reliable_identity_sample']=valid
    nodes={m:aggregate([r for vs in by_case.values() for r in vs if r['method']==m]) for m in methods}
    return nodes,nodes[methods[0]]['identity_a_cosine_formal_common']['n']


def increment_table(doc,methods,records,title):
    old.para(doc,title,10,True)
    table=doc.add_table(rows=1,cols=6);table.style='Table Grid'
    for c,text in zip(table.rows[0].cells,['新增路线','动作变化','A脸变化','共同脸n','背景变化','耗时变化/秒']):old.font(c.paragraphs[0].add_run(text),8,True)
    for m in methods:
        record=records[f'{m}_minus_native_skeleton'];identity=record['identity_a_cosine_reliable']
        values=[old.CN[m],old.number(record['pose_pck_002']['mean'],4),old.number(identity['mean'],4),str(identity['n']),
                old.number(record['background_mae_source']['mean'],2),old.number(record['seconds']['mean'],1)]
        for c,text in zip(table.add_row().cells,values):old.font(c.paragraphs[0].add_run(text),8)
    old.para(doc,'均为该路线减去A＋骨架；动作、身份正值是提高，背景、耗时正值是增加。身份使用每一对方案共同可评的脸，n可能不同，不能横向只按身份变化排名；n<3只记录观察。',8)


def factor_table(doc,records,title):
    old.para(doc,title,10,True)
    t=doc.add_table(rows=1,cols=6);t.style='Table Grid'
    for c,text in zip(t.rows[0].cells,['增加条件','原路线→新路线','动作变化','A脸变化','共同脸n','背景变化']):old.font(c.paragraphs[0].add_run(text),8,True)
    names={'B':'B照片','B_normal':'法线','B_depth':'深度'}
    for role,a,b in factor_edges():
        rec=records[f'{b}_minus_{a}']['metrics'];idn=rec['identity_a_cosine_reliable']
        label=old.CN[a].replace('原生 A＋','')+' → '+old.CN[b].replace('原生 A＋','')
        vals=[names[role],label,old.number(rec['pose_pck_002']['mean'],4),old.number(idn['mean'],4),str(idn['n']),
              old.number(rec['background_mae_source']['mean'],2)]
        for c,text in zip(t.add_row().cells,vals):old.font(c.paragraphs[0].add_run(text),7.5)
    old.para(doc,'每行只增加一种条件。身份是该行前后共同可靠脸的变化，n<3不能得出稳定身份结论；不同背景条件的结果分别列出，不把法线、深度混成一种条件。',8)


def visual_table(doc,rows,review,methods,title):
    old.para(doc,title,10,True)
    flags=(('geometry_color_transfer','几何染色'),('skeleton_overlay','骨架线覆盖'),
           ('extra_person','多余人物/肢体'),('face_color_artifact','其他脸部色块'))
    table=doc.add_table(rows=1,cols=5);table.style='Table Grid'
    for cell,label in zip(table.rows[0].cells,['路线',*[label for _,label in flags]]):
        old.font(cell.paragraphs[0].add_run(label),8,True)
    for method in methods:
        variants=[r for r in rows if r['method']==method]
        counts=[sum(bool(({flag,'extra_limbs'} if flag=='extra_person' else {flag}) &
                         set(review[r['image']].get('visual_flags',[]))) for r in variants) for flag,_ in flags]
        vals=[old.CN[method],*[str(n)+'/'+str(len(variants)) for n in counts]]
        for cell,label in zip(table.add_row().cells,vals):old.font(cell.paragraphs[0].add_run(label),8)
    old.para(doc,'逐图人工记录，不是自动检测分数；同一张图可同时出现多种异常。0只表示未记录该类明显异常，不代表编辑已完全成功。异常图全部保留并参与原有指标计算。',8)


def compress_figure(fig):
    p=Path(fig['path']);tmp=p.with_suffix('.optimized.png')
    with Image.open(p) as im:im.save(tmp,optimize=True,compress_level=9)
    tmp.replace(p);fig['sha256']=sha(p)


def build_report(args):
    root=args.root;parent=root.parent
    rows=json.loads((root/'evaluation/evaluated_samples.json').read_text())
    expected=require_protocol_outputs(args.protocol,rows)
    summary=json.loads((root/'evaluation/metric_summary.json').read_text())
    review=json.loads((root/'evaluation/face_review.json').read_text())
    if not summary['formal_identity_review_complete']:raise ValueError('Supplementary face review is incomplete')
    obs=json.loads((root/'final_observations.json').read_text())
    pairs=read_jsonl(root/'data/pairs.jsonl');by_id={p['pair_id']:p for p in pairs}
    selection=json.loads((root/'data/report_selection.json').read_text())
    ids=selection['pair_ids'];high=selection['highres_pair_ids']
    old.CN.update(COMBINATION_LABELS);old.CN['native_geometry']='原生 A＋法线＋深度'
    figures=[]
    for pid in ids:
        for methods,suffix in [(COMBINATION_METHODS[:4],'native_combinations_1'),(COMBINATION_METHODS[4:],'native_combinations_2'),
                              (('native_ab','native_geometry','simple_body','full_body'),'legacy_pose'),(old.FACE_FIGURE,'legacy_face')]:
            f=old.figure(root,rows,by_id[pid],512,methods,suffix)
            if f is None:raise ValueError(f'Missing figure inputs: {pid}/{suffix}')
            figures.append(f)
    for pid in high:
        for methods,suffix in [(COMBINATION_METHODS[:4],'native_combinations_1'),(COMBINATION_METHODS[4:],'native_combinations_2'),
                              (('native_ab','native_geometry','full_body','full_body_face'),'legacy_1024')]:
            f=old.figure(root,rows,by_id[pid],1024,methods,suffix)
            if f is None:raise ValueError(f'Missing figure inputs: {pid}/{suffix}')
            figures.append(f)
    figures.append(old.sampling_steps_figure(root,rows,by_id['iper_00']))
    if len(figures)!=37 or any(f is None for f in figures):raise ValueError('Expected 37 complete report figures')
    for f in figures:compress_figure(f)
    doc=Document();sec=doc.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7)
    sec.left_margin=sec.right_margin=Cm(1.3);sec.top_margin=sec.bottom_margin=Cm(1.1)
    old.para(doc,'Qwen基座与控制模块：价值核对（组合补测）',17,True)
    old.para(doc,'2026年10月8日｜只做推理，不训练｜512完整比较＋1024固定复核',9)
    old.para(doc,obs['headline'],11,True)
    old.para(doc,'为什么补测',12,True)
    old.para(doc,'上一轮原生A＋骨架已表现较强，但没有充分测试把照片、骨架、法线和深度一起交给基座。本轮以A＋骨架为起点，分别加入B照片、法线、深度，补齐八种组合；法线与深度分别测试，不再只看两者打包的表现。')
    old.para(doc,'怎样理解这些输入',12,True)
    old.para(doc,'A决定人物身份、服装和场景；B照片提供目标动作细节；骨架提供关节布局；法线的颜色表示表面朝向，深度的灰度表示远近，二者都不是要生成的肤色。Body Adapter是训练过的动作控制模块，Face Adapter是面部控制模块。Adapter各列复用上一轮结果，本轮只新增原生组合，没有训练，也没有加载LoRA。')
    old.para(doc,'实际完成量',12,True)
    old.para(doc,'512沿用12对（iPER、TikTok各6对），1024沿用原4对（各2对）。八种组合共128张：16张A＋骨架复用、112张新增。连同旧路线与40步附录，正式索引共252张。另8张五图开发检查独立记录。全部新图30步、种子42；详细动作和保留要求不变，只补充图片角色说明。')
    old.para(doc,'重点结论',12,True)
    for text in obs['conclusions']:old.para(doc,text)
    old.para(doc,'12对只覆盖6个A人物；1024复核只覆盖2个A人物。这里比较既有条件的输入组合，不排查或修正条件图估计。原生组合与Adapter的差别还包含训练数据和训练量；结果只能证明这些案例上的表现，不能证明普遍最优。',9)
    for resolution in (512,1024):
        doc.add_page_break();old.para(doc,f'八种原生组合｜{resolution}、30步',13,True)
        for dataset in ('iper','tiktok'):
            group=f'{dataset}_{resolution}_steps30';nodes=summary['native_combination_groups'][group]
            n=nodes['native_skeleton']['identity_a_cosine_formal_common']['n']
            quality_table(doc,nodes,COMBINATION_METHODS,f'{dataset}｜{6 if resolution==512 else 2}对',n)
            increment_table(doc,COMBINATION_METHODS[1:],summary['paired_differences'][group],'相对A＋骨架的变化')
        old.para(doc,'PCK越高表示关节位置越贴近B，NME越低表示位置误差越小；两者不能单独评价动作自然、三维方向或人脸。身份用实际检测脸，背面、误检、遮挡及过小脸说明原因并保留原始记录；不会因低分剔除。背景误差与A比较，B不是理想生成图，不计算对B的PSNR、SSIM或人体重建误差。',9)
    doc.add_page_break();old.para(doc,'动作分数之外：明显异常出现了多少次',13,True)
    for resolution in (512,1024):
        for dataset in ('iper','tiktok'):
            subset=[r for r in rows if r['dataset']==dataset and r['resolution']==resolution and r['steps']==30]
            visual_table(doc,subset,review,COMBINATION_METHODS,f'{dataset}｜{resolution}、30步')
    old.para(doc,'出现额外人物时，人脸身份选择实际接受动作编辑的人物，选择框和原因在逐图标注中保留。姿态仍沿用原评测器的人物选择逻辑，因此此类失败不能只依据PCK或身份高低判断。',9)
    doc.add_page_break();old.para(doc,'单独增加B照片、法线、深度的影响',13,True)
    for resolution in (512,1024):
        for dataset in ('iper','tiktok'):
            group=f'{dataset}_{resolution}_steps30'
            factor_table(doc,summary['combination_factor_differences'][group],f'{dataset}｜{resolution}、30步')
    doc.add_page_break();old.para(doc,'原生组合与既有Adapter的同样本对照',13,True)
    for resolution in (512,1024):
        for dataset in ('iper','tiktok'):
            subset=[r for r in rows if r['dataset']==dataset and r['resolution']==resolution and r['steps']==30]
            methods=('native_skeleton','native_skeleton_depth','simple_body','full_body','full_body_face')+ (('simple_body1024',) if resolution==1024 else ())
            nodes,n=shared_nodes(subset,methods)
            quality_table(doc,nodes,methods,f'{dataset}｜{resolution}：骨架、骨架＋深度与控制方案',n)
    old.para(doc,'八种原生路线与各Adapter的逐对身份、动作和背景增量均保留在组合评测文件中。本报告不给不同指标强行合成一个总分，也不自动替换默认权重。完整Body/Face的1024结果仍是512训练权重用于1024推理。',9)
    # Historical statistics are copied from their immutable source, so extending
    # the common-face grid cannot silently change the old report's identities.
    legacy=json.loads((root/'archive/evaluation/metric_summary.json').read_text())
    doc.add_page_break();old.para(doc,'原有路线数据（保留原口径）',13,True)
    for ds in ('iper','tiktok'):old.metric_table(doc,legacy,f'{ds}_512_steps30',f'{ds}｜原有512对照',True,3)
    for ds in ('iper','tiktok'):old.geometry_adapter_table(doc,legacy,f'{ds}_512_steps30')
    for ds in ('iper','tiktok'):old.face_increment_table(doc,legacy,f'{ds}_512_steps30')
    old.supplementary_tables(doc,legacy,True,3);bfs=old.add_bfs(doc)
    old.para(doc,'原表的身份共同子集与新八组合的共同子集可能不同，不能直接跨表相减；少于3对只作有限观察。原有40步与外部LoRA历史参考继续单列，没有新增40步或LoRA实验。耗时含采样和控制准备，不含模型加载及已缓存的上游几何、ArcFace/DINO提取；二次编辑含两段采样。',9)
    kinds={'native_combinations_1':'原生组合①：骨架／照片／单独法线／单独深度',
           'native_combinations_2':'原生组合②：多条件共同输入','legacy_pose':'原有动作对照',
           'legacy_face':'原有面部对照','legacy_1024':'原有1024对照','sampling_steps':'原有30／40步附录'}
    blocks=[]
    for pid in ids:
        f=[x for x in figures if x['pair_id']==pid and x['resolution']==512]
        blocks.extend([f[:2],f[2:]])
    for pid in high:
        f=[x for x in figures if x['pair_id']==pid and x['resolution']==1024 and x['steps']==30]
        blocks.append(f[:2])
    context=[x for x in figures if x['part']=='legacy_1024'];blocks.extend([context[:2],context[2:]])
    blocks.append([x for x in figures if x['part']=='sampling_steps'])
    for block in blocks:
        doc.add_page_break();old.para(doc,'固定案例对照｜整图＋人脸',12,True)
        old.para(doc,'第一行整图，第二行实际检测脸，列宽一致；仅用于展示的不可评脸会标明。无损图片已直接嵌入，放大不能恢复原图缺失的细节。',8)
        for fig in block:
            old.para(doc,f'{fig["pair_id"]}｜{fig["resolution"]}／{fig["steps"]}步｜{kinds[fig["part"]]}',9,True)
            doc.add_picture(fig['path'],width=Cm(18.2))
        for pid in dict.fromkeys(x['pair_id'] for x in block):
            old.para(doc,pid+'：'+obs.get('figure_notes',{}).get(pid,'固定展示样本，完整保留各路线结果。'),9)
    doc.settings._element.append(OxmlElement('w:doNotAutoCompressPictures'))
    args.output.parent.mkdir(parents=True,exist_ok=True)
    tmp=args.output.with_name(args.output.stem+'_building.docx');doc.save(tmp)
    with zipfile.ZipFile(tmp) as z:
        if z.testzip() is not None:raise ValueError('Word ZIP integrity failed')
        media=[x for x in z.namelist() if x.startswith('word/media/')]
        if len(media)!=37 or {hashlib.sha256(z.read(x)).hexdigest() for x in media}!={f['sha256'] for f in figures}:
            raise ValueError('Word embedded-media mismatch')
    tmp.replace(args.output)
    save(root/'report_figures/manifest.json',figures)
    save(root/'report_manifest.json',{'docx':str(args.output),'sha256':sha(args.output),'revision':'condition_combinations_v1',
        'expected_outputs':len(expected),'evaluated_outputs':len(rows),'embedded_figures':37,'word_zip_verified':True,
        'formal_identity_review_complete':True,'previous_report':str(root/'archive/report_before_supplement.docx'),
        'original_experiment':str(parent),'protocol_sha256':sha(args.protocol),'summary_sha256':sha(root/'evaluation/metric_summary.json'),
        'face_review_sha256':sha(root/'evaluation/face_review.json'),'observations_sha256':sha(root/'final_observations.json'),
        'figures':figures,'bfs_separate_reference':bfs,'office_visual_render_performed':False})
    print(str(args.output),flush=True)
