#!/usr/bin/env python3
"""Build a readable Word report from the locked Qwen native/control audit.

Images are embedded, with whole-image and actual-detected-face rows. Figures
are chosen from the precommitted selection, never sorted by result quality.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import math
from pathlib import Path
import sys
import zipfile

PROJECT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(PROJECT))
import numpy as np
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from scripts.eval.score_qwen21_native_value import DEFAULT_ROOT, METHODS, FONT, rgb, save, sha, read_jsonl
from src.data.letterbox import Letterbox

CN = {'native_ab':'原生 A＋B','native_ab_crop':'原生 A＋B＋A脸',
      'native_skeleton':'原生 A＋骨架','native_geometry':'原生 A＋法线/深度',
      'simple_body':'简化 Body 20k','full_body':'完整 Body 6000',
      'full_body_face':'完整 Body＋Face','full_body_native_face':'完整 Body→原生修脸',
      'simple_body1024':'简化 1024 Body 3k'}
POSE_FIGURE = ('native_ab','native_skeleton','native_geometry','simple_body')
FACE_FIGURE = ('native_ab_crop','full_body','full_body_face','full_body_native_face')
BFS_PATH = Path('/home/shangguanrz/nihaochen/qwen21_full/bfs_compare_30_20261004/results.json')


def font(run, size=10, bold=False):
    run.font.name = 'Noto Sans CJK SC'
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Noto Sans CJK SC')
    run.font.size = Pt(size)
    run.bold = bold


def para(doc, text, size=10, bold=False):
    p = doc.add_paragraph()
    font(p.add_run(text),size,bold)
    p.paragraph_format.space_after=Pt(4)
    p.paragraph_format.line_spacing=1.05
    return p


def number(value, digits=3):
    return '无法判断' if value is None or not math.isfinite(value) else f'{value:.{digits}f}'


def summary_value(node, key):
    value = node.get(key)
    return value.get('mean') if isinstance(value,dict) else value


def metric_table(doc, summary, key, title, manual_complete, minimum_identity_bold_n=0):
    if key not in summary['groups']:
        return
    nodes = summary['groups'][key]
    methods = [m for m in (*METHODS,'simple_body1024') if m in nodes]
    manual_complete=all(nodes[m].get('face_manual_review_complete',manual_complete) for m in methods)
    para(doc,title,11,True)
    identity = 'identity_a_cosine_formal_common' if manual_complete else 'identity_a_cosine_common'
    columns = [('pose_pck_002','动作位置↑',3,True),
               ('pose_torso_aligned_pck_002','动作结构↑',3,True),
               (identity,'A身份↑',3,True),
               ('identity_a_cosine_raw','原始A相似度',3,None),
               ('background_mae_source','背景误差↓',2,False),
               ('seconds','秒/张↓',1,False),
               ('peak_cuda_gib','GiB↓',1,False)]
    table=doc.add_table(rows=1,cols=1+len(columns));table.style='Table Grid'
    for cell,text in zip(table.rows[0].cells,['路线',*[c[1] for c in columns]]):
        font(cell.paragraphs[0].add_run(text),8,True)
    best={}
    for metric,_,_,up in columns:
        values=[summary_value(nodes[m],metric) for m in methods]
        values=[v for v in values if v is not None and math.isfinite(v)]
        best[metric]=(max(values) if up else min(values)) if values and up is not None else None
    for method in methods:
        cells=table.add_row().cells
        font(cells[0].paragraphs[0].add_run(CN[method]),8)
        for cell,(metric,_,digits,_) in zip(cells[1:],columns):
            value=summary_value(nodes[method],metric)
            text=number(value,digits)
            if metric=='pose_torso_aligned_pck_002':
                text+=f' (n={nodes[method][metric]["n"]})'
            if metric=='identity_a_cosine_raw':
                text+=f' (n={nodes[method][metric]["n"]})'
            comparable = metric != identity or nodes[method][identity]['n'] >= minimum_identity_bold_n
            font(cell.paragraphs[0].add_run(text),8,value is not None and value==best[metric] and comparable)
    n=nodes[methods[0]][identity]['n']
    n_images=nodes[methods[0]]['n_images']
    note=f'每条路线 {n_images} 对；身份使用各路线同一批可评脸（n={n}）。n 是可评价数量，不是训练数量。'
    if not manual_complete:
        note+='目前身份子集仅通过自动检查，尚待逐图确认，不能据此作正式身份结论。'
    if n<3:
        note+='共同可靠脸过少，身份优劣无法可靠判断；逐样本原始分数及失败原因另有保留。'
    if n==0:
        note+='n=0表示没有案例同时满足所有路线的可靠脸条件，并非各路线都没有生成脸；后面的两两比较使用各自共同子集。'
    para(doc,note+' 原始A相似度保留全部检出脸（另列n），可能包含误检和过小脸，只作辅助，不加粗排名。其他加粗表示每列最优，不代表综合获胜。',8)


def geometry_adapter_table(doc,summary,group):
    records=summary.get('paired_differences',{}).get(group,{})
    comparisons=(('simple_body_minus_native_skeleton','简化Body vs 原生骨架'),
                 ('full_body_minus_native_skeleton','完整Body vs 原生骨架'))
    if not any(key in records for key,_ in comparisons):return
    para(doc,'原生骨架与Adapter｜'+group,10,True)
    table=doc.add_table(rows=1,cols=7);table.style='Table Grid'
    headers=['比较','原生A身份','Adapter A身份','身份变化','共同脸n','动作变化','背景变化']
    for cell,text in zip(table.rows[0].cells,headers):font(cell.paragraphs[0].add_run(text),8,True)
    for key,label in comparisons:
        if key not in records:continue
        record=records[key];ident=record['identity_a_cosine_reliable']
        vals=[label,number(ident.get('first_mean'),3),number(ident.get('second_mean'),3),
              number(ident['mean'],4),str(ident['n']),number(record['pose_pck_002']['mean'],4),
              number(record['background_mae_source']['mean'],3)]
        for cell,text in zip(table.add_row().cells,vals):font(cell.paragraphs[0].add_run(text),8)
    para(doc,'身份只在这两个方案均经人工核实、可靠可评的同一批人脸上比较；与主表八方案共同子集分开。动作、背景差值使用相同固定配对。变化为Adapter减去原生骨架：身份、动作越高越好，背景变化越低越好。共同脸不足时无法判断。',8)


def face_increment_table(doc,summary,group):
    records=summary.get('paired_differences',{}).get(group,{})
    comparisons=(('full_body_face_minus_full_body','完整Body＋Face'),
                 ('full_body_native_face_minus_full_body','完整Body→原生修脸'))
    if not any(key in records for key,_ in comparisons):return
    para(doc,'相对完整Body的面部增量｜'+group,10,True)
    table=doc.add_table(rows=1,cols=6);table.style='Table Grid'
    for cell,text in zip(table.rows[0].cells,['方案','A身份变化','可评脸n','身份提高数','动作变化','背景误差变化']):font(cell.paragraphs[0].add_run(text),8,True)
    for key,label in comparisons:
        if key not in records:continue
        record=records[key];identity=record['identity_a_cosine_reliable'];pose=record['pose_pck_002'];bg=record['background_mae_source']
        vals=[label,number(identity['mean'],4),str(identity['n']),str(identity['improved_n']),number(pose['mean'],4),number(bg['mean'],3)]
        for cell,text in zip(table.add_row().cells,vals):font(cell.paragraphs[0].add_run(text),8)
    para(doc,'这里分别固定Body与该方案都可靠、且人工确认可评价的脸。因此两行n可能不同，也可能大于八路线共同子集；不能跨行只看均值排名。正身份变化是收益；动作变化、背景误差变化也必须同时看。',8)


def crop_actual_face(image, record, selection):
    manual_box=selection.get('manual_face_box_native')
    if manual_box is not None:
        if len(manual_box)!=4 or not all(isinstance(v,(int,float)) and math.isfinite(v) for v in manual_box):
            raise ValueError('Manual display box must contain four finite coordinates')
        x0,y0,x1,y1=manual_box
        if not 0<=x0<x1<=image.width or not 0<=y0<y1<=image.height:
            raise ValueError('Manual display box lies outside the native image')
        box=[int(x0),int(y0),int(x1),int(y1)]
        crop=image.crop(box)
        return Letterbox.from_size(crop.size,512).image(crop,fill=(240,240,240)),'人工定位，仅用于展示',box
    if selection.get('reviewed') and selection.get('is_real_face') is not True:
        canvas=Image.new('RGB',(512,512),(240,240,240))
        ImageDraw.Draw(canvas).text((22,220),'未检出可靠脸',font=ImageFont.truetype(FONT,38),fill='black')
        return canvas,'不展示已排除的错检',None
    index=selection.get('selected_index')
    label=('实际检测脸' if selection.get('reliable') else '不可可靠评价，供查看') if selection.get('reviewed') else '检测候选，待核对'
    if index is None:
        index=record.get('raw_largest_index')
        label='检测不可靠，供查看'
    if index is None:
        canvas=Image.new('RGB',(512,512),(240,240,240))
        ImageDraw.Draw(canvas).text((30,225),'未检出可靠人脸',font=ImageFont.truetype(FONT,34),fill='black')
        return canvas,'未检出脸',None
    x0,y0,x1,y1=record['detections'][index]['box_native']
    side=max(x1-x0,y1-y0)*1.45
    cx,cy=(x0+x1)/2,(y0+y1)/2
    # Clip only at native-image bounds; never substitute B geometry for G.
    box=(max(0,int(cx-side/2)),max(0,int(cy-side/2)),min(image.width,int(cx+side/2)),min(image.height,int(cy+side/2)))
    crop=image.crop(box)
    return Letterbox.from_size(crop.size,512).image(crop,fill=(240,240,240)),label,list(box)


def figure(root, rows, pair, resolution, methods, suffix, steps=30):
    selected={r['method']:r for r in rows if r['pair_id']==pair['pair_id'] and r['resolution']==resolution and r['steps']==steps}
    if any(m not in selected for m in methods):
        return None
    anchor=selected[methods[0]]
    inputs=[('A：身份/服装/背景','source',anchor),('B：仅目标动作','target',anchor)]
    specs=inputs+[(CN[m],'generated',selected[m]) for m in methods]
    if len(specs)>6:
        raise ValueError('Figures must not exceed six columns')
    canvas=Image.new('RGB',(512*len(specs),1132),'white');draw=ImageDraw.Draw(canvas);details=[]
    for col,(title,role,row) in enumerate(specs):
        detection=row['face_detection'][role]
        original=rgb(detection['path'])
        display=Letterbox.from_size(original.size,512).image(original,fill=(128,128,128))
        canvas.paste(display,(col*512,52))
        draw.text((col*512+8,8),title,font=ImageFont.truetype(FONT,28),fill='black')
        crop,label,box=crop_actual_face(original,detection,row['face_selection'][role])
        canvas.paste(crop,(col*512,616))
        draw.text((col*512+8,575),label,font=ImageFont.truetype(FONT,25),fill='black')
        details.append({'label':title,'path':detection['path'],'sha256':detection['sha256'],
                        'face_crop_box_native':box,'face_crop_method':label})
    path=root/'report_figures'/f'{pair["pair_id"]}_{resolution}_steps{steps}_{suffix}.png'
    path.parent.mkdir(parents=True,exist_ok=True);canvas.save(path)
    return {'pair_id':pair['pair_id'],'dataset':pair['dataset'],'resolution':resolution,
            'steps':steps,'part':suffix,'methods':list(methods),'path':str(path),'sha256':sha(path),'details':details}


def read_observations(root):
    path=root/'final_observations.json'
    if not path.exists():
        return {'headline':'实验数据已整理；逐图观察结论尚未填写。',
                'conclusions':['暂无经人工复核的效果结论，不将自动指标当作综合画质判断。'],
                'figure_notes':{},'limitations':'本轮只有固定少量案例，不能证明某一类方法普遍必要或无用。'}
    return json.loads(path.read_text())


def sampling_steps_figure(root,rows,pair):
    configurations=(('native_ab',30),('native_ab',40),('native_ab_crop',30),('native_ab_crop',40))
    selected={(r['method'],r['steps']):r for r in rows if r['pair_id']==pair['pair_id'] and r['resolution']==1024}
    if any(key not in selected for key in configurations):return None
    anchor=selected[configurations[0]]
    specs=[('A：身份/服装/背景','source',anchor),('B：仅目标动作','target',anchor)]
    specs += [(f'{CN[method]} {steps}步','generated',selected[(method,steps)]) for method,steps in configurations]
    canvas=Image.new('RGB',(512*len(specs),1132),'white');draw=ImageDraw.Draw(canvas);details=[]
    for col,(title,role,row) in enumerate(specs):
        detection=row['face_detection'][role];image=rgb(detection['path'])
        whole=Letterbox.from_size(image.size,512).image(image,fill=(128,128,128))
        canvas.paste(whole,(col*512,52));draw.text((col*512+6,8),title,font=ImageFont.truetype(FONT,26),fill='black')
        crop,label,box=crop_actual_face(image,detection,row['face_selection'][role])
        canvas.paste(crop,(col*512,616));draw.text((col*512+6,575),label,font=ImageFont.truetype(FONT,25),fill='black')
        details.append({'label':title,'path':detection['path'],'sha256':detection['sha256'],'face_crop_box_native':box,'face_crop_method':label})
    path=root/'report_figures'/f'{pair["pair_id"]}_1024_steps30_vs40.png';path.parent.mkdir(parents=True,exist_ok=True);canvas.save(path)
    return {'pair_id':pair['pair_id'],'dataset':pair['dataset'],'resolution':1024,'steps':'30/40',
            'part':'sampling_steps','methods':[f'{m}_{s}' for m,s in configurations],
            'path':str(path),'sha256':sha(path),'details':details}


def add_bfs(doc):
    if not BFS_PATH.exists():
        return None
    data=json.loads(BFS_PATH.read_text())['summary']
    para(doc,'独立参考：同输入下的原生修脸与外部 LoRA',11,True)
    columns=('body','nolora','bfs','alternative')
    labels={'body':'原 Body','nolora':'原生修脸，无LoRA','bfs':'BFS LoRA','alternative':'BFS alternative'}
    t=doc.add_table(rows=1,cols=4);t.style='Table Grid'
    for cell,text in zip(t.rows[0].cells,['方案','A身份↑','动作PCK↑','脸外变化↓']):font(cell.paragraphs[0].add_run(text),8,True)
    keys=('common_identity_mean','eligible_pose_pck_002_mean','outside_head_mae_255_mean')
    for method in columns:
        cells=t.add_row().cells;font(cells[0].paragraphs[0].add_run(labels[method]),8)
        for cell,key in zip(cells[1:],keys):font(cell.paragraphs[0].add_run(number(data[method].get(key),3)),8)
    para(doc,'来自已有30个案例中的24个可靠脸：第一阶段是简化Body，第二阶段768分辨率。原生修脸已显示较强身份恢复能力；这两份外部LoRA的身份均值低于不加LoRA，但动作更接近第一阶段。条件、数据与本轮不同，不能放进本轮主表排名，也不能推广为所有LoRA都无用。',8)
    return {'path':str(BFS_PATH),'sha256':sha(BFS_PATH),'common_n':24}


def supplementary_tables(doc,summary,manual_complete,minimum_identity_bold_n=0):
    for dataset in ('iper','tiktok'):
        metric_table(doc,summary,f'{dataset}_1024_steps30',f'1024复核｜{dataset}｜固定2对，30步',manual_complete,minimum_identity_bold_n)
    for dataset in ('iper','tiktok'):
        key=f'{dataset}_1024_steps40'
        if key in summary['groups']:
            metric_table(doc,summary,key,f'基座辅助检查｜{dataset}｜1024、40步',manual_complete,minimum_identity_bold_n)
    para(doc,'1024和40步只在预先固定的4对上复核，不能与主比较12对的均值直接相减。完整Body/Face仍是512训练权重用于1024推理；“简化1024 Body 3k”才来自原1024续训。40步表独立展示，不混入30步排名。',9)
    differences=summary.get('sampling_step_differences',{})
    if differences:
        para(doc,'30→40步｜同一案例、同一原生路线',10,True)
        table=doc.add_table(rows=1,cols=7);table.style='Table Grid'
        for cell,text in zip(table.rows[0].cells,['数据/路线','A身份变化','共同脸n','动作位置变化','动作结构变化','背景误差变化','耗时变化/秒']):font(cell.paragraphs[0].add_run(text),8,True)
        for group,methods in sorted(differences.items()):
            for method,record in methods.items():
                identity=record['identity_a_cosine_reliable']
                values=[group.split('_')[0]+' / '+CN[method],number(identity['mean'],4),str(identity['n']),
                        number(record['pose_pck_002']['mean'],4),number(record['pose_torso_aligned_pck_002']['mean'],4),
                        number(record['background_mae_source']['mean'],3),number(record['seconds']['mean'],1)]
                for cell,text in zip(table.add_row().cells,values):font(cell.paragraphs[0].add_run(text),8)
        para(doc,'均为40步减30步；身份只比较两个采样节点都可靠且人工确认可评的同一批脸。前两张表的身份共同子集不同，不能直接相减。身份、动作正值表示提高；背景、耗时正值表示增加。',8)


def main():
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('--root',type=Path,default=DEFAULT_ROOT)
    p.add_argument('--output',type=Path,default=Path('/home/shangguanrz/nihaochen/docs/Qwen基座与控制模块_价值核对_20261008.docx'))
    p.add_argument('--protocol',type=Path,help='Build the supplementary combinations revision, preserving legacy defaults')
    p.add_argument('--allow-incomplete',action='store_true',help='Create an explicitly provisional preview before all images and reviews are ready')
    args=p.parse_args();root=args.root
    if args.protocol:
        from scripts.reports.build_qwen21_combination_report import build_report
        return build_report(args)
    rows=json.loads((root/'evaluation/evaluated_samples.json').read_text())
    summary=json.loads((root/'evaluation/metric_summary.json').read_text())
    pairs=read_jsonl(root/'data/pairs.jsonl')
    selection_path=root/'data/report_selection.json'
    if not selection_path.exists():
        raise FileNotFoundError('Precommitted data/report_selection.json required; do not select images by score')
    selection=json.loads(selection_path.read_text())
    ids=selection.get('pair_ids',selection.get('examples',[]))
    if ids and isinstance(ids[0],dict):ids=[x['pair_id'] for x in ids]
    if not 4<=len(ids)<=6:
        raise ValueError('Report must show four to six preselected cases')
    by_id={x['pair_id']:x for x in pairs};obs=read_observations(root)
    manual_complete=summary['formal_identity_review_complete']
    highres_ids=selection.get('highres_pair_ids',[p['pair_id'] for p in pairs if p.get('highres_check')])
    expected={(pair['pair_id'],512,30,method) for pair in pairs for method in METHODS}
    expected|={(pair_id,1024,30,method) for pair_id in highres_ids for method in (*METHODS,'simple_body1024')}
    expected|={(pair_id,1024,40,method) for pair_id in highres_ids for method in ('native_ab','native_ab_crop')}
    present={(r['pair_id'],r['resolution'],r['steps'],r['method']) for r in rows}
    missing=sorted(expected-present)
    if not args.allow_incomplete:
        if missing:raise ValueError(f'Incomplete audit: {len(missing)} planned outputs missing, e.g. {missing[:3]}')
        if not manual_complete:raise ValueError('Human face review is incomplete; final report requires all fixed outputs reviewed')
        if not (root/'final_observations.json').exists():raise FileNotFoundError('Human final_observations.json is required for the final report')
    figures=[]
    for pair_id in ids:
        pair=by_id[pair_id]
        for methods,suffix in ((POSE_FIGURE,'pose'),(FACE_FIGURE,'face')):
            value=figure(root,rows,pair,512,methods,suffix)
            if value:figures.append(value)
    for pair_id in highres_ids:
        pair=by_id[pair_id]
        value=figure(root,rows,pair,1024,('native_ab','native_skeleton','native_geometry','full_body_face'),'1024_key_routes')
        if value:figures.append(value)
    # Fixed first iPER high-resolution case; not selected after inspecting results.
    if 'iper_00' in highres_ids:
        value=sampling_steps_figure(root,rows,by_id['iper_00'])
        if value:figures.append(value)
    doc=Document();sec=doc.sections[0]
    sec.page_width=Cm(21);sec.page_height=Cm(29.7)
    sec.left_margin=sec.right_margin=Cm(1.45);sec.top_margin=sec.bottom_margin=Cm(1.35)
    para(doc,'Qwen基座与控制模块：价值核对',17,True)
    para(doc,'2026年10月8日｜不同人物A/B｜只做推理，不训练',9)
    if args.allow_incomplete and (missing or not manual_complete or not (root/'final_observations.json').exists()):
        para(doc,'临时预览：部分固定输出或人工审核尚未完成，不能作为最终结果报告。',11,True)
    para(doc,obs['headline'],11,True)
    para(doc,'为什么做这次实验',12,True)
    para(doc,'此前的原生基线通常只收到A图和一句泛化指令，并未真正看到B的动作。这不足以说明基座不能完成任务。本轮先给基座明确的图片角色、逐图核对的详细动作描述和保留要求，再与控制模块公平比较。')
    para(doc,'怎样理解各条路线',12,True)
    para(doc,'原生Qwen直接读参考图片并编辑，不加控制模块。Body Adapter把B的身体几何送入生成过程，帮助改变动作；Face Adapter从A提取身份、外观，从B提取面部几何，帮助保脸和转头。A脸图则直接作为Qwen的一张额外参考，机制与Face Adapter不同。二次原生修脸是在Body结果上再编辑一次，因此成本也算两次。')
    para(doc,'正式比较12对不同人物A/B（iPER、TikTok各6对），实际包含6个A人物，每个A对应多个B动作；另4对独立案例打磨详细提示词。模板和正式动作描述在看结果前锁定。各路线512、30步、种子42；旧检查点保持原样，未加载LoRA，未替换默认方案。')
    para(doc,'已核对的iPER Adapter训练清单与本轮A/B外观不重叠；TikTok验证原视频组与8548对训练池分开。这只说明我们自己的训练划分，不代表Qwen预训练从未见过这些人物或图片。',9)
    para(doc,'重点结论',12,True)
    for text in obs.get('conclusions',[]):para(doc,text)
    para(doc,'结论限于固定12对（6个A人物）与1024复核4对（2个A人物）：简化版与完整版的训练数据、训练量不同；真正按A体型和相机重定向尚未测试；完整架构1024结果只是512权重用于推理，不能据此判断普遍泛化或架构本身优劣。',9)
    doc.add_page_break()
    para(doc,'主要数据｜相同512分辨率、30步',13,True)
    for dataset,title in (('iper','iPER：6对不同人物'),('tiktok','TikTok：6对不同人物')):
        metric_table(doc,summary,f'{dataset}_512_steps30',title,manual_complete)
    para(doc,'动作位置是PCK@0.02：B可见关节中，生成关节落在目标附近的比例；512正方形图的容差约14像素，补边竖图更小，要求很严格，低分或0分并不等于没有改变动作。“动作结构”先将躯干居中并统一缩放到同一长度，再算PCK，不旋转身体、不单独拉伸四肢，仅辅助减小体型与位置差异的影响；需要双肩和双髋同时有效，因此另列n。躯干估计不稳时也会影响它，它不评价是否站在正确位置，不能替代原始PCK。A身份是人脸特征相似度，越高越像A，并非身份正确率或百分比；背景误差越低越保留A背景。检测失败计入动作失败，人脸低分不因分数低而删除。',9)
    para(doc,'B是另一个人，不是理想生成图，所以不算对B的PSNR、SSIM或人体重建误差。背景在A/B人体区域并集外、双方有效画面交集内与A比较；未使用生成图的人体分割，iPER部分人体轮廓来自几何渲染，服装、影子或人物移动到评价区域时都会影响该误差。所有原始身份分数、与B的相似度、检测框和排除原因均保留。',9)
    para(doc,'耗时包括采样与控制准备，已缓存的上游几何、ArcFace/DINO特征提取及模型加载未计入；二次编辑累计两段采样。这里不是从原始照片到结果的一条龙总成本。',8)
    doc.add_page_break();para(doc,'补充检查与历史参考',13,True)
    for dataset in ('iper','tiktok'):geometry_adapter_table(doc,summary,f'{dataset}_512_steps30')
    for dataset in ('iper','tiktok'):face_increment_table(doc,summary,f'{dataset}_512_steps30')
    if obs.get('resolution_sampling_observation'):
        para(doc,'1024／40步：固定4对的辅助观察',11,True)
        para(doc,obs['resolution_sampling_observation'],10)
    supplementary_tables(doc,summary,manual_complete)
    bfs=add_bfs(doc)
    # Pair the two 512 panels on one page; retain the original 3072px PNG
    # inside Word so zooming does not depend on a reduced screenshot.
    sec=doc.add_section(WD_SECTION_START.NEW_PAGE)
    sec.page_width=Cm(21);sec.page_height=Cm(29.7)
    sec.left_margin=sec.right_margin=Cm(1.3);sec.top_margin=sec.bottom_margin=Cm(1.1)
    blocks=[[fig for fig in figures if fig['pair_id']==case and fig['resolution']==512] for case in ids]
    high=[fig for fig in figures if fig['part']=='1024_key_routes']
    blocks+=[high[index:index+2] for index in range(0,len(high),2)]
    auxiliary=[fig for fig in figures if fig['part']=='sampling_steps']
    if auxiliary:blocks.append(auxiliary)
    blocks=[block for block in blocks if block]
    for index,block in enumerate(blocks):
        if index:doc.add_page_break()
        para(doc,'固定案例对照｜整图＋人脸',12,True)
        para(doc,'每张图第一行整图，第二行实际人脸，列宽一致。原始图片直接嵌入，可放大查看；放大不能恢复原图没有的细节。',9)
        seen=set()
        for fig in block:
            kind={'pose':'动作与外观','face':'面部方案','1024_key_routes':'高分辨率复核','sampling_steps':'30／40步辅助检查'}[fig['part']]
            para(doc,f'{fig["dataset"]}｜{fig["pair_id"]}｜{fig["resolution"]}／{fig["steps"]}步｜{kind}',10,True)
            doc.add_picture(fig['path'],width=Cm(18.2))
            seen.add(fig['pair_id'])
        for case in sorted(seen):
            note=obs.get('figure_notes',{}).get(case,'预先固定的展示案例，未根据结果更换。')
            para(doc,case+'：'+note,9)
    no_compression=OxmlElement('w:doNotAutoCompressPictures')
    doc.settings._element.append(no_compression)
    args.output.parent.mkdir(parents=True,exist_ok=True);doc.save(args.output)
    with zipfile.ZipFile(args.output) as z:
        if z.testzip() is not None:raise ValueError('Word archive failed integrity check')
        media=[name for name in z.namelist() if name.startswith('word/media/')]
        embedded={hashlib.sha256(z.read(name)).hexdigest() for name in media}
        if embedded != {fig['sha256'] for fig in figures}:raise ValueError('Embedded figures differ from manifest')
    save(root/'report_figures/manifest.json',figures)
    save(root/'report_manifest.json',{'docx':str(args.output),'sha256':sha(args.output),
         'provisional':bool(missing or not manual_complete or not (root/'final_observations.json').exists()),
         'expected_outputs':len(expected),'evaluated_outputs':len(present),'missing_planned_outputs':missing,
         'embedded_figures':len(figures),'selected_main_cases':ids,'highres_cases':highres_ids,
         'source_summary':str(root/'evaluation/metric_summary.json'),'summary_sha256':sha(root/'evaluation/metric_summary.json'),
         'evaluated_samples_sha256':sha(root/'evaluation/evaluated_samples.json'),
         'locked_pairs_sha256':sha(root/'data/pairs.jsonl'),'selection_sha256':sha(selection_path),
         'observations':str(root/'final_observations.json'),
         'observations_sha256':sha(root/'final_observations.json') if (root/'final_observations.json').exists() else None,
         'face_review_sha256':sha(root/'evaluation/face_review.json') if (root/'evaluation/face_review.json').exists() else None,
         'formal_identity_review_complete':manual_complete,
         'bfs_separate_reference':bfs,'word_zip_verified':True,'figures':figures})
    print(str(args.output),flush=True)


if __name__=='__main__':main()
