#!/usr/bin/env python3
"""Short, embedded-image Word report for the fixed three-arm BFS audit."""
from __future__ import annotations
import argparse
import hashlib
import json
import zipfile
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from scripts.eval.eval_qwen21_bfs_skeleton import DEFAULT, sha256
from scripts.reports.build_qwen21_native_value_report import font, para, number

OUTPUT=Path('/home/shangguanrz/nihaochen/docs/BFS_LoRA直接增强骨架编辑_对比结果_20261010.docx')
METHODS=('native','bfs_body','bfs_head')
LABELS={'native':'原生','bfs_body':'BFS Body','bfs_head':'BFS Head'}
SELECTION=(('iper_00',512),('tiktok_00',512),('iper_00',1024),('iper_01',1024),
           ('tiktok_00',1024),('tiktok_01',1024))


def value(node,key):
    v=node.get(key)
    return v.get('mean') if isinstance(v,dict) else v


def table(doc,nodes,title):
    para(doc,title,10,True)
    columns=[('pose_pck_002','动作PCK↑',3,True),('pose_nme_diagonal','位置误差↓',4,False),
             ('identity_a_cosine_formal_common','A脸相似度↑',3,True),
             ('background_mae_source','背景误差↓',2,False),('seconds','秒/张↓',1,False),
             ('peak_cuda_gib','显存GiB↓',1,False),('face_detected_n','检出脸',0,None)]
    tab=doc.add_table(rows=1,cols=8);tab.style='Table Grid'
    for cell,text in zip(tab.rows[0].cells,['方案',*[c[1] for c in columns]]):
        font(cell.paragraphs[0].add_run(text),8,True)
    for method in METHODS:
        node=nodes[method];cells=tab.add_row().cells
        font(cells[0].paragraphs[0].add_run(LABELS[method]),8)
        for col,(key,_,digits,up) in enumerate(columns,1):
            v=value(node,key)
            vals=[value(nodes[m],key) for m in METHODS if value(nodes[m],key) is not None]
            comparable=up is not None and bool(vals)
            if key=='identity_a_cosine_formal_common':
                n=node[key]['n'];comparable=comparable and n>=3
                text=number(v,digits)+f' (n={n})'
            elif key=='face_detected_n':text=f'{int(v)}/{node["n_images"]}'
            else:text=number(v,digits)
            if method=='native' and key in ('seconds','peak_cuda_gib'):text+='*'
            best=(max(vals) if up else min(vals)) if vals else None
            bold=comparable and v is not None and abs(v-best)<1e-10
            font(cells[col].paragraphs[0].add_run(text),8,bold)
    for row in tab.rows:
        for cell in row.cells:
            for p in cell.paragraphs:p.paragraph_format.space_after=Pt(1)
        trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--root',type=Path,default=DEFAULT)
    parser.add_argument('--output',type=Path,default=OUTPUT)
    args=parser.parse_args();root=args.root
    stats=json.loads((root/'evaluation/metric_summary.json').read_text())
    observations=json.loads((root/'evaluation/observations.json').read_text())
    panels=json.loads((root/'report_figures/manifest.json').read_text())
    assert len(panels)==16 and stats['formal_identity_review_complete']
    doc=Document();sec=doc.sections[0]
    sec.page_width=Cm(21);sec.page_height=Cm(29.7)
    sec.top_margin=Cm(1.4);sec.bottom_margin=Cm(1.4);sec.left_margin=Cm(1.4);sec.right_margin=Cm(1.4)
    normal=doc.styles['Normal'];normal.font.name='Noto Sans CJK SC';normal.font.size=Pt(10)
    settings=doc.settings.element
    compression=OxmlElement('w:doNotAutoCompressPictures');settings.append(compression)
    para(doc,'BFS LoRA 能否直接增强“A＋骨架”？',17,True)
    para(doc,'固定案例对比 · 2026年10月10日',9)
    para(doc,'本轮做了什么',11,True)
    para(doc,'每次都只给 Qwen 一张人物原图 A 和一张目标骨架图，使用上一轮同一份详细提示词。比较不加载 LoRA、加载 BFS Body V1.0、加载 BFS Head V1.1 三种情况。Body 与 Head 分别使用，强度都是 1.0。',10)
    para(doc,'本轮没有训练，也没有启用项目自训的 Body／Face Adapter，没有二次修脸。BFS LoRA 是外部现成权重，加载后改变基座的计算；项目 Adapter 则是另外训练的控制分支，两者不是同一模块。',10)
    para(doc,'先看结论',11,True)
    for line in observations['conclusions']:para(doc,line,10)
    para(doc,'关键数据',11,True)
    for dataset,res in (('iper',512),('tiktok',512),('iper',1024),('tiktok',1024)):
        nodes=stats['groups'][f'{dataset}_{res}_steps30']
        n=nodes['native']['n_images']
        table(doc,nodes,f'{"iPER" if dataset=="iper" else "TikTok"} · {res}×{res} · {n}对')
    para(doc,'读表：PCK 越高表示关键点更接近目标；位置误差、背景误差越低越好。A脸相似度只比较三组都可可靠评价的同一批人脸，n 是有效配对数；n<3只供参考，不加粗排名。同一指标在三组中按方向标粗最优值，可能并列。检出脸只表示检测成功，不等于保持得好。',8.5)
    para(doc,'* 原生耗时和显存为核验后复用的历史测量。其余为本轮测量；峰值是 PyTorch 已分配显存。1024的姿态和背景缩至512评价，人脸使用原尺寸。',8.5)
    doc.add_page_break()
    para(doc,'如何理解这些结果',14,True)
    for line in observations['interpretation']:para(doc,line,10)
    para(doc,'比较范围与限制',11,True)
    para(doc,'512：12对不同人物 A/B；1024：预先固定的4对。每个输入都比较三组，共48张正式结果，其中16张原生结果复用、32张LoRA结果新生成。另有独立开发样本，只检查加载和运行正确性，不参与排名。',10)
    para(doc,'生成固定为30步采样、随机种子42、true CFG 1.0，关闭KV缓存；完整A在前、骨架在后，等比例缩放并补边。三组提示词逐字相同，包含具体目标动作及A的外观保持要求，因此比较的是“骨架＋详细动作文字”的整体用法。',10)
    para(doc,'B是另一个人，只用于动作评价与展示，未交给模型；不将B当理想结果计算PSNR、SSIM。这批数据规模小，1024每个数据集只有两对，不能据此证明泛化。结论只针对直接挂载BFS增强骨架编辑，不代表BFS官方换头、换人的全部能力。',10)
    para(doc,'正确性核验',11,True)
    checks=[]
    for method in METHODS[1:]:
        check=json.loads((root/'development'/method/'preflight.json').read_text())
        checks.append(f'{LABELS[method]}：零强度预测最大差 {check["zero_strength_max_abs"]:.6g}；强度1产生非零影响，完整采样通过。')
    for line in checks:para(doc,line,9)
    para(doc,'加载参数逐项核对；所有模型参数冻结；自训Adapter调用数为0；参考图严格为A和骨架；全部失败或低分结果保留。另外复跑一个案例的512、1024原生结果，与历史图片逐像素完全一致。完整指标、检测与人工核对、提示词及16组对照图另存于实验结果目录。',9)
    para(doc,'展示图读法',11,True)
    para(doc,'下列六组案例在实验前固定。每组第一行是整图，第二行是该图实际检测到的人脸放大，宽度与整图列一致；脸图只用于观察，没有额外输入模型。B列用于对照动作，三种生成结果中均没有输入B照片。',10)
    embedded=[]
    for i,(pair,res) in enumerate(SELECTION):
        if i%2==0:doc.add_page_break()
        panel=next(p for p in panels if p['pair_id']==pair and p['resolution']==res)
        if sha256(panel['path'])!=panel['sha256']:raise ValueError('Panel changed')
        para(doc,f'{pair} · {res}×{res}',11,True)
        caption=observations.get('captions',{}).get(f'{pair}_{res}','')
        if caption:para(doc,caption,9)
        doc.add_picture(panel['path'],width=Cm(18.1));embedded.append(panel)
    args.output.parent.mkdir(parents=True,exist_ok=True);doc.save(args.output)
    with zipfile.ZipFile(args.output) as archive:
        bad=archive.testzip()
        media=[x for x in archive.namelist() if x.startswith('word/media/')]
        if bad or len(media)!=6:raise AssertionError('Word embedded media validation failed')
        hashes={hashlib.sha256(archive.read(x)).hexdigest() for x in media}
        if hashes != {p['sha256'] for p in embedded}:raise AssertionError('Word images are not lossless panel copies')
    manifest={'document':str(args.output),'sha256':sha256(args.output),'embedded_figures':embedded,
              'media_count':len(media),'zip_valid':True,'observations':str(root/'evaluation/observations.json'),
              'summary_sha256':sha256(root/'evaluation/metric_summary.json')}
    (root/'word_report_manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2))
    print(json.dumps({'report':str(args.output),'images':len(media),'bytes':args.output.stat().st_size}),flush=True)


if __name__=='__main__':
    main()
