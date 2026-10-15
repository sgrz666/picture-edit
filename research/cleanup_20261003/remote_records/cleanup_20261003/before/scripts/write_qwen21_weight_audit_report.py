#!/usr/bin/env python3
"""Add provenance/semantic evidence and write the standalone Word audit report."""
from __future__ import annotations
from collections import Counter
from datetime import datetime
from pathlib import Path
import argparse
import hashlib
import json
import sys
from zoneinfo import ZoneInfo

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
import torch
from docx import Document
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from scripts.audit_qwen21_weight_compatibility import RUN, NIHAO, SOURCES, file_sha, read_checkpoint, write_json


def table(doc, headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Light Shading Accent 1'
    for cell, text in zip(t.rows[0].cells, headers): cell.text = str(text)
    repeat = OxmlElement('w:tblHeader')
    t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        for cell, text in zip(t.add_row().cells, row): cell.text = str(text)
    for row in t.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs: run.font.size = Pt(9)
    return t


def supplemental_evidence(root):
    summary = json.loads((root / 'summary.json').read_text())
    inventory = json.loads((root / 'parameter_inventory.json').read_text())
    manifest = json.loads((RUN / 'migration_manifest.json').read_text())
    copied = []
    for row in manifest['files']:
        a = ROOT / 'src/pose_control/v6' / row['relative_path']
        b = ROOT / 'src/pose_control/qwen21_full' / row['relative_path']
        copied.append({'file': row['relative_path'], 'sha256': file_sha(a),
                       'equal_to_copy_and_snapshot': file_sha(a) == file_sha(b) == row['source_sha256']})
    if not all(x['equal_to_copy_and_snapshot'] for x in copied):
        raise AssertionError('Preserved architecture changed; re-audit required')
    write_json(root / 'transitive_source_check.json', copied)
    cache = next(iter(sorted((NIHAO / 'qwen21_iper_cache').glob('source_*.pt'))))
    cached = read_checkpoint(cache)
    latent_shape = list(cached['reference_latent'].shape)
    if latent_shape != [1, 1024, 64]:
        raise AssertionError('Historical 512 cache evidence differs')
    manifests = [ROOT / 'experiments/qwen21_iper_body/full_20k/run_manifest.json',
                 NIHAO / 'qwen21_face/face_train_3k/manifest.json',
                 NIHAO / 'iper1024/body_3k_final/run_manifest.json',
                 NIHAO / 'iper1024/face_2500/manifest.json',
                 NIHAO / 'iper1024/continuation_manifest.json',
                 NIHAO / 'iper1024/body_3k_finetune/run_manifest.json']
    provenance = {
        'historical_records': [{'path': str(path), 'sha256': file_sha(path), 'content': json.loads(path.read_text())} for path in manifests],
        'resolution_512_evidence': {'checkpoint_resolution_field_missing': True,
                                   'runtime_default_resolution': 512,
                                   'runtime_source': str(ROOT / 'src/pose_control/qwen21/runtime.py'),
                                   'runtime_sha256': file_sha(ROOT / 'src/pose_control/qwen21/runtime.py'),
                                   'cache_source': str(cache), 'cache_sha256': file_sha(cache), 'reference_latent_shape': latent_shape,
                                   'interpretation': '512 is corroborated by historical cache root in both training manifests, default resolution and 32x32 native token grid, not by filename alone.'},
        'body1024_origin': 'Current checkpoint records resolution=1024, step=3000, output=body_3k_final. Its args record resume-to-self and init_checkpoint=None; the earlier finetune manifest names Body512 but does not by itself prove the entire intervening chain. No cumulative-step claim is made.',
        'face1024_origin': 'Checkpoint explicitly records init_checkpoint=Face512 selected 2500; current stage step=2500.',
        'selected_body1024_matches_final_best': file_sha(SOURCES['body1024']) == file_sha(NIHAO / 'iper1024/body_3k_final/checkpoint_best.pt'),
        'selected_face512_resolves_to': str(SOURCES['face512'].resolve()),
        'comparison_baseline_is_not_automatically_initialization': True,
    }
    write_json(root / 'provenance.json', provenance)
    target = read_checkpoint(RUN / 'face_512_100/checkpoint_last.pt')['adapter_state_dict']
    false_friends = []
    for row in inventory:
        if row['structurally_compatible']:
            continue
        family = ('legacy.face_preparer.', 'legacy.control_interface.face_adapter.') if row['source_id'].startswith('face') else ('legacy.control_interface.control_core.', 'output_bridges.geometry.')
        candidates = [name for name, value in target.items() if name.startswith(family) and list(value.shape) == row['shape']]
        false_friends.append({'source_id': row['source_id'], 'source': row['source'], 'shape': row['shape'],
                              'same_shape_target_count': len(candidates), 'same_shape_examples': candidates[:8],
                              'approved_mapping': False,
                              'reason': 'Learned feature basis, input semantics and/or surrounding operations differ. Shape coincidence is insufficient.'})
    write_json(root / 'nontransferable_shape_matches.json', false_friends)
    module_review = [
        ('几何编码器', '输入Normal、姿态热图、部位、人体系数与任务标记相同；默认关闭depth、采用native编码。', '可继承共有且有训练证据的部分；私有双人/接触编码默认保留接收方。'),
        ('Reasoner', '实现及配置相同；低分辨率、高分辨率特征、人员token、有效掩码均比较。', '单人主干可继承；不能把双人分支参数存在误认为已接受双人训练。'),
        ('条件转换', '两版都使用geometry_feature和geometry_highres，并通过geometry_high_downsample合并。', '几何转换及高分辨率降采样可继承；交互转换没有单人学习证据时保留原值。'),
        ('Body控制核心和六个输出头', '简化版在六层读取Qwen隐藏状态，由512维MLP输出4096维残差；完整版保留1536维旧注意力核心及六组旧头，再经新接口注入32层。', '不能按名字或形状直接替换；另行适配不在本轮范围。'),
        ('Face身份、纹理与融合', '简化版为ArcFace向量的LayerNorm/Linear及裁剪图CNN；完整版含DINO、Resampler、StableAnimator融合和细致几何。', '本轮没有找到可批准的直接加载项。512/1024的Face检查点都不进入白名单。'),
        ('Face阶段编码与输出头', '即使阶段编码同为6×512，也分属不同计算路径；简化版输出512→4096，完整版旧头512→1536后接1536→4096接口。', '不能因为都是六组、尺寸相近就直接继承。'),
        ('外观绑定、完整Face/Hand、双人控制核心、新Qwen接口', '简化版权重没有对应的完整实现，或没有对应计算意义的权重。', '保留qwen21_full接收方原值；不删除，不用简化模块替换。'),
    ]
    write_json(root / 'module_review.json', [{'module': a, 'evidence': b, 'decision': c} for a,b,c in module_review])
    return summary, inventory, module_review, provenance


def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument('--root', type=Path, default=RUN / 'weight_audit')
    p.add_argument('--out', type=Path, default=NIHAO / 'docs/Qwen简化版到完整版_权重兼容性核对.docx')
    args = p.parse_args()
    summary, inventory, modules, provenance = supplemental_evidence(args.root)
    if not summary['passed']:
        raise AssertionError('Cannot produce a passing report from failed local checks')
    sources = summary['sources']
    names = {'body512':'512 Body', 'face512':'512 Face', 'body1024':'1024 Body', 'face1024':'1024 Face'}
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(.7)
    section.left_margin = section.right_margin = Inches(.65)
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(10.5)
    style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '微软雅黑')
    doc.add_heading('简化版Qwen → qwen21_full\n权重兼容性核对', 0)
    doc.add_paragraph(datetime.now(ZoneInfo('Asia/Shanghai')).strftime('%Y年%m月%d日') + '｜只核对与局部验证，不训练、不出图')
    doc.add_heading('一、先说结论', level=1)
    doc.add_paragraph('之前的训练并没有白做：Body里共有的几何编码、Reasoner和几何条件转换权重可以利用。'
                      '但不能把整个简化版检查点直接塞进完整版；两版Body控制核心和Face架构不同。')
    table(doc, ['来源','检查点步数','结构兼容的参数表','建议加载的参数表','本轮结论'], [
        (names[key], value['step'], value['compatible_tensors'], value['recommended_tensors'],
         '共有模块可局部继承' if key.startswith('body') else '未找到可批准的直接继承项')
        for key,value in sources.items()])
    doc.add_paragraph('上表的“参数表”指一个权重张量，可能包含成千上万个数字，不是一个独立模块，也不是训练图片数量。'
                      '“结构兼容”只说明能合理加载；建议清单进一步排除了没有非零梯度证据的参数，并保留未训练的任务/角色行。')
    doc.add_paragraph('512和1024的Body都保留为候选。本轮没有比较画质，不能判定1024权重一定更好，也不能把兼容性通过说成Face彩斑已经解决。')
    doc.add_heading('二、具体能搬什么，不能搬什么', level=1)
    table(doc, ['模块','实际核对依据','处理建议'], modules)
    doc.add_paragraph('特别说明：简化版也使用了Reasoner的高分辨率几何特征。它们不是“因为简化就完全没有训练”的部分。'
                      '但完整架构还会使用人员token、外观绑定及双人等路径，因此继承共享模块后，整套系统的输出仍需另行验证。')
    doc.add_heading('三、哪些参数确实有训练证据', level=1)
    stats=[]
    for key,value in sources.items():
        e=value['training_evidence']
        comparable = sum(row['same_as_baseline'] is not None for row in inventory if row['source_id'] == key)
        difference = f"{value['changed_from_comparable_baseline']}/{comparable}" if comparable else '—（无可比基线）'
        stats.append((names[key],e.get('nonzero_gradient_evidence',0),e.get('zero_moments_only',0),e.get('no_optimizer_state',0),difference))
    table(doc,['来源','非零梯度证据','仅零动量记录','无优化器记录','不同数／可比数'],stats)
    doc.add_paragraph('这里仅读取AdamW留下的统计记录，没有恢复优化器、执行反向传播或更新参数。'
                      '非零二阶动量说明历史上有过非零梯度，不保证学到了有用能力；没有优化器记录，也不能断言该参数整个历史从未被训练。'
                      '零动量参数仍可能因权重衰减而变化，所以“数值变了”不等于“学到了东西”。')
    doc.add_paragraph('比较参考：512 Body与原DeepGen Body20k比较；1024 Body/Face与对应512权重比较。'
                      '这只是数值参照，不自动证明初始化关系。512 Face没有结构等价的旧Face初始化可比，上表该项不作有效变化结论。')
    embedding_rows=[]
    whitelists=json.loads((args.root/'load_whitelists.json').read_text())
    for key,value in whitelists.items():
        for entry in value['entries']:
            if 'rows' in entry:
                embedding_rows.append((names[key],entry['source'],str(entry['rows']),'其余行保留完整版原值'))
    table(doc,['来源','分行继承的参数表','加载行（从0编号）','其余行'],embedding_rows)
    doc.add_paragraph('共享几何编码器被用于单人和双人；即使保留了双人专用参数，换掉共享编码器也可能影响双人表现。'
                      '本轮双人合成输入仅验证计算兼容，不能当成双人训练或效果验收。')
    doc.add_heading('四、验证做到了哪一步', level=1)
    checks=[]
    for key,value in sources.items():
        if not key.startswith('body'): continue
        full=value['local_checks']; selected=value['whitelist_protection']['selective_single_person_checks']
        maximum=max(x['max_abs'] for case in full+selected for x in case['compared_fields'])
        checks.append((names[key],len(full),len(selected),f'{maximum:.8g}','白名单外参数及未选中的行保持不变'))
    table(doc,['候选','完整共有模块对照','建议子集对照','最大绝对差','保护检查'],checks)
    doc.add_paragraph('完整共有模块对照包含：单人、双人、混合人员有效性/接触条件，以及4对验证样本分别按512、1024入口读取。'
                      '逐级检查几何编码、Reasoner低/高分辨率输出、人员token与掩码、条件转换；CPU FP32容差固定为绝对1e-6、相对1e-5，未放宽。')
    doc.add_paragraph('另外，用随机latent检查完整架构在32×32和64×64目标token网格下的条件准备与外观绑定接口。'
                      '没有加载Qwen基座，也没有生成整图。两种RGB尺寸仍使用现有512几何图；本轮没有重建1024几何。'
                      '共享模块之后两版的空间处理与控制核心不同，本轮没有强行缩放不同输出来声称整条流程等价。')
    doc.add_paragraph('原47个控制源文件重新核对一致；所有被审计的检查点前后校验值一致；内存接收方最后恢复到原始状态。'
                      '最终50对测试集没有参与。两个Face检查点只完成参数、结构和训练证据审查，没有可批准映射，因此未做不合理的跨结构输出对齐。')
    doc.add_heading('五、来源记录及不能确认的部分', level=1)
    doc.add_paragraph('四份来源均记录了真实路径、符号链接指向、文件校验值和检查点内的步数。512检查点没有显式分辨率字段，'
                      '通过历史训练配置中的缓存位置、原流程默认512以及对应缓存的32×32参考token网格交叉确认，而非只看文件名。'
                      '1024检查点明确记录分辨率；Face明确记录从512已选权重初始化。')
    doc.add_paragraph('1024 Body的已选文件已与final目录的best文件核对；但当前配置被续训记录覆盖，显示“从本目录恢复”、初始检查点为空。'
                      '较早的微调配置记录了512起点，但不足以单独证明整个中间链条，因此这里不把3000步擅自累加成一个已证明的总训练步数。')
    doc.add_heading('六、后续若要尝试继承，按什么规则', level=1)
    doc.add_paragraph('先选择512 Body或1024 Body中的一份，按对应白名单整组加载共有模块里通过验证的参数；不混搭两份来源。'
                      '部分角色/任务参数只复制清单列明的行。完整控制核心、外观绑定、Face、Hand、双人专用参数和新Qwen接口保留接收方原值。'
                      '简化版Face暂不加载，不继承任何优化器状态，也不沿用旧检查点步数作为新训练进度。')
    doc.add_paragraph('本轮只交付未启用的白名单，没有保存混合权重。后续若执行真实迁移，需要另做验证图和训练实验；'
                      '当前Face彩斑仍需要独立排查。可加载不等于有效果，更不等于完整方案已优于之前版本。')
    doc.add_heading('交付文件',level=1)
    table(doc,['文件','用途'],[
        ('load_whitelists.json','每份来源的建议映射、允许复制的行、明确不加载的键；未启用'),
        ('parameter_inventory.json','四份检查点每个张量的形状、数据类型、训练证据及数值变化'),
        ('body512_audit.json / body1024_audit.json','逐样本、逐输出的局部对照及参数保护记录'),
        ('provenance.json / source_fingerprints.json','来源、历史配置、分辨率证据、文件校验值'),
        ('module_review.json / nontransferable_shape_matches.json','模块含义与同形状却不能直接迁移的案例'),
    ])
    doc.add_paragraph('上述机器可读文件保存在 nihaochen/qwen21_full/weight_audit。报告只展示重点，没有将数百行参数表全部铺开。')
    args.out.parent.mkdir(parents=True,exist_ok=True)
    doc.save(args.out)
    write_json(args.root/'report_manifest.json',{'path':str(args.out),'sha256':file_sha(args.out),'tables':len(doc.tables),
              'passed':True,'no_images_or_training':True,'face_direct_transfer_approved':False})
    print(json.dumps({'report':str(args.out),'tables':len(doc.tables)},ensure_ascii=False))


if __name__=='__main__':
    main()
