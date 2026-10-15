#!/usr/bin/env python3
"""Write a compact, image-rich Word report from fixed Face validation/test outputs."""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt
from PIL import Image


FACE_ROOT = Path('/home/shangguanrz/nihaochen/qwen21_face')
DEFAULT_OUT = Path('/home/shangguanrz/nihaochen/docs/Qwen21_面部保持优化_结果_20260928.docx')
MODES = ('body', 'body_crop_role', 'body_face', 'body_face_crop')
LABELS = {'body': 'Body', 'body_crop_role': 'Body＋A 脸图',
          'body_face': 'Body＋Face', 'body_face_crop': 'Body＋Face＋A 脸图'}


def arguments():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--validation', type=Path, default=FACE_ROOT / 'face_eval_2500_metrics/summary.json')
    parser.add_argument('--reference-probe', type=Path, default=FACE_ROOT / 'reference_probe_eval_strict/summary.json')
    parser.add_argument('--test', type=Path, default=FACE_ROOT / 'test_four_arm_metrics/summary.json')
    parser.add_argument('--groups', type=Path, default=FACE_ROOT / 'test_four_arm_metrics/group_summary.json')
    parser.add_argument('--cross', type=Path, default=FACE_ROOT / 'cross_person_eval/per_sample.json')
    parser.add_argument('--cross-panels', type=Path, default=FACE_ROOT / 'cross_person_eval')
    parser.add_argument('--panels', type=Path, default=FACE_ROOT / 'test_panels')
    parser.add_argument('--out', type=Path, default=DEFAULT_OUT)
    return parser.parse_args()


def metric(summary, name, mode, subset):
    return summary['metrics'][name][mode][subset]


def wrap_panel(path: Path, columns: int, tile_width: int) -> Path:
    """Give each comparison tile enough width on a portrait Word page."""
    output = path.with_name(path.stem + '_report_2rows.png')
    with Image.open(path) as source:
        source = source.convert('RGB')
        if source.width % tile_width:
            raise ValueError(f'panel width does not match tile width: {path}')
        tile_count = source.width // tile_width
        rows = (tile_count + columns - 1) // columns
        sheet = Image.new('RGB', (tile_width * columns, source.height * rows), 'white')
        for index in range(tile_count):
            tile = source.crop((index * tile_width, 0, (index + 1) * tile_width, source.height))
            sheet.paste(tile, ((index % columns) * tile_width, (index // columns) * source.height))
        sheet.save(output, optimize=True)
    return output


def add_table(doc, summary, title):
    doc.add_heading(title, level=2)
    doc.add_paragraph('ArcFace 源脸相似度只统计源脸、目标脸和四组生成脸共同可靠可见的样本；其余整图指标统计全部 50 对。')
    fields = [
        ('源脸相似度 ↑', 'identity_cosine', 'strict_joint', True),
        ('姿态 PCK@0.02 ↑', 'pose_pck_002', 'all', True),
        ('背景 MAE ↓', 'background_mae_source', 'all', False),
        ('PSNR ↑', 'psnr', 'all', True),
    ]
    table = doc.add_table(rows=1, cols=1 + len(MODES))
    table.style = 'Table Grid'
    for cell, label in zip(table.rows[0].cells, ['指标', *(LABELS[x] for x in MODES)]):
        cell.text = label
    for label, name, subset, higher in fields:
        row = table.add_row().cells
        row[0].text = label
        values = [metric(summary, name, mode, subset)['mean'] for mode in MODES]
        best = max(values) if higher else min(values)
        for index, (mode, value) in enumerate(zip(MODES, values), start=1):
            n = metric(summary, name, mode, subset)['n']
            text = f'{value:.3f}' + (f' (n={n})' if name == 'identity_cosine' else '')
            run = row[index].paragraphs[0].add_run(text)
            run.bold = abs(value - best) < 1e-10


def main():
    args = arguments()
    val = json.loads(args.validation.read_text())
    reference = json.loads(args.reference_probe.read_text())
    test = json.loads(args.test.read_text())
    test_rows = json.loads((args.test.parent / 'per_sample.json').read_text())
    groups = json.loads(args.groups.read_text())
    cross = json.loads(args.cross.read_text())
    doc = Document()
    normal = doc.styles['Normal']
    normal.font.name = 'Arial'
    normal.font.size = Pt(10)
    doc.add_heading('Qwen-Image-2.1 面部保持优化：阶段结果', 0)
    doc.add_paragraph('2026 年 9 月 28 日｜单人 iPER｜30 步、512×512、种子 42')
    doc.add_heading('先看懂四组结果', level=1)
    doc.add_paragraph('A 是要编辑的原图：希望保留 A 中的人、衣服和背景。B 展示希望模仿的动作；系统从 B 提取动作和头部位置等几何信息，不把 B 的脸当作人物身份参考。四组都使用同一个 Qwen21 基座和已有的 Body 身体控制，区别在于是否额外使用 Face 和 A 脸图。')
    doc.add_paragraph('Body Adapter 是已有的身体姿态控制模块，负责让 A 中的人摆出 B 的动作。Face Adapter 是本轮另外训练的面部控制模块：参考 A 的脸，并结合 B 的目标头部位置，帮助生成结果保留 A 的面貌。A 脸图是从 A 单独裁出的人脸照片，额外作为 Qwen21 的第二张参考图片；添加它本身不需要训练新模块。')
    doc.add_paragraph('Face Adapter 确实另外训练了六个残差输出头，与 Body Adapter 分别接在 Qwen21 的同六个层后：层索引 3、8、13、18、23、28（从 0 开始）。每层的 Body 残差与 Face 残差分别计算，再作用于生成图；Face 残差仅在目标脸部区域生效。训练 Face 时，Qwen21 基座与 Body Adapter 保持冻结。')
    doc.add_paragraph('注意：Face 分支内部也会使用 A 的人脸特征，但表中“＋A 脸图”专指把脸部裁剪图额外送入 Qwen21 原有的图片参考通道。因此，“Body＋Face”仍只给该通道一张完整 A 图。')
    doc.add_paragraph('双图方案在推理时把［完整 A 图，A 的脸部裁剪图］一起传给 Qwen21，并配一段说明两张图角色的英文提示词。提示词中的“目标动作几何”实际由 Body Adapter 提供；B 的照片和脸没有作为 Qwen21 的图片参考。')
    doc.add_paragraph('双图提示词原文：Image 1 is the full source photo. Image 2 is a crop of the same person’s face from image 1. Change the person in image 1 to match the supplied target pose geometry. Preserve the identity and facial features from image 2, and preserve the clothing and background from image 1.')
    doc.add_paragraph('单图组的提示词：Change the person pose to match the target pose. Preserve identity, clothing, and background.')
    explanation = doc.add_table(rows=1, cols=2)
    explanation.style = 'Table Grid'
    explanation.rows[0].cells[0].text = '结果名称'
    explanation.rows[0].cells[1].text = '实际给了模型什么、启用了什么'
    for name, description in (
        ('Body', '完整 A 图＋B 的动作几何；只启用已有身体控制。'),
        ('Body＋A 脸图', 'Body 的条件＋从 A 裁出的人脸图，作为第二张图片参考；不启用 Face 分支。'),
        ('Body＋Face', 'Body 的条件＋新训练的 Face 分支；Qwen21 的图片参考仍只有完整 A 图。'),
        ('Body＋Face＋A 脸图', '同时启用 Body、Face 分支，并额外提供 A 的人脸图作为第二张图片参考。'),
    ):
        cells = explanation.add_row().cells
        cells[0].text = name
        cells[1].text = description
    doc.add_paragraph('若 A 的脸无法可靠识别，或目标脸不可见，面部增强会跳过并退回 Body。下文的真实目标图只用于训练监督和离线评分，推理时不会把它交给生成模型。')
    doc.add_heading('实验设置', level=1)
    doc.add_paragraph('冻结 Qwen21 基座与既有 Body 20k，单独训练 Face 残差分支。Face 身份和纹理来自 A 图；B 只提供姿态与头部几何。无有效脸时退回 Body。训练外观 69 种，验证外观 10 种；原有 50 对测试只在验证选定方案后使用。')
    doc.add_paragraph('额外 A 脸图先单独测试，再与 Face 分支合用。Face 在第 3、8、13、18、23、28 层注入目标脸区域；本轮选择 2.5k 检查点。')
    ref_same = metric(reference, 'identity_cosine', 'body_crop_same', 'strict_joint')
    ref_role = metric(reference, 'identity_cosine', 'body_crop_role', 'strict_joint')
    doc.add_paragraph(f'仅加 A 脸图的独立实验：同一提示词时源脸相似度 {ref_same["mean"]:.3f}，明确说明两图同源后为 {ref_role["mean"]:.3f}（共同可靠脸 n={ref_role["n"]}）。')
    add_table(doc, val, '验证集：50 对')
    doc.add_paragraph('验证集 2.5k Face＋脸图相对 Body：可靠脸源身份余弦 +0.053；姿态 PCK 不变；背景 MAE 增加约 0.006。3k 检查点身份指标较低，因此提前停止。减弱 Body 脸部残差的验证消融没有超过所选方案。')
    add_table(doc, test, '最终测试集：50 对')
    identity = groups['groups']['reliable_visible_face']['comparisons']['body_face_crop']['identity_cosine']
    source_target = [row['source_target_identity_cosine'] for row in test_rows
                     if row['strict_reference_usable'] and row['source_target_identity_cosine'] is not None]
    body_pose = metric(test, 'pose_pck_002', 'body', 'all')['mean']
    selected_pose = metric(test, 'pose_pck_002', 'body_face_crop', 'all')['mean']
    body_background = metric(test, 'background_mae_source', 'body', 'all')['mean']
    selected_background = metric(test, 'background_mae_source', 'body_face_crop', 'all')['mean']
    doc.add_paragraph(f'选定方案相对 Body：可靠脸源身份余弦 +{identity["delta"]:.3f}（n={identity["n"]}）；姿态 PCK {body_pose:.3f}→{selected_pose:.3f}；背景 MAE {body_background:.3f}→{selected_background:.3f}。同一批可靠脸的 Source–真实 Target 余弦约为 {sum(source_target)/len(source_target):.3f}，生成脸仍有明显差距。')
    crop_increment = []
    for row in test_rows:
        if row['strict_reference_usable']:
            a, b = row['variants']['body_crop_role'], row['variants']['body_face_crop']
            if a['identity_cosine'] is not None and b['identity_cosine'] is not None:
                crop_increment.append(b['identity_cosine'] - a['identity_cosine'])
    doc.add_paragraph(f'Face 在双图方案上再增加约 {sum(crop_increment)/len(crop_increment):.3f} 的源脸余弦；这项增量较小，在测试集上不能作为已证实的稳定收益。')
    doc.add_heading('测试分组与不同人 A/B 检查', level=2)
    seen = groups['groups']['reliable_face_person_seen']
    unseen = groups['groups']['reliable_face_person_unseen']
    doc.add_paragraph(f"测试外观均未进入 Face 训练，但人物编号可能重复。50 对中已见人物编号 39 对，未见人物编号 11 对；其中可靠脸分别为 {seen['pairs']} 对和 {unseen['pairs']} 对。样本少的组不作强结论。")
    for group_name, data in (('已见人物编号', seen), ('未见人物编号', unseen)):
        comparison = data['comparisons']['body_face_crop']['identity_cosine']
        doc.add_paragraph(f"{group_name}：Face＋脸图相对 Body 的可靠脸源身份余弦差值 {comparison['delta']:+.3f}，n={comparison['n']}。" if comparison['delta'] is not None else f'{group_name}：可靠脸不足，无法判断。')
    valid_cross = [row for row in cross if row['face_valid']]
    source_over_b = sum(row['scores']['body_face_crop']['source_cosine'] > row['scores']['body_face_crop']['B_face_cosine'] for row in valid_cross)
    doc.add_paragraph(f'另检查了 {len(cross)} 组不同人物编号的 A/B 案例，其中 Face 有效 {len(valid_cross)} 组；选定方案有 {source_over_b}/{len(valid_cross)} 组生成脸与 A 的余弦高于与 B 的余弦。B 图像只用于展示；生成条件来自 A 的图像与身份、B 的几何。这是小规模串脸检查，并非证明绝无串脸。')
    old = groups['groups']['prior_audit_all_front_32']['comparisons']['body_face_crop']['identity_cosine']
    doc.add_paragraph(f'沿用旧面部重评的六图均正面、短边≥32 像素子集时，选定方案相对 Body 的源脸余弦增量为 {old["delta"]:+.3f}（n={old["n"]}）。明显转头且满足本轮严格可靠脸门槛的样本很少，应以并排图逐例查看。')
    doc.add_heading('典型样本：全图及脸部放大', level=1)
    for index in (0, 10, 42):
        matches = sorted(args.panels.glob(f'{index:03d}_*.jpg'))
        if matches:
            doc.add_paragraph(f'测试样本 {index:03d}')
            doc.add_picture(str(wrap_panel(matches[0], 3, 384)), width=Inches(6.3))
    cross_picture = args.cross_panels / '000_comparison.png'
    if cross_picture.exists():
        doc.add_paragraph('不同人 A/B 案例 000：B 仅提供目标几何')
        doc.add_picture(str(wrap_panel(cross_picture, 3, 512)), width=Inches(6.3))
    doc.add_heading('结论范围', level=1)
    doc.add_paragraph('面部相似度受转头、背面、小脸和检测误差影响；数值结论仅用于共同可靠可见脸子集。脸部放大图仍需逐例查看表情、眼镜及头向。iPER 的测试外观未见过，不等于所有人物身份未见过。手部与双人交互不在本轮范围。')
    args.out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.out)
    print(args.out)


if __name__ == '__main__':
    main()
