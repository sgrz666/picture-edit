#!/usr/bin/env python3
"""Word report with embedded full-pixel panels and explicit training-data labels."""
import json
from pathlib import Path
import sys
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
from PIL import Image,ImageDraw,ImageFont
from docx import Document
from docx.enum.section import WD_ORIENT,WD_SECTION
from docx.shared import Inches
from scripts.run_qwen21_face_single_sample import OUT
from scripts.write_qwen21_full_report import configure_doc,add_table

def read(name,default=None):
    p=OUT/name
    return json.loads(p.read_text()) if p.exists() else default

def number(x):return '无法判断' if x is None else f'{x:.3f}'

def panel(name,rows,labels):
    side=512;header=48;gap=8
    canvas=Image.new('RGB',((side+gap)*len(rows)-gap,side*2+header*2),'white')
    draw=ImageDraw.Draw(canvas);font=ImageFont.truetype('/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc',22)
    for i,(row,label) in enumerate(zip(rows,labels)):
        x=i*(side+gap);image=Image.open(row['image']).convert('RGB')
        draw.text((x+5,10),label,font=font,fill='black')
        canvas.paste(image.resize((side,side),Image.Resampling.LANCZOS),(x,header))
        info=row.get('face_detection',{}).get('generated',{})
        idx=info.get('raw_largest_index')
        if idx is not None:
            box=info['detections'][idx]['box'];box=[v/512 for v in box]
        else:box=row['face_boxes_normalized']['target']
        cx,cy=(box[0]+box[2])/2,(box[1]+box[3])/2
        extent=max(box[2]-box[0],box[3]-box[1])*1.35
        crop=image.crop(tuple(round(v) for v in ((cx-extent/2)*image.width,(cy-extent/2)*image.height,
                                                (cx+extent/2)*image.width,(cy+extent/2)*image.height)))
        canvas.paste(crop.resize((side,side),Image.Resampling.LANCZOS),(x,header*2+side))
        draw.text((x+5,header+side+10),'对应脸部放大｜非额外生成',font=font,fill='black')
    directory=OUT/'report_panels';directory.mkdir(exist_ok=True)
    png=directory/f'{name}.png';canvas.save(png)
    jpg=directory/f'{name}.jpg';canvas.save(jpg,quality=95,subsampling=0,optimize=True)
    return jpg

def metric_table(doc,rows,labels):
    fields=[('姿态 PCK ↑','pose_pck_002'),('背景误差 ↓','background_mae_source'),('源脸相似度（原始）↑','identity_cosine_raw')]
    values=[]
    for label,key in fields:values.append([label,*[number(r['metrics'].get(key)) for r in rows]])
    values.append(['可靠源脸相似度 ↑',*[f"{number(r['metrics'].get('identity_cosine_reliable'))}（n={int(r['metrics'].get('identity_cosine_reliable') is not None)}）" for r in rows]])
    table=add_table(doc,['指标',*labels],values,font_size=8)
    for j,(_,key) in enumerate(fields,1):
        numbers=[r['metrics'].get(key) for r in rows];valid=[v for v in numbers if v is not None]
        if valid:
            best=min(valid) if key=='background_mae_source' else max(valid)
            for i,value in enumerate(numbers,1):
                if value==best:
                    for p in table.rows[j].cells[i].paragraphs:
                        for run in p.runs:run.bold=True
    values=[r['metrics'].get('identity_cosine_reliable') for r in rows]
    valid=[v for v in values if v is not None]
    if valid:
        for i,value in enumerate(values,1):
            if value==max(valid):
                for paragraph in table.rows[len(fields)+1].cells[i].paragraphs:
                    for run in paragraph.runs:run.bold=True

def main():
    selection=read('selection.json');decision=read('training_decision.json')
    if not selection or not decision:raise RuntimeError('Selection and completed training decision required')
    diagnosis=read('diagnosis_metrics.json',[]);body=read('body_stage_metrics.json',[]);face=read('face_stage_metrics.json',[])
    visual=read('visual_review.json',{})
    doc=Document();configure_doc(doc)
    doc.add_heading('Qwen 完整版：面部彩斑排查与单样本训练',0)
    doc.add_paragraph('本轮检查：面部彩斑来自哪里，以及完整 Face Adapter 能否在一对反复训练的图片上学会正常生成。Body Adapter 负责动作，Face Adapter 负责身份、脸部外观与目标面部几何；Qwen 基座始终冻结。')
    pair=selection['selected'];doc.add_paragraph(f"本次使用训练集中的 {pair['appearance']}：A={pair['source']['stem']}，Target={pair['target']['stem']}。这是训练样本上的学习诊断，展示图片不是未见过的测试结果，不能说明泛化能力。最终 50 对测试集未使用。")
    doc.add_paragraph('所有 512 展示均用 30 步采样、种子 42、true CFG 1.0。图中的 100／500／1000 等是新增训练更新次数；每个检查点都重新完成同样的 30 步生成。输入完整 A 图，不增加第二张 A 脸图，Target RGB 不进入推理条件。')
    doc.add_paragraph('Face 训练从未训练过 Qwen 面部输出接口的 Body 检查点开始：旧版完整 Face 模块仍保留，新增输出接口最初为零。因此 Face 0 步应当等于冻结后的 Body，能直接观察 Face 学习是否带来额外收益。此次没有混入简化版 Qwen 的继承候选。')
    doc.add_paragraph('Body 从已有完整版 Body 100 步权重出发，图中的“新增”不包含此前的训练。这里训练的是两个 Adapter 及其接口，Qwen 大模型本身没有重新训练。')
    doc.add_paragraph('本轮保留旧版完整 Face 的身份、外观、面部几何编码、融合模块和六组残差头，另用输出接口连接 Qwen。Face 训练前 20 步只学习输出接口，其后再学习原 Face 分支；Body、共享输入接口和基座在 Face 阶段都不更新。')
    doc.add_heading('目前得到的结论',1)
    for text in visual.get('conclusions',[]):doc.add_paragraph(text)
    if not decision['body_gate_passed']:
        doc.add_paragraph('Body 在允许的训练上限内未通过动作与头部位置门槛，因此按计划未启动 Face 2000 步长训。当前结果不能回答 Face 长训是否能够学会；它说明这对图片的 Body 位置问题尚未排除。')
    else:
        doc.add_paragraph(f"Body 已通过位置门槛，新增训练 {decision['body_added_steps']} 步；Face 实际新增 {read('face_train/summary.json',{}).get('optimizer_steps',0)} 步。是否成功还需结合完整生成图、姿态与背景约束判断。")
    doc.add_heading('彩斑诊断：每次只改变一个因素',1)
    if diagnosis:
        labels=['Body','当前 Face','Face×0.25','噪声进度','关闭缓存','进度＋关缓存']
        metric_table(doc,diagnosis,labels)
        doc.add_paragraph('所有诊断列均从相同的当前完整 Face 100 步检查点出发。降低强度、统一控制时机、关闭缓存用于诊断；组合列单独标明，不能把它当成单因素实验。新训练和生成统一使用噪声进度，Face 强度仍为 1。')
    doc.add_heading('训练过程与正确性',1)
    audit=read('evaluator_input_audit.json',{})
    if audit:
        doc.add_paragraph('排查发现：旧 DeepGen Face 调用与此前完整版 Face 短训，把 [0,1] 脸图传给要求 [-1,1] 的身份、感知网络。新实验修正这个输入约定，保留网络权重和损失系数。因此新旧结果同时包含输入修复、控制时机统一和训练增加的影响，不能把变化全部归因于训练步数。')
        doc.add_paragraph(f"同一固定 RGB 测试图：正确输入的 TorchScript 身份特征与原 ONNX 特征余弦相似度为 {audit['normalized_TorchScript_vs_ONNX_cosine']:.6f}；旧输入方式为 {audit['legacy_zero_one_vs_ONNX_cosine']:.6f}。这是输入接口核对，不是生成脸身份得分。")
    statuses=read('correctness.json',{})
    add_table(doc,['检查','结果'],[
        ['原生关闭控制、目标 token 和 Face 区域边界','通过' if statuses.get('passed') else '未完成'],
        ['身份与 DINO 外观来自 A、面部几何来自 B','通过' if statuses.get('passed') else '未完成'],
        ['身份与感知损失实际传到 Face 输出接口','通过' if statuses.get('passed') else '未完成'],
        ['初始零 Face 输出等同 Body','通过' if statuses.get('initial_face_vs_body',{}).get('equal') else '未完成'],
        ['最终测试集','本轮未使用']])
    for branch in ('body','face'):
        summary=read(f'{branch}_train/summary.json',{})
        if summary:
            doc.add_paragraph(f"{branch.title()}：{summary['optimizer_steps']} 次新增优化更新，唯一训练对数 {summary['unique_training_pairs']}；训练耗时 {summary['seconds']/60:.1f} 分钟，峰值显存 {summary['peak_cuda_gib']:.1f} GiB。冻结参数检查：{'通过' if summary['frozen_parameters_unchanged'] and summary['base_versions_unchanged'] else '未通过'}。")
            probe_files=sorted((OUT/f'{branch}_train').glob('diagnostic_probe_*.json'))
            add_table(doc,['新增训练步','固定 8 组噪声的平均生成损失 ↓'],[[json.loads(p.read_text())['step'],number(json.loads(p.read_text())['mean_flow'])] for p in probe_files])
    regional=read('regional_probe.json',[])
    if regional:
        doc.add_paragraph('为避免整图损失掩盖小面积的人脸问题，还在完全相同的 8 组噪声下单独检查脸区和脸外区域。下面是训练式单步诊断，输入含带噪的监督 Target，不能当成从 A 编辑出的完整图片。')
        add_table(doc,['Face 新增训练步','脸区预测误差 ↓','脸外预测误差 ↓'],[
            [r['added_step'],number(sum(v['face_region_flow_mse'] for v in r['fixed_noise_rows'])/8),
             number(sum(v['nonface_region_flow_mse'] for v in r['fixed_noise_rows'])/8)] for r in regional])
    recovery=read('recovery_summary.json',{})
    if recovery:
        doc.add_paragraph(recovery['explanation'])
    doc.add_paragraph('完整生成的耗时和显存逐张保存在数据记录中。部分中间检查点在训练继续运行时出图，资源负载不同，因此这里不把各训练步的出图耗时当成速度优劣比较。')
    if body:metric_table(doc,body,[f"Body {r['added_step']}" for r in body])
    if face:metric_table(doc,face,[f"Face {r['added_step']}" for r in face])
    doc.add_paragraph('n 表示能够可靠计算身份指标的样本数，本轮每列最多 1。检测不到可靠脸时写“无法判断”，不把检测失败换成零相似度。脸图是对应整图的裁剪和放大，放大不会增加原始细节。')
    doc.add_heading('局限与下一步',1)
    doc.add_paragraph('单样本成功只证明这一例能够学会；单样本失败也不能直接断言整个结构无效。保持本轮上限，不根据这些结果扩大训练、不替换默认权重。旧 DeepGen、简化 Qwen、完整版历史权重均保留。')
    for text in visual.get('next_steps',[]):doc.add_paragraph(text)
    doc.add_paragraph(visual.get('resolution_1024_explanation','1024 检查只在 512 面部验收通过后进行。'))
    section=doc.add_section(WD_SECTION.NEW_PAGE);section.orientation=WD_ORIENT.LANDSCAPE
    section.page_width=Inches(11.69);section.page_height=Inches(8.27)
    section.left_margin=section.right_margin=Inches(.35)
    pictures=[]
    groups=[]
    if diagnosis:
        anchor=diagnosis[0];reference_rows=[]
        for which in ('source','target'):
            reference_rows.append({**anchor,'image':anchor[which],
                                   'face_detection':{'generated':anchor['face_detection'][which]},
                                   'face_boxes_normalized':{'target':anchor['face_boxes_normalized'][which]}})
        groups.append(('reference',reference_rows+[anchor]+([body[-1]] if body else []),
                       ['A：身份与背景','Target：监督与目标几何','训练前 Body']+(['冻结后的 Body'] if body else [])))
    if diagnosis:
        groups.extend([('diagnosis_1',diagnosis[:3],['Body','当前 Face 100步','Face 强度 0.25']),
                       ('diagnosis_2',[diagnosis[i] for i in (1,3,4,5)],['当前 Face','噪声进度','关闭缓存','噪声进度＋关缓存'])])
        overlays=[{**r,'image':str(OUT/'diagnosis'/f"{r['case']}_overlay.png")} for r in diagnosis[:2]]
        groups.append(('diagnosis_alignment',overlays,['Body：红为目标，绿为检测','当前 Face：同样的目标脸区']))
    if body:groups.append(('body_training',body,[f"Body 新增 {r['added_step']}步" for r in body]))
    if face:
        groups.append(('face_training_early',face[:3],[f"Face 新增 {r['added_step']}步" for r in face[:3]]))
        groups.append(('face_training_late',face[3:],[f"Face 新增 {r['added_step']}步" for r in face[3:]]))
    for i,(name,rows,labels) in enumerate(groups):
        if i:doc.add_page_break()
        doc.add_heading({'reference':'固定训练案例与 Body 输出','diagnosis_1':'Face 开关与强度对照','diagnosis_2':'控制时机与缓存对照','body_training':'同一训练样本：Body 学习过程',
                         'diagnosis_alignment':'脸部位置核对：红色目标框与绿色检测框',
                         'face_training_early':'同一训练样本：Face 早期学习过程','face_training_late':'同一训练样本：Face 后期学习过程'}[name],1)
        picture=panel(name,rows,labels);pictures.append(str(picture))
        with Image.open(picture) as pixels:
            width=min(10.8,6.3*pixels.width/pixels.height)
        doc.add_picture(str(picture),width=Inches(width))
    destination=Path('/home/shangguanrz/nihaochen/docs/Qwen完整版_面部彩斑排查与单样本训练.docx')
    doc.save(destination)
    check=Document(destination)
    if len(check.inline_shapes)!=len(pictures):raise AssertionError('Embedded picture count differs')
    (OUT/'report_manifest.json').write_text(json.dumps({'word':str(destination),'embedded_panels':pictures,'bytes':destination.stat().st_size},ensure_ascii=False,indent=2))
    print(destination)

if __name__=='__main__':main()
