#!/usr/bin/env python3
"""Write the bounded Face repair report from saved evidence; no model execution."""
from __future__ import annotations

import argparse
import hashlib
import json
import math
from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.shared import Inches, Pt

from scripts.write_qwen21_full_report import add_table, configure_doc

DEFAULT_ROOT = Path('/home/shangguanrz/nihaochen/qwen21_full/face_repair')
DEFAULT_PREVIOUS = DEFAULT_ROOT.parent / 'face_single_sample'
DEFAULT_OUTPUT = Path('/home/shangguanrz/nihaochen/docs/Qwen完整版_Face彩斑修复验证.docx')
FONT = '/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc'
RUNS = ('roi_same_aux', 'roi_aligned_aux')


def read(path, default=None):
    path = Path(path)
    return json.loads(path.read_text()) if path.is_file() else default


def finite(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value)


def number(value, digits=3):
    return f'{value:.{digits}f}' if finite(value) else '无法判断'


def by_step(rows, step):
    matching = [row for row in rows if row.get('added_step') == step]
    if len(matching) != 1:
        raise ValueError(f'Expected one record at added step {step}, found {len(matching)}')
    return matching[0]


def bold_best(table, table_row, values, *, higher):
    valid = [value for value in values if finite(value)]
    if not valid:
        return
    best = max(valid) if higher else min(valid)
    for index, value in enumerate(values, 1):
        if finite(value) and value == best:
            for paragraph in table.rows[table_row].cells[index].paragraphs:
                for run in paragraph.runs:
                    run.bold = True


def metric_table(doc, rows, labels):
    fields = (
        ('源脸相似度（可靠）↑', 'identity_cosine_reliable', True),
        ('源脸相似度（原始）↑', 'identity_cosine_raw', True),
        ('姿态 PCK ↑', 'pose_pck_002', True),
        ('背景保持误差 ↓', 'background_mae_source', False),
        ('整图 PSNR ↑', 'psnr', True),
        ('整图 SSIM ↑', 'ssim', True),
    )
    data = []
    for label, key, _ in fields:
        values = []
        for row in rows:
            value = row['metrics'].get(key)
            text = number(value)
            if key == 'identity_cosine_reliable':
                text += f'（n={int(finite(value))}）'
            values.append(text)
        data.append([label, *values])
    table = add_table(doc, ['指标', *labels], data, font_size=8)
    for index, (_, key, higher) in enumerate(fields, 1):
        bold_best(table, index, [row['metrics'].get(key) for row in rows], higher=higher)
    return table


def panel(directory, name, rows, labels):
    """Keep each 512 whole image and its enlarged crop in equal-width columns."""
    side, header, gap = 512, 48, 8
    canvas = Image.new('RGB', ((side + gap) * len(rows) - gap, 2 * (side + header)), 'white')
    draw = ImageDraw.Draw(canvas)
    font = ImageFont.truetype(FONT, 21)
    for index, (row, label) in enumerate(zip(rows, labels, strict=True)):
        x = index * (side + gap)
        with Image.open(row['image']) as opened:
            image = opened.convert('RGB')
        draw.text((x + 5, 10), label, font=font, fill='black')
        canvas.paste(image.resize((side, side), Image.Resampling.LANCZOS), (x, header))
        info = row.get('face_detection', {}).get('generated', {})
        detection_index = info.get('reliable_index')
        if detection_index is None:
            detection_index = info.get('raw_largest_index')
        if detection_index is not None:
            # All metrics in this experiment use the common 512 image grid.
            box = [float(value) / 512 for value in info['detections'][detection_index]['box']]
        else:
            box = row['face_boxes_normalized']['target']
        cx, cy = (box[0] + box[2]) / 2, (box[1] + box[3]) / 2
        extent = max(box[2] - box[0], box[3] - box[1]) * 1.35
        bounds = tuple(round(value) for value in (
            (cx - extent / 2) * image.width, (cy - extent / 2) * image.height,
            (cx + extent / 2) * image.width, (cy + extent / 2) * image.height))
        crop = image.crop(bounds)
        draw.text((x + 5, header + side + 10), '对应脸部放大｜非额外生成', font=font, fill='black')
        canvas.paste(crop.resize((side, side), Image.Resampling.LANCZOS), (x, 2 * header + side))
    directory.mkdir(parents=True, exist_ok=True)
    png, jpg = directory / f'{name}.png', directory / f'{name}.jpg'
    canvas.save(png)
    canvas.save(jpg, quality=95, subsampling=0, optimize=True)
    return png, jpg


def reference_row(anchor, which):
    row = dict(anchor, image=anchor[which])
    row['face_detection'] = {'generated': anchor['face_detection'][which]}
    row['face_boxes_normalized'] = {'target': anchor['face_boxes_normalized'][which]}
    return row


def probe_mean(probe, key):
    values = [row.get(key) for row in probe.get('rows', [])]
    valid = [value for value in values if finite(value)]
    return sum(valid) / len(valid) if valid and len(valid) == len(values) else None


def paragraphs(doc, values):
    if isinstance(values, str):
        values = [values]
    if isinstance(values, list):
        for value in values:
            if isinstance(value, str):
                doc.add_paragraph(value)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--root', type=Path, default=DEFAULT_ROOT)
    parser.add_argument('--previous', type=Path, default=DEFAULT_PREVIOUS)
    parser.add_argument('--output', type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    root, previous = args.root, args.previous
    body = by_step(read(previous / 'body_stage_metrics.json', []), 500)
    old_face_rows = read(previous / 'face_stage_metrics.json', [])
    old300, old2000 = by_step(old_face_rows, 300), by_step(old_face_rows, 2000)
    samples, summaries = {}, {}
    for run in RUNS:
        summaries[run] = read(root / run / 'summary.json')
        if not summaries[run] or summaries[run].get('optimizer_steps') != 300:
            raise RuntimeError(f'Completed 300-step training evidence required for {run}')
        for step in (100, 300):
            path = root / 'generations' / run / f'step_{step:06d}' / 'metrics.json'
            row = read(path)
            if not row or not Path(row['image']).is_file():
                raise RuntimeError(f'Complete saved generation and metrics required: {path}')
            if row.get('control_forwards') != 30:
                raise RuntimeError(f'Expected complete 30-step sampling: {path}')
            samples[(run, step)] = row
    roi100, roi300 = samples[('roi_same_aux', 100)], samples[('roi_same_aux', 300)]
    aligned100, aligned300 = samples[('roi_aligned_aux', 100)], samples[('roi_aligned_aux', 300)]
    review = read(root / 'final_review.json', {})
    acceptance = read(root / 'acceptance.json', {})
    protection = read(root / 'provenance.json', {})
    selection = read(previous / 'selection.json', {})
    high_resolution_root = root / 'generations_1024' / 'roi_same_aux'
    high_resolution_paths = {
        'body': high_resolution_root / 'step_000000' / 'body' / 'metrics.json',
        'native_grid': high_resolution_root / 'step_000300' / 'metrics.json',
        'checkpoint_grid': high_resolution_root / 'step_000300' / 'checkpoint_grid' / 'metrics.json',
    }
    high_resolution = {name: read(path) for name, path in high_resolution_paths.items()}
    for name, row in high_resolution.items():
        if row is not None and (row.get('resolution') != 1024 or row.get('control_forwards') != 30
                                or not Path(row['image']).is_file()):
            raise RuntimeError(f'Incomplete or mismatched 1024 generation record: {high_resolution_paths[name]}')
    scatter_audit_path = root / 'resolution_face_scatter_audit.json'
    scatter_audit = read(scatter_audit_path, {})
    peak_ratios = [row.get('native64_peak_vs32') for row in scatter_audit.get('rows', [])]
    peak_ratios = [ratio for ratio in peak_ratios if finite(ratio)]
    highest_peak_ratio = max(peak_ratios) if peak_ratios else None
    docs = Document()
    configure_doc(docs)
    for section in docs.sections:
        section.footer.paragraphs[0].text = 'Qwen 完整版｜Face 彩斑修复验证'
    docs.add_heading('Qwen 完整版：Face 彩斑修复验证', 0)
    docs.add_paragraph('这次工作针对一个明确问题：仅使用 Body 时脸部基本正常，加入 Face 后反而出现彩斑。我们保留完整控制结构，先修正“训练到底重视什么”，再用同一对图片检查是否真正恢复正常。')
    if acceptance.get('512_results'):
        docs.add_paragraph('本次结果：512 下两组修复训练均改善了这个案例；1024 下原先出现的绿色斑点和明显重影，在修正 Face 控制映射后消失。脸部细节仍偏糊，其他人物和转头情况尚待验证。')
    docs.add_heading('先说明每种情况做了什么', 1)
    add_table(docs, ['名称', '实际做法'], [
        ['Body', 'Body Adapter 负责把 A 中人物的动作调整为 B 的几何动作，同时保留 A 的外观与背景。本轮固定已有单样本 Body 500 步权重。'],
        ['旧 Face', '完整 Face Adapter 从 A 的人脸身份和外观特征、B 的面部几何学习脸部控制；此前训练后的结果出现彩斑。'],
        ['ROI Face', '仍然训练同一个完整 Face Adapter，但将目标脸区的生成误差单独计算，避免小面积的脸部问题被整图平均掩盖。'],
        ['对齐 ROI Face', '在 ROI Face 的基础上，把训练时送入身份与感知网络的脸部，按五个关键点统一位置与角度。这个操作用于计算训练监督，Qwen 的原生参考仍只有完整 A 图。'],
    ], font_size=9)
    docs.add_paragraph('Body 和 Face 都是 Adapter。Qwen 大模型基座没有重新训练，Face 阶段的 Body 也保持冻结。原完整 Face 的身份编码、DINO 外观、面部几何、融合模块和六组残差头都保留，六组覆盖 Qwen 32 层的方式也没有改变。')
    docs.add_paragraph('本轮两组各新增训练 300 步，从同一个固定 Body 和零输出的 Face 接口独立开始，没有把第一组的权重叠加到第二组，也没有继续使用已经产生彩斑的 Face 2000 步权重。')
    docs.add_heading('为什么优先改训练目标', 1)
    docs.add_paragraph('这张 512 图片中的脸只有约 39×39 像素。在模型用于生成的网格上，脸只占 6／1024 个位置，约 0.59%。以前把整张图片的误差平均起来，脸变坏也可能被其他区域的改善掩盖。之前的诊断已经看到：整图误差略降，但脸区误差明显增大。')
    docs.add_paragraph('这次直接要求 Face 把脸区生成正确，并继续约束身份、视觉外观和脸外稳定性。身份网络用于辅助训练和评价，不能用“识别分数提高”替代肉眼确认脸部正常。')
    docs.add_heading('比较范围与固定设置', 1)
    pair = selection.get('selected', {})
    appearance = pair.get('appearance', body['metadata'].get('source_appearance', '001_13'))
    docs.add_paragraph(f'本轮仍是训练集中的一个固定案例，外观编号 {appearance}，A 与 Target 是同一人物的不同动作。两组各反复使用这一对图片 300 次。表格的 n=1 表示只有一张可靠可评的生成脸，不是一个大规模测试结果。最终 50 对测试样本没有参与。')
    docs.add_paragraph('本轮修复训练和主要出图为 512×512；随后另做同一案例的 1024 推理检查，没有进行 1024 训练。所有出图均为 30 步生成采样、种子 42、true CFG 1.0，控制强度为 1。图中的 100／300 是训练更新次数；每张结果仍单独进行完整的 30 步生成。只给 Qwen 完整 A 图，A 提供身份与外观，B 只提供动作和脸部几何，Target RGB 只作训练监督及离线评价。')
    docs.add_heading('完整出图结果', 1)
    metric_rows = [body, old300, old2000, roi300, aligned300]
    labels = ['Body', '旧 Face 300', '旧 Face 2000', 'ROI Face 300', '对齐 ROI 300']
    metric_table(docs, metric_rows, labels)
    docs.add_paragraph('↑ 表示越大越好，↓ 表示越小越好；每行加粗实际最优值，相同最优值同时加粗。源脸相似度比较生成脸与 A；可靠指标要求人脸检测有效。彩斑可能仍被检测为人脸，因此可靠检测也不能替代图像检查。')
    docs.add_paragraph('以上数值只说明这一个训练案例。整图 PSNR、SSIM 能补充检查画面变化，但小面积的人脸错误仍可能被整图指标掩盖；判断修复同时看放大的脸、身份分数、姿态和背景。')
    docs.add_heading('1024 推理检查：还需要修正 Face 控制映射', 1)
    if high_resolution['body'] and high_resolution['native_grid']:
        high_rows = [high_resolution['body'], high_resolution['native_grid']]
        high_labels = ['1024 Body', '512 ROI 权重直用 1024']
        if high_resolution['checkpoint_grid']:
            high_rows.append(high_resolution['checkpoint_grid'])
            high_labels.append('1024 控制网格接口修复')
        high_common_n = int(all(finite(row['metrics'].get('identity_cosine_reliable')) for row in high_rows))
        metric_table(docs, high_rows, high_labels)
        body_raw = high_resolution['body']['metrics'].get('identity_cosine_raw')
        direct_raw = high_resolution['native_grid']['metrics'].get('identity_cosine_raw')
        docs.add_paragraph(f'将同一个 512 ROI Face 300 步权重直接用于 1024 后，图中再次出现绿色点状失真和面部重影。原始源脸分数从 Body 的 {number(body_raw)} 变为 {number(direct_raw)}；这些原始分数保留用于追溯，不能作为可靠的身份优劣结论。')
        docs.add_paragraph(f'人脸评价仍沿用共同的 512 图像网格和原来 32 像素的短边门槛。本次 1024 Body 和直接迁移结果的人脸短边在该网格上均不足 32 像素，未达到可靠评价标准。当前高分辨率方案的共同可靠身份样本 n={high_common_n}；n=0 表示没有足够可靠的共同脸可供身份比较，不能说身份一定提升或下降。肉眼看到的绿点与重影仍应作为失败现象记录。')
        if high_resolution['checkpoint_grid']:
            fixed = high_resolution['checkpoint_grid']
            docs.add_paragraph('已额外保存控制网格接口修复后的同一案例完整生成与指标，列在上表最后一列；它沿用原 512 检查点，没有增加训练或更换人物。该结果用于检查接口修复后的变化，是否消除可见失真仍需结合对应放大图与逐图记录判断；若共同可靠 n=0，继续保留“无法可靠判断身份”的结论。')
            docs.add_paragraph(f"接口修复后的原始源脸参考分为 {number(fixed['metrics'].get('identity_cosine_raw'))}，可靠有效数 n={int(finite(fixed['metrics'].get('identity_cosine_reliable')))}。单独检出可靠脸也不能补成与不可靠 Body 的共同身份比较。")
            if acceptance.get('1024', {}).get('checkpoint_grid_manual'):
                docs.add_paragraph('修正方法是先按 Face 训练时的网格计算控制，再平滑映射到实际生成尺寸，并限制在目标脸框内；Body 继续使用实际尺寸。人工逐图检查已未见原先明显的绿点和严重重影，姿态与背景保持稳定；这次仅检查推理，没有增加 1024 训练。')
        else:
            docs.add_paragraph('1024 控制网格接口修复后的完整生成与指标尚未保存，本报告将这一项标为“未验证”；不预先宣称高分辨率问题已经解决。')
    else:
        high_common_n = None
        docs.add_paragraph('1024 的 Body 对照或直接迁移完整记录尚不齐全，当前仅保留 512 修复结论；高分辨率验收标为“未验证”。')
    if highest_peak_ratio is not None:
        docs.add_paragraph(f'CPU 控制分支审计发现：同一套权重、同样的面部特征和控制时刻，从 512 对应的较粗网格改为 1024 的细网格直接采样，出现此前较粗网格未触达的局部高幅位置，最大的局部残差尖峰约为原来的 {highest_peak_ratio:.2f} 倍。这提示高分辨率控制映射需要单独检查与适配，不表示整张脸的控制都放大了相同倍数，也不能只凭这一项确定所有失真的原因。')
    docs.add_paragraph('上述 1024 结果仍来自同一对训练图片，只是 512 短训权重的高分辨率推理检查；不能称为已经训练好的 1024 Face 模型。512 上两组 300 步修复结果及其 n=1 的单案例结论继续保留。')
    docs.add_heading('观察结论与验收', 1)
    conclusions = review.get('conclusions', review.get('summary', []))
    if conclusions:
        paragraphs(docs, conclusions)
    else:
        docs.add_paragraph('人工检查结论尚未写入记录；本报告展示已保存的结果，不预先认定对齐方案一定更好，也不根据分数单独宣称彩斑已消失。')
    if acceptance:
        paragraphs(docs, acceptance.get('conclusions', acceptance.get('explanation', [])))
        if isinstance(acceptance.get('checks'), dict):
            entries = []
            for name, value in acceptance['checks'].items():
                if isinstance(value, bool):
                    entries.append([name, '通过' if value else '未通过'])
                elif isinstance(value, str):
                    entries.append([name, value])
            if entries:
                add_table(docs, ['检查记录', '结果'], entries, font_size=9)
    docs.add_heading('训练诊断与资源', 1)
    probe_rows = []
    regional = read(previous / 'regional_probe.json', [])
    for label, step in [('Body／Face 0', 0), ('旧 Face 300', 300), ('旧 Face 2000', 2000)]:
        matches = [entry for entry in regional if entry['added_step'] == step]
        if matches:
            probe_rows.append((label, probe_mean({'rows': matches[0]['fixed_noise_rows']}, 'face_region_flow_mse')))
    for run, label in [('roi_same_aux', 'ROI Face 300'), ('roi_aligned_aux', '对齐 ROI Face 300')]:
        probe = read(root / run / 'diagnostic_probe_000300.json', {})
        probe_rows.append((label, probe_mean(probe, 'face_region_flow_mse')))
    if probe_rows:
        values = [value for _, value in probe_rows]
        table = add_table(docs, ['诊断项', *[label for label, _ in probe_rows]],
                          [['固定 8 组噪声的脸区误差 ↓', *[number(value) for value in values]]], font_size=8)
        bold_best(table, 1, values, higher=False)
        docs.add_paragraph('这里的噪声诊断包含带噪 Target，目的是检查训练目标是否改善；它不是真正从 A 开始的完整生成，也不能替代上面的出图验收。所有方案比较的是同一脸区误差，而不是混用不同定义的总损失。')
    resource_rows = []
    for run, label in [('roi_same_aux', 'ROI Face'), ('roi_aligned_aux', '对齐 ROI Face')]:
        summary = summaries[run]
        resource_rows.append([
            label, summary['optimizer_steps'], f"{summary['seconds']/60:.1f} 分钟",
            f"{summary['peak_cuda_gib']:.1f} GiB",
            '通过' if summary.get('frozen_parameters_unchanged') and summary.get('base_versions_unchanged') else '未通过',
        ])
    add_table(docs, ['实验', '新增训练步', '训练耗时', '训练峰值显存', '冻结参数'], resource_rows, font_size=9)
    docs.add_paragraph('部分中间结果在训练继续运行时生成，因此各检查点的出图耗时不作为速度优劣比较。完整图片、检查点、配置和逐项数据另存于实验目录，旧结果与默认权重保留。')
    if protection:
        paragraphs(docs, protection.get('explanation', []))
    docs.add_heading('这次结果能说明什么', 1)
    docs.add_paragraph('若脸恢复正常且身份、姿态和背景满足验收，只能说明这一个案例上的问题得到改善。下一步仍需在固定验证样本中确认其他人物和转头情况，最后再用未参与选择的测试集评价。此次没有完成整个 iPER 训练，也没有证明泛化能力已经提高。')
    paragraphs(docs, review.get('next_steps', []))
    docs.add_heading('后续 Face 训练应如何使用修复', 1)
    docs.add_paragraph('新的修复训练应显式启用脸区生成目标和五点对齐：--face-flow-loss roi --face-aux-alignment five_point。前者让训练直接重视脸部误差，后者统一辅助身份与感知网络的脸部位置；这两个选项不删除或替换完整 Face Adapter。')
    docs.add_paragraph('本报告中的 ROI Face 对照使用 roi 与 box，保留旧方框裁剪；“对齐 ROI Face”使用 roi 与 five_point。两者是独立训练的候选。历史选项继续用于复现旧结果，新的修复实验应记录所用选项。512 权重用于 1024 时，本次已检查的接口模式为 --face-scatter-grid checkpoint；其他人物和转头情况仍需验证。')

    section = docs.add_section(WD_SECTION.NEW_PAGE)
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Inches(11.69), Inches(8.27)
    section.left_margin = section.right_margin = Inches(.35)
    section.top_margin = section.bottom_margin = Inches(.40)
    groups = [
        ('reference', '固定训练案例：人物身份与目标动作',
         [reference_row(body, 'source'), reference_row(body, 'target'), body],
         ['A：身份、衣服与背景', 'Target：监督与目标动作', '固定 Body 输出']),
        ('roi_repair', '改变脸区训练目标之后，完整出图发生了什么',
         [body, old300, roi300], ['Body', '旧 Face：新增 300 步', 'ROI Face：新增 300 步']),
        ('alignment_comparison', '对齐辅助训练是否带来额外收益',
         [body, roi300, aligned300], ['Body', 'ROI Face：新增 300 步', '对齐 ROI：新增 300 步']),
        ('learning_steps', '同一案例的学习过程：每列均为完整生成',
         [roi100, roi300, aligned100, aligned300],
         ['ROI Face：训练 100 步', 'ROI Face：训练 300 步', '对齐 ROI：训练 100 步', '对齐 ROI：训练 300 步']),
    ]
    panels = []
    for index, (name, title, rows, image_labels) in enumerate(groups):
        if index:
            docs.add_page_break()
        docs.add_heading(title, 1)
        png, jpg = panel(root / 'report_panels', name, rows, image_labels)
        with Image.open(jpg) as image:
            width = min(10.8, 5.95 * image.width / image.height)
        docs.add_picture(str(jpg), width=Inches(width))
        caption = docs.add_paragraph('第一行是整图，第二行是对应脸部裁剪并放大；每列上下宽度相同。放大不增加原始细节。本页全部来自这一对训练图片，生成采样始终为 30 步。')
        caption.paragraph_format.space_after = Pt(0)
        for run in caption.runs:
            run.font.size = Pt(8)
        panels.append({'name': name, 'png_original': str(png), 'embedded_jpeg': str(jpg),
                       'source_images': [row['image'] for row in rows]})
    existing_resolution_panels = [
        ('resolution1024_review', '1024 直接迁移：保留可见失败现象',
         root / 'report_panels' / 'resolution1024_review.png',
         '第一行是整图，第二行是对应放大脸图，上下列宽相同。这里使用 512 短训权重进行 1024 推理，没有 1024 训练；直接迁移出现的绿点与重影不能忽略。共同可靠身份样本 n=0，原始身份分只作为参考。'),
        ('resolution1024_fixed_review', '1024 网格接口修复后的检查',
         root / 'report_panels' / 'resolution1024_fixed_review.png',
         '仍是同一对训练图片、同一个 512 检查点的 1024 推理检查。第一行整图，第二行对应脸部放大；这张图用于检查接口修复后的实际变化，不代表泛化或完整 iPER 训练完成。'),
    ]
    for name, title, image_path, caption_text in existing_resolution_panels:
        docs.add_page_break()
        docs.add_heading(title, 1)
        if not image_path.is_file():
            docs.add_paragraph('对应完整对照图尚未保存，该图像检查标为“未验证”。')
            continue
        if name == 'resolution1024_fixed_review' and not high_resolution['checkpoint_grid']:
            docs.add_paragraph('对应接口修复指标尚未保存，该结果的验收标为“未验证”。')
        with Image.open(image_path) as image:
            width = min(10.8, 5.95 * image.width / image.height)
        docs.add_picture(str(image_path), width=Inches(width))
        caption = docs.add_paragraph(caption_text)
        caption.paragraph_format.space_after = Pt(0)
        for run in caption.runs:
            run.font.size = Pt(8)
        panels.append({'name': name, 'png_original': str(image_path), 'embedded_image': str(image_path),
                       'image_sha256': hashlib.sha256(image_path.read_bytes()).hexdigest(),
                       'source_metrics': [str(path) for path in high_resolution_paths.values() if path.is_file()]})
    args.output.parent.mkdir(parents=True, exist_ok=True)
    docs.save(args.output)
    checked = Document(args.output)
    if len(checked.inline_shapes) != len(panels):
        raise AssertionError('Some panel images were not embedded into Word')
    if any(shape.height > Inches(6.3) for shape in checked.inline_shapes):
        raise AssertionError('Embedded panel exceeds the 6.3 inch page-height limit')
    manifest = {
        'word': str(args.output), 'word_sha256': hashlib.sha256(args.output.read_bytes()).hexdigest(),
        'bytes': args.output.stat().st_size, 'embedded_panels': panels,
        'training_cases': 1, 'final_50_test_pairs_used': False,
        'completed_new_face_steps': {run: summaries[run]['optimizer_steps'] for run in RUNS},
        'resolution_1024': {
            'training_performed': False,
            'common_reliable_identity_n': high_common_n,
            'checkpoint_grid_complete_generation_available': high_resolution['checkpoint_grid'] is not None,
            'metrics': {name: str(path) for name, path in high_resolution_paths.items() if path.is_file()},
            'face_scatter_audit': str(scatter_audit_path) if scatter_audit_path.is_file() else None,
            'max_local_peak_ratio_in_cpu_audit': highest_peak_ratio,
        },
    }
    (root / 'report_manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2))
    print(args.output)


if __name__ == '__main__':
    main()
